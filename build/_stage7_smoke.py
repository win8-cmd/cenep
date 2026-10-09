"""阶段 7 冒烟脚本（临时，不属于产品代码）。"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))

from openpyxl import load_workbook  # noqa: E402

from cenep.calculation.engine import calculation_engine  # noqa: E402
from cenep.reports.excel_exporter import SHEET_NAMES, ExcelExporter  # noqa: E402

import test_v2_integration as t  # noqa: E402


def main() -> None:
    project = t.build_project()
    result = calculation_engine.calculate(project)
    target = Path(__file__).resolve().parent / "stage7_empty.xlsx"
    out = ExcelExporter().export(project, result, target)
    wb = load_workbook(out)
    print("declared", len(SHEET_NAMES), "actual", len(wb.sheetnames), wb.sheetnames == SHEET_NAMES)
    bad = [
        f"{ws.title}!{c.coordinate}"
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    print("formulas:", bad)
    for name in ("月度电费分析", "负荷数据质量", "电价版本与来源", "方案电费对比", "计算假设与警告"):
        ws = wb[name]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value]
        print("---", name, len(texts))
        for line in texts[:6]:
            print("    ", line[:110])


if __name__ == "__main__":
    main()
