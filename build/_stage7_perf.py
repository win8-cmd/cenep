"""V2.4 阶段 7 性能实测脚本（规格书 §8.3；临时工具，不属于产品代码）。

用法::

    python build/_stage7_perf.py            # 打印 JSON 实测结果

测量项（全部为**实测**，不使用估算值）：

1. 单年 8760 小时时序仿真（§8.3 第一条：< 2 秒，V2.0 已建立的 §86 预算）；
2. 35040 个 15 分钟点的基础消纳计算（§8.3 第二条：< 5 秒）；
3. 四场景账单对比（V2.3 §7.1）；
4. 账单复算与校准（12 张账单，V2.3 §7.4）；
5. 电价计划时段检测与校验（V2.3 §4.2、§7.3）；
6. **大项目**加载 / 保存耗时与文件体积（35040 点曲线 + 12 张账单）；
7. 25 年全周期仿真与 100 候选扫描推算（§86 的 30 秒口径）。
"""

from __future__ import annotations

import json
import sys
import tempfile
import time
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tests"))

import numpy as np  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from cenep.application.scenario_service import ScenarioBillService  # noqa: E402
from cenep.application.tariff_service import TariffService  # noqa: E402
from cenep.calculation import economic_v2 as econ  # noqa: E402
from cenep.calculation.engine import calculation_engine  # noqa: E402
from cenep.calculation.self_consumption import analyze_self_consumption  # noqa: E402
from cenep.calculation.tariff_plan_engine import (  # noqa: E402
    effective_period_prices,
    period_hours_of,
    validate_tariff_plan,
)
from cenep.domain.bill_models import BillSourceType  # noqa: E402
from cenep.domain.enums import LoadDataSourceType, Resolution  # noqa: E402
from cenep.domain.load_data import HighFrequencyLoadDataset, TimeSeriesPoint  # noqa: E402
from cenep.domain.models import Project  # noqa: E402
from cenep.domain.timeseries import Resolution  # noqa: E402
from cenep.infrastructure.project_file import load_project, save_project  # noqa: E402
from cenep.policy.tariff_plan_store import TariffPlanStore  # noqa: E402
from cenep.reports.excel_exporter import ExcelExporter  # noqa: E402
from cenep.reports.pdf_exporter import PdfExporter  # noqa: E402

OFFICIAL_110_ID = "HUBEI_COMMERCIAL_2026_01_TWO_PART_KV110"

RESULTS: dict[str, object] = {}


def _timed(key: str, fn, repeat: int = 1) -> object:
    """跑 ``repeat`` 次取**最快一次**（避免后台进程抖动影响判定），并记录耗时。"""
    best = float("inf")
    value = None
    for _ in range(repeat):
        start = time.perf_counter()
        value = fn()
        best = min(best, time.perf_counter() - start)
    RESULTS[key] = {"seconds": round(best, 6), "repeat": repeat}
    print(f"{key}: {best:.6f} s")
    return value


def _build_axis(project: Project, resolution: Resolution, year: int = 2025):
    return econ.build_time_axis(year, resolution, ())


def _load_series(axis, pv: bool) -> np.ndarray:
    n = axis.point_count
    if pv:
        return np.array(
            [
                (
                    600.0 * max(0.0, np.sin(np.pi * ((i % 24) - 6) / 12))
                    if 6 <= (i % 24) <= 18
                    else 0.0
                )
                for i in range(n)
            ],
            dtype=float,
        )
    return np.array([400.0 if 7 <= (i % 24) <= 18 else 120.0 for i in range(n)], dtype=float)


def _stage_project(base_year: int = 2025) -> Project:
    project = Project()
    project.analysis_period = 1
    project.basic_info.project_name = "阶段 7 性能项目"
    project.timeseries.enabled = True
    project.timeseries.base_year = base_year
    project.pv.pv_capacity_kwp = 1000.0
    project.pv.equivalent_hours = 1100.0
    project.pv.performance_ratio = 1.0
    project.storage.storage_power_kw = 500.0
    project.storage.storage_energy_kwh = 2000.0
    project.timeseries.dispatch.soc_min = 0.10
    project.timeseries.dispatch.soc_max = 1.00
    project.timeseries.load.annual_energy_kwh = 1_500_000.0
    project.load.annual_load_kwh = 1_500_000.0
    project.timeseries.pv.equivalent_hours = 1100.0
    return project


def _dataset(project: Project, interval_minutes: int, points: int, profile_id: str) -> HighFrequencyLoadDataset:
    """受控负荷数据集（白天 400 kW / 夜间 120 kW，按间隔换算成 kWh）。"""
    delta_hours = interval_minutes / 60.0
    step = timedelta(minutes=interval_minutes)
    start = datetime(2025, 1, 1)
    stamps = [start + step * i for i in range(points)]
    load_points = []
    for index in range(points):
        hour = (index * interval_minutes // 60) % 24
        power = 400.0 if 7 <= hour <= 18 else 120.0
        load_points.append(TimeSeriesPoint(timestamp=stamps[index], load_kwh=power * delta_hours))
    return HighFrequencyLoadDataset(
        profile_id=profile_id,
        name=f"{interval_minutes} 分钟受控负荷",
        source_type=LoadDataSourceType.HIGH_FREQUENCY_IMPORT,
        resolution=Resolution.from_interval_minutes(interval_minutes) or Resolution.HOURLY,
        interval_minutes=interval_minutes,
        period_start=stamps[0],
        period_end=stamps[-1],
        points=load_points,
        estimated=False,
        coverage_ratio=1.0,
        missing_intervals=0,
    )


def _bills(project: Project, count: int = 12, year: int = 2026) -> None:
    """``count`` 张真实结构（受控数值）的月账单，用于复算与文件体积测量。

    全部落在 **2026-01**：湖北官方 2026-01 版电价计划的生效区间就是
    2026-01-01 ~ 2026-01-31；复算按**计费日期**选取版本，跨月会被正确拒绝
    （V2.3 §4.2 T06：早于生效日或已过期都不得用于正式复算）。
    多月账单用不同计量点区分（同一计量点同账期会触发重复判定）。
    """
    service = __import__(
        "cenep.application.bill_service", fromlist=["BillService"]
    ).BillService(project)
    start = datetime(year, 1, 1).date()
    end = datetime(year, 1, 31).date()
    for index in range(1, count + 1):
        service.create_bill(
            meter_id=f"METER-{index:03d}",
            billing_period_start=start,
            billing_period_end=end,
            voltage_level="交流 10kV",
            tariff_structure="two_part",
            energy_total_kwh=7_000_000.0,
            energy_sharp_kwh=600_000.0,
            energy_peak_kwh=1_500_000.0,
            energy_flat_kwh=3_400_000.0,
            energy_valley_kwh=1_500_000.0,
            energy_charge_yuan=5_000_000.0,
            basic_capacity_charge_yuan=200_000.0,
            government_fund_charge_yuan=130_000.0,
            bill_total_yuan=5_330_000.0,
            billing_demand_kw=3_500.0,
            contract_capacity_kva=5_000.0,
            source_type=BillSourceType.MANUAL,
        )


def main() -> None:
    RESULTS["environment"] = {
        "python": sys.version.split()[0],
        "platform": sys.platform,
    }
    import platform as _platform

    RESULTS["environment"]["impl"] = _platform.python_implementation()

    # ---------- 1. 单年 8760 小时时序仿真（§8.3 第 1 条） ----------
    project = _stage_project()
    axis_hourly = _build_axis(project, Resolution.HOURLY)
    RESULTS["axis_hourly_points"] = axis_hourly.point_count
    sim = _timed(
        "single_year_8760_hourly_simulation",
        lambda: econ.simulate_year(project, axis_hourly, 1),
        repeat=3,
    )
    RESULTS["single_year_8760_points"] = sim.axis.point_count

    # ---------- 2. 35040 个 15 分钟点的基础消纳计算（§8.3 第 2 条） ----------
    axis_15 = _build_axis(project, Resolution.QUARTER_HOURLY)
    RESULTS["axis_15min_points"] = axis_15.point_count
    load_15 = _load_series(axis_15, pv=False)
    pv_15 = _load_series(axis_15, pv=True)

    def _consume():
        return analyze_self_consumption(
            load_15,
            pv_15,
            interval_minutes=15,
            timestamps=list(axis_15.timestamps),
            load_source_type=LoadDataSourceType.HIGH_FREQUENCY_IMPORT,
        )

    consumption = _timed("consumption_35040_points_15min", _consume, repeat=3)
    RESULTS["consumption_self_consumption_rate"] = consumption.self_consumption_rate
    RESULTS["consumption_max_balance_error_kwh"] = consumption.max_interval_balance_error_kwh

    # ---------- 2b. 35040 点全项目时序仿真（信息项，非 §8.3 硬指标） ----------
    sim15 = _timed(
        "single_year_35040_15min_simulation",
        lambda: econ.simulate_year(project, axis_15, 1),
        repeat=2,
    )
    RESULTS["single_year_35040_points"] = sim15.axis.point_count

    # ---------- 3. 四场景账单对比（V2.3 §7.1） ----------
    store = TariffPlanStore()
    plan = store.get_plan(OFFICIAL_110_ID)
    client = _stage_project()
    client.analysis_period = 1
    from cenep.application.load_profile_service import LoadProfileService

    load_service = LoadProfileService(client)
    load_service.register(_dataset(client, 60, 8760, "PERF-8760"))
    service = ScenarioBillService(client, store=store, load_profile_service=load_service)
    scenario = _timed("four_scenario_bill_compare_8760", lambda: service.compare(OFFICIAL_110_ID), repeat=2)
    RESULTS["four_scenario_unique_benefit_yuan"] = scenario.unique_annual_benefit_yuan
    RESULTS["four_scenario_dedup_verified"] = scenario.dedup_verified

    # ---------- 4. 账单复算与校准（V2.3 §7.4，12 张账单） ----------
    bill_project = _stage_project()
    _bills(bill_project)
    RESULTS["bill_count"] = len(bill_project.bills)
    tariff_service = TariffService(bill_project, store=TariffPlanStore())
    calibration = _timed(
        "bill_recompute_and_calibrate_12_bills",
        lambda: tariff_service.recompute_all(OFFICIAL_110_ID),
        repeat=3,
    )
    RESULTS["calibration_bills"] = calibration.bill_count
    RESULTS["calibration_passed"] = calibration.calibration_passed

    # ---------- 5. 电价计划时段检测与校验（V2.3 §4.2、§7.3） ----------
    def _validate_plan():
        validation = validate_tariff_plan(plan)
        prices = effective_period_prices(plan)
        hours = period_hours_of(plan)
        return validation, prices, hours

    validation, prices, hours = _timed("tariff_plan_period_detection", _validate_plan, repeat=5)
    RESULTS["tariff_plan_has_error"] = validation.has_error
    RESULTS["tariff_plan_issue_count"] = len(validation.issues)
    RESULTS["tariff_plan_period_prices"] = {str(k): v for k, v in prices.items()}
    RESULTS["tariff_plan_period_hours"] = hours

    # ---------- 6. 大项目加载 / 保存耗时与体积（35040 点 + 12 张账单） ----------
    big = _stage_project()
    _bills(big)
    LoadProfileService(big).register(_dataset(big, 15, 35040, "PERF-35040"))
    RESULTS["big_project_datasets"] = len(big.load_datasets)
    RESULTS["big_project_points"] = big.load_datasets[0].points.__len__()

    with tempfile.TemporaryDirectory() as tmp:
        target = Path(tmp) / "大项目.nep"

        def _save():
            return save_project(big, target)

        saved = _timed("big_project_save_35040_plus_12_bills", _save, repeat=3)
        RESULTS["big_project_file_bytes"] = saved.stat().st_size
        RESULTS["big_project_file_mb"] = round(saved.stat().st_size / 1024 / 1024, 3)

        loaded = _timed("big_project_load", lambda: load_project(saved), repeat=3)
        RESULTS["big_project_loaded_points"] = len(loaded.load_datasets[0].points)
        RESULTS["big_project_loaded_bills"] = len(loaded.bills)

        # 报告导出（32 张表 / 21 章）
        result_big = calculation_engine.calculate(loaded)
        excel_path = _timed(
            "big_project_excel_export",
            lambda: ExcelExporter().export(loaded, result_big, Path(tmp) / "大项目.xlsx"),
            repeat=1,
        )
        RESULTS["big_project_excel_bytes"] = excel_path.stat().st_size
        sheets = load_workbook(excel_path).sheetnames
        RESULTS["big_project_excel_sheets"] = len(sheets)

        pdf_path = _timed(
            "big_project_pdf_export",
            lambda: PdfExporter().export(loaded, result_big, Path(tmp) / "大项目.pdf"),
            repeat=1,
        )
        RESULTS["big_project_pdf_bytes"] = pdf_path.stat().st_size

    # ---------- 7. 25 年全周期 + 100 候选推算（§86 的 30 秒口径） ----------
    full = _stage_project()
    full.analysis_period = 25
    project25 = _stage_project()
    project25.analysis_period = 25
    project25.pv.annual_degradation_rate = 0.005
    project25.storage.annual_degradation_rate = 0.02
    project25.timeseries.load.annual_growth_rate = 0.02
    project25.timeseries.tariff.annual_growth_rate = 0.01
    _timed("full_period_25_year_simulation", lambda: econ.simulate_project(project25), repeat=1)

    per = RESULTS["single_year_8760_hourly_simulation"]["seconds"]
    RESULTS["projected_100_candidates_seconds"] = round(per * 100, 3)
    RESULTS["budgets"] = {
        "single_year_8760_budget_s": 2.0,
        "consumption_35040_budget_s": 5.0,
        "scan_100_candidates_budget_s": 30.0,
    }
    print(json.dumps(RESULTS, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
