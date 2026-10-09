"""电费账单 Excel 模板、列映射与导入预览（V2.1 §5.3、§5.5、§9.1）。

职责（V2.1 §1：``data/`` 负责读取、列映射、质量检查、导入预览）
----------------------------------------------------------------
1. **模板生成**：``CENEP_电费账单导入模板.xlsx``，含 ``填写说明`` / ``月账单`` /
   ``分时电价`` / ``数据字典`` 四张工作表（§5.3），枚举字段带下拉选项，
   示例行明确标记为"示例行"且默认**不导入**。
2. **列映射**：支持用户重命名列、中英文列名、列顺序变化（§5.5）。
3. **导入预览**：逐行给出"工作表 + 行号 + 字段 + 原因"的中文问题，
   并统计识别到的记录数与有效/警告/无效数量（§5.3、§5.5）。
4. **重复识别**：以"项目 + 账期 + 计量点"组合识别潜在重复（§5.5），
   由 :func:`apply_bill_import` 按 跳过 / 替换 / 保留 处理。

设计边界
--------
* 本模块**不做电价计算**，也不预填任何电价数值（V2.1 §0.2、§4.1）。
  模板中的 ``分时电价`` 工作表只提供**表头框架**：V2.1 的导入器不消费它，
  留待 V2.3 的版本化电价计划（§2.3、阶段 5）。预览会对该工作表有数据的情况给出提示。
* 单元格解析失败一律产出**中文行级错误**，不抛裸异常（§0.2）。
* 复用 :func:`cenep.data.importer.normalize_header` / :func:`parse_timestamp` /
  :func:`list_sheets` / :func:`read_table_from_sheet`，不另建读取与表头归一化实现。
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from pydantic import Field, ValidationError as PydanticValidationError

from ..calculation.bill_calculator import apply_quality
from ..calculation.errors import ValidationError
from ..domain.base import _Model
from ..domain.bill_models import (
    MONTH_FORMAT,
    BillTolerance,
    ElectricityBill,
    label_of,
    make_bill_id,
)
from ..domain.enums import BillQualityStatus, BillSourceType, DuplicateStrategy, TariffStructure
from .importer import list_sheets, normalize_header, parse_timestamp, read_table_from_sheet

logger = logging.getLogger(__name__)

__all__ = [
    "BILL_COLUMNS",
    "BILL_COLUMN_ALIASES",
    "DEFAULT_SHEET",
    "EXAMPLE_MARKER",
    "MARKER_COLUMN",
    "REQUIRED_BILL_FIELDS",
    "SHEET_BILLS",
    "SHEET_DICT",
    "SHEET_HELP",
    "SHEET_TARIFF",
    "TEMPLATE_FILE_NAME",
    "VOLTAGE_LEVEL_OPTIONS",
    "BillColumn",
    "BillImportPreview",
    "BillImportResult",
    "BillImportRow",
    "ColumnMapping",
    "apply_bill_import",
    "bill_template_bytes",
    "build_bill_from_row",
    "build_bill_template",
    "parse_bill_source",
    "parse_choice",
    "parse_tariff_structure",
    "preview_bill_import",
    "resolve_bill_columns",
]

# --------------------------------------------------------------------------- #
# 模板常量
# --------------------------------------------------------------------------- #
#: 模板文件名（V2.1 §5.3 规定）
TEMPLATE_FILE_NAME = "CENEP_电费账单导入模板.xlsx"
#: 工作表名（§5.3）
SHEET_HELP = "填写说明"
SHEET_BILLS = "月账单"
SHEET_TARIFF = "分时电价"
SHEET_DICT = "数据字典"
#: 默认导入的工作表
DEFAULT_SHEET = SHEET_BILLS
#: 示例行标记列与标记文本（§5.3：示例行不得被默认当作真实数据导入）
MARKER_COLUMN = "数据标记"
EXAMPLE_MARKER = "示例行"

#: 必填字段：只有账期是必需的——账单可能只提供总额、也可能只提供分时电量（§2.1）
REQUIRED_BILL_FIELDS: tuple[str, ...] = ("billing_period_start", "billing_period_end")

#: 电压等级下拉候选（仅"等级名"提示，**不含任何电价数值**）
VOLTAGE_LEVEL_OPTIONS: tuple[str, ...] = (
    "不满1kV",
    "1-10kV",
    "20kV",
    "35kV",
    "110kV",
    "220kV及以上",
)


@dataclass(frozen=True)
class BillColumn:
    """模板列定义：模板表头、模型字段、单位、必填与说明（模板与映射的唯一来源）。"""

    field: str
    header: str
    unit: str = ""
    required: bool = False
    description: str = ""
    aliases: tuple[str, ...] = dataclass_field(default_factory=tuple)

    @property
    def label(self) -> str:
        """中文列名（含单位时原样返回表头）。"""
        return self.header


#: ``月账单`` 工作表的列（顺序即模板列顺序，§5.3）
BILL_COLUMNS: tuple[BillColumn, ...] = (
    BillColumn("", MARKER_COLUMN, "", False, "示例行请填『示例行』；正式数据请清空本列", ("标记", "数据类型")),
    BillColumn(
        "billing_period_start", "账期起(YYYY-MM-DD)", "日期", True,
        "账单计费周期起始日（含）。例：2026-01-01", ("账期开始", "起始日期", "开始日期", "billingstart", "period_start"),
    ),
    BillColumn(
        "billing_period_end", "账期止(YYYY-MM-DD)", "日期", True,
        "账单计费周期结束日（含）。跨月账期如实填写，系统会标注跨月", ("账期结束", "结束日期", "billingend", "period_end"),
    ),
    BillColumn("billing_month", "账单月份(YYYY-MM)", "YYYY-MM", False, "留空则按账期起始日推导", ("月份", "month")),
    BillColumn("meter_id", "计量点编号", "", False, "用于重复判定（项目 + 账期 + 计量点）", ("计量点", "表号", "meter")),
    BillColumn("customer_name", "客户名称", "", False, "账单户名", ("户名", "客户", "用户名")),
    BillColumn("voltage_level", "电压等级", "", False, "用于后续电价匹配，如 10kV", ("供电电压", "电压")),
    BillColumn("tariff_structure", "计费方式", "", False, "单一制 / 两部制 / 未知", ("电价制式", "tariffstructure")),
    BillColumn("contract_capacity_kva", "合同容量(kVA)", "kVA", False, "两部制按容量计费时的合同容量", ("合同容量kva", "容量")),
    BillColumn(
        "billing_demand_kw", "账单计费需量(kW)", "kW", False,
        "**账单上的计费需量**，不是负荷曲线最大值（V2.1 §2.1）", ("计费需量", "需量", "demand"),
    ),
    BillColumn("energy_total_kwh", "总购电量(kWh)", "kWh", False, "账单总电量", ("总电量", "购电量", "用电量合计")),
    BillColumn("energy_sharp_kwh", "尖峰电量(kWh)", "kWh", False, "尖峰时段电量", ("尖峰",)),
    BillColumn("energy_peak_kwh", "高峰电量(kWh)", "kWh", False, "高峰时段电量", ("高峰",)),
    BillColumn("energy_flat_kwh", "平段电量(kWh)", "kWh", False, "平段电量", ("平段", "平电量")),
    BillColumn("energy_valley_kwh", "低谷电量(kWh)", "kWh", False, "低谷电量", ("低谷", "谷电量")),
    BillColumn(
        "energy_offpeak_kwh", "深谷电量(kWh)", "kWh", False,
        "深谷电量（**不是低谷**：低谷请填上一列）", ("深谷", "低谷深谷"),
    ),
    BillColumn("energy_charge_yuan", "电度电费合计(元)", "元", False, "电度电费合计；若填了下面 5 个明细列，二者应相符", ("电度电费", "电量电费")),
    BillColumn("market_purchase_charge_yuan", "代理购电电费(元)", "元", False, "电度电费明细：市场化 / 代理购电电费", ("购电费", "电能量电费")),
    BillColumn("transmission_distribution_charge_yuan", "输配电费(元)", "元", False, "电度电费明细：输配电费", ("输配电",)),
    BillColumn("line_loss_charge_yuan", "线损费(元)", "元", False, "电度电费明细：线损费", ("线损",)),
    BillColumn("system_operation_charge_yuan", "系统运行费(元)", "元", False, "电度电费明细：系统运行费", ("系统运行",)),
    BillColumn("government_fund_charge_yuan", "政府性基金及附加(元)", "元", False, "电度电费明细：政府性基金及附加", ("政府性基金", "基金及附加")),
    BillColumn("basic_capacity_charge_yuan", "基本电费-按容量(元)", "元", False, "两部制按容量计费的基本电费", ("容量电费", "基本电费容量")),
    BillColumn("demand_charge_yuan", "基本电费-按需量(元)", "元", False, "两部制按需量计费的基本电费", ("需量电费", "基本电费需量")),
    BillColumn("power_factor_adjustment_yuan", "功率因数调整电费(元)", "元", False, "允许负值（返还 / 奖励）", ("力率调整", "功率因数")),
    BillColumn("other_charge_yuan", "其他费用(元)", "元", False, "其他费用", ("其它费用",)),
    BillColumn("vat_yuan", "增值税(元)", "元", False, "增值税", ("税金", "税额")),
    BillColumn("adjustment_charge_yuan", "调整/补退费(元)", "元", False, "允许负值（V2.1 §2.1）", ("补退费", "退补电费")),
    BillColumn("bill_total_yuan", "账单总额(元)", "元", False, "账单应收总额（以账单为准）", ("总额", "应收电费", "电费合计")),
    BillColumn(
        "source_type", "数据来源", "", False,
        "手动录入 / Excel导入 / 估算；Excel 导入的行一律记为 excel", ("来源", "数据类型", "sourcetype"),
    ),
    BillColumn("notes", "备注", "", False, "自由填写", ("说明", "note")),
)

#: ``分时电价`` 工作表的列（**只建表头，不参与 V2.1 导入**；§5.3、阶段 5 再实装）
TARIFF_SHEET_COLUMNS: tuple[tuple[str, str], ...] = (
    (MARKER_COLUMN, "示例行请填『示例行』"),
    ("账期起(YYYY-MM-DD)", "与月账单对应"),
    ("账期止(YYYY-MM-DD)", "与月账单对应"),
    ("计量点编号", "与月账单对应"),
    ("尖峰单价(元/kWh)", "账单实际尖峰单价，可选填"),
    ("高峰单价(元/kWh)", "账单实际高峰单价，可选填"),
    ("平段单价(元/kWh)", "账单实际平段单价，可选填"),
    ("低谷单价(元/kWh)", "账单实际低谷单价，可选填"),
    ("深谷单价(元/kWh)", "账单实际深谷单价，可选填"),
    ("备注", "自由填写"),
)

#: 字段 → 可接受的列名别名（已归一化）。模板表头、字段名本身与声明的别名都可用。
BILL_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    column.field: tuple(
        dict.fromkeys(
            [
                normalize_header(column.header),
                normalize_header(column.field),
                *(normalize_header(alias) for alias in column.aliases),
            ]
        )
    )
    for column in BILL_COLUMNS
    if column.field
}

#: 计费方式的可读文本（下拉与解析共用；**不含任何电价数值**）
TARIFF_STRUCTURE_TEXT: dict[str, TariffStructure] = {
    "单一制": TariffStructure.SINGLE_PART,
    "singlepart": TariffStructure.SINGLE_PART,
    "single_part": TariffStructure.SINGLE_PART,
    "两部制": TariffStructure.TWO_PART,
    "twopart": TariffStructure.TWO_PART,
    "two_part": TariffStructure.TWO_PART,
    "未知": TariffStructure.UNKNOWN,
    "unknown": TariffStructure.UNKNOWN,
}

#: 数据来源的可读文本
BILL_SOURCE_TEXT: dict[str, BillSourceType] = {
    "手动录入": BillSourceType.MANUAL,
    "手工录入": BillSourceType.MANUAL,
    "manual": BillSourceType.MANUAL,
    "excel": BillSourceType.EXCEL,
    "excel导入": BillSourceType.EXCEL,
    "导入": BillSourceType.EXCEL,
    "估算": BillSourceType.ESTIMATED,
    "estimated": BillSourceType.ESTIMATED,
}


# --------------------------------------------------------------------------- #
# 列映射（§5.5：支持列名重命名 / 中英文 / 列顺序变化）
# --------------------------------------------------------------------------- #
@dataclass
class ColumnMapping:
    """列映射结果：``字段 → 实际表头``。"""

    mapping: dict[str, str] = dataclass_field(default_factory=dict)
    unmapped_headers: list[str] = dataclass_field(default_factory=list)
    messages: list[str] = dataclass_field(default_factory=list)

    def header_of(self, field_name: str) -> str | None:
        return self.mapping.get(field_name)


def _match_strength(header_norm: str, aliases: tuple[str, ...]) -> tuple[int, int]:
    """返回 ``(强度, 匹配长度)``，强度：完全相等 3 > 表头含别名 2 > 表头是别名前缀 1 > 不匹配 0。

    * "表头含别名"用**最长别名**优先，避免 ``谷电量`` 这类短别名抢走 ``深谷电量``；
    * "表头是别名前缀"让用户写 ``账期起`` 也能对上 ``账期起(YYYY-MM-DD)``（真实账单常见写法），
      但要求表头至少 2 个字符，避免单字误匹配。
    """
    best = (0, 0)
    for alias in aliases:
        if not alias:
            continue
        if header_norm == alias:
            return (3, len(alias))
        if alias in header_norm and (2, len(alias)) > best:
            best = (2, len(alias))
        if len(header_norm) >= 2 and header_norm in alias and (1, len(header_norm)) > best:
            best = (1, len(header_norm))
    return best


def resolve_bill_columns(
    headers: Iterable[Any],
    *,
    required: tuple[str, ...] = REQUIRED_BILL_FIELDS,
    explicit: dict[str, str] | None = None,
) -> ColumnMapping:
    """把模板列映射到模型字段（V2.1 §5.5）。

    规则：**精确匹配优先，其次最长别名包含匹配**；用户显式指定的映射优先级最高。
    同一字段被多个表头命中时保留得分最高者，并在 ``messages`` 中给出中文提示（不静默改判）。

    :param headers: 文件的实际表头
    :param required: 必须存在的字段
    :param explicit: 用户手工指定的映射（``字段 → 表头``）
    :raises ValidationError: 表头为空或缺少必需列（中文提示并列出实际表头）
    """
    header_list = [str(h) for h in headers if h is not None and str(h).strip()]
    if not header_list:
        raise ValidationError("工作表没有可用表头，无法识别账单列", field="bill.import.header")

    mapping = ColumnMapping()
    used_headers: set[str] = set()

    # 1) 用户显式指定的映射优先
    for field_name, header in (explicit or {}).items():
        if header not in header_list:
            raise ValidationError(
                f"指定的列「{header}」在文件中不存在；可用表头：{'、'.join(header_list[:10])}",
                field=f"bill.import.{field_name}",
            )
        mapping.mapping[field_name] = header
        used_headers.add(header)

    # 2) 自动匹配
    for field_name, aliases in BILL_COLUMN_ALIASES.items():
        if field_name in mapping.mapping:
            continue
        scored = [
            (*_match_strength(normalize_header(header), aliases), header)
            for header in header_list
            if header not in used_headers
        ]
        scored = [item for item in scored if item[0] > 0]
        if not scored:
            continue
        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
        best = scored[0]
        mapping.mapping[field_name] = best[2]
        used_headers.add(best[2])
        # 只有**同强度**的竞争者才算真正歧义；精确命中后不该再提示"也像是"
        others = [item[2] for item in scored[1:] if item[0] == best[0] and item[1] == best[1]]
        if others:
            mapping.messages.append(
                f"列「{'、'.join(others)}」与「{best[2]}」都可能是「{label_of(field_name)}」，"
                f"已按「{best[2]}」识别；如需改判请在导入界面手工指定列映射"
            )

    mapping.unmapped_headers = [h for h in header_list if h not in used_headers]

    missing = [f for f in required if f not in mapping.mapping]
    if missing:
        raise ValidationError(
            "账单表缺少必需列："
            + "、".join(f"{label_of(f)}（{BILL_COLUMN_ALIASES[f][0]}）" for f in missing)
            + f"；实际表头：{'、'.join(header_list[:10])}",
            field=f"bill.import.{missing[0]}",
        )
    logger.debug("账单列映射：%s", mapping.mapping)
    return mapping


def _explicit_mapping_passthrough(explicit: dict[str, str] | None) -> dict[str, str] | None:
    """规范化用户显式映射：去掉空表头项，返回 ``None`` 表示没有显式映射。"""
    if not explicit:
        return None
    cleaned = {k: v for k, v in explicit.items() if str(v or "").strip()}
    return cleaned or None


# --------------------------------------------------------------------------- #
# 单元格解析（中文行级错误）
# --------------------------------------------------------------------------- #
def _loc(sheet: str, row: int | None, field_name: str = "") -> str:
    """错误定位前缀：``工作表「月账单」第 5 行「高峰电量(kWh)」``（§5.5）。"""
    text = f"工作表「{sheet}」第 {row} 行" if row else f"工作表「{sheet}」"
    if field_name:
        text += f"「{label_of(field_name)}」"
    return text


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _as_text(value: Any) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


_NUMBER_JUNK = ("¥", "￥", "元", "kWh", "kwh", "KWH", "kVA", "kva", "kW", "kw", ",", "，", " ", "\u3000")


def _as_float(value: Any, *, sheet: str, row: int, field_name: str, errors: list[str]) -> float | None:
    """解析数值单元格；空 → ``None``（= 未提供，**不是 0**）；非法 → 记录中文错误。"""
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        errors.append(f"{_loc(sheet, row, field_name)}：不能填 TRUE/FALSE，请填数值或留空")
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    cleaned = text
    for junk in _NUMBER_JUNK:
        cleaned = cleaned.replace(junk, "")
    cleaned = cleaned.replace("（", "").replace("）", "").replace("(", "").replace(")", "")
    try:
        return float(cleaned)
    except ValueError:
        errors.append(f"{_loc(sheet, row, field_name)}：『{text}』不是有效数值，请填数值或留空")
        return None


def _as_date(value: Any, *, sheet: str, row: int, field_name: str, errors: list[str]) -> date | None:
    """解析日期单元格（复用 :func:`parse_timestamp` 的多格式容错）。"""
    if _is_blank(value):
        errors.append(f"{_loc(sheet, row, field_name)}：账期为必填项，请填写日期（如 2026-01-01）")
        return None
    stamp = parse_timestamp(value)
    if stamp is None and isinstance(value, str):
        # 账单里常见的紧凑写法，属"账单导入"专属补充格式（不改动共享解析器）
        for fmt in ("%Y%m%d", "%Y年%m月%d日"):
            try:
                stamp = datetime.strptime(value.strip(), fmt)
                break
            except ValueError:
                continue
    if stamp is None:
        errors.append(
            f"{_loc(sheet, row, field_name)}：『{value}』不是有效日期，"
            "请使用 YYYY-MM-DD（也支持 2026/1/1、20260101、Excel 日期单元格）"
        )
        return None
    return stamp.date()


def _as_month(value: Any, *, sheet: str, row: int, field_name: str, errors: list[str]) -> str | None:
    """解析 ``YYYY-MM`` 或 ``YYYY-MM-DD``（Excel 把文本转成日期时自动归一）。"""
    if _is_blank(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime(MONTH_FORMAT)

    text = str(value).strip()
    parts = [part for part in re.split(r"[-/.年]", text) if part.strip()]
    if len(parts) >= 2 and parts[0].strip().isdigit() and parts[1].strip().isdigit():
        year, month = int(parts[0]), int(parts[1])
        if 1900 <= year <= 2200 and 1 <= month <= 12:
            return f"{year:04d}-{month:02d}"
    compact = re.sub(r"\D", "", text)
    if len(compact) in (6, 8) and compact[:4].isdigit() and 1900 <= int(compact[:4]) <= 2200:
        month = int(compact[4:6])
        if 1 <= month <= 12:
            return f"{compact[:4]}-{month:02d}"
    errors.append(
        f"{_loc(sheet, row, field_name)}：『{text}』不是有效月份，请使用 YYYY-MM（如 2026-01）或留空"
    )
    return None


def parse_choice(value: Any, table: dict[str, Any]) -> Any | None:
    """把用户填写的可读文本解析成枚举成员；无法识别返回 ``None``。

    Excel 导入与手动录入**共用**这一套口径（模板下拉、中文标签、机器值都接受），
    避免同一个枚举出现两套解析规则。
    """
    if _is_blank(value):
        return None
    if isinstance(value, (BillSourceType, TariffStructure)):
        return value
    return table.get(normalize_header(value))


def parse_tariff_structure(value: Any) -> TariffStructure | None:
    """解析计费方式（``单一制`` / ``single_part`` / ``two_part`` …）。"""
    return parse_choice(value, TARIFF_STRUCTURE_TEXT)


def parse_bill_source(value: Any) -> BillSourceType | None:
    """解析账单来源（``手动录入`` / ``excel`` / ``估算`` …）。"""
    return parse_choice(value, BILL_SOURCE_TEXT)


def _allowed_texts(table: dict[str, Any]) -> str:
    """枚举允许值的**中文**说明，用于报错提示。"""
    return "、".join(sorted({key for key in table if not key.isascii()}))


def _as_enum(
    value: Any, table: dict[str, Any], *, sheet: str, row: int, field_name: str, errors: list[str]
) -> Any:
    """按"可读文本表"解析枚举；未知取值 → 中文错误并给出允许值。"""
    if _is_blank(value):
        return None
    parsed = parse_choice(value, table)
    if parsed is not None:
        return parsed
    errors.append(
        f"{_loc(sheet, row, field_name)}：『{value}』不是有效取值，允许值：{_allowed_texts(table)}"
    )
    return None


# --------------------------------------------------------------------------- #
# 预览结果模型
# --------------------------------------------------------------------------- #
class BillImportRow(_Model):
    """导入预览中的一行（V2.1 §5.3、§5.5）。"""

    row_number: int = Field(ge=1, description="源文件中的行号（与 Excel 行号一致）")
    status: BillQualityStatus = BillQualityStatus.VALID
    bill: ElectricityBill | None = Field(default=None, description="解析成功的账单（解析失败为 None）")
    duplicate_of: str | None = Field(default=None, description="重复对象说明（已存在账单 ID 或文件内行号）")
    messages: list[str] = Field(default_factory=list, description="中文问题与提示")
    raw: dict[str, str] = Field(default_factory=dict, description="原始单元格文本（供界面显示）")

    @property
    def importable(self) -> bool:
        """是否可作为有效行导入（警告行可导入；无效行不可）。"""
        return self.status is not BillQualityStatus.INVALID and self.bill is not None


class BillImportPreview(_Model):
    """导入预览：识别到的记录数、有效/警告/无效统计与逐行问题（V2.1 §5.3、§5.5）。"""

    file_name: str
    sheet_name: str
    available_sheets: list[str] = Field(default_factory=list)
    column_mapping: dict[str, str] = Field(default_factory=dict)
    unmapped_headers: list[str] = Field(default_factory=list)
    total_rows: int = 0
    example_rows_skipped: int = 0
    valid_count: int = 0
    warning_count: int = 0
    invalid_count: int = 0
    duplicate_count: int = 0
    existing_bill_ids: list[str] = Field(
        default_factory=list,
        description="预览时项目里已有的账单编号（供『保留两条』生成不冲突的新编号）",
    )
    rows: list[BillImportRow] = Field(default_factory=list)
    messages: list[str] = Field(default_factory=list)

    @property
    def importable_rows(self) -> list[BillImportRow]:
        return [row for row in self.rows if row.importable]

    def counts_text(self) -> str:
        """一行中文统计，用于导入界面（§5.3 要求显示识别到的数量）。"""
        return (
            f"工作表「{self.sheet_name}」识别到 {self.total_rows} 条记录："
            f"有效 {self.valid_count}、警告 {self.warning_count}、无效 {self.invalid_count}"
            + (f"；疑似重复 {self.duplicate_count} 条" if self.duplicate_count else "")
            + (f"；已跳过示例行 {self.example_rows_skipped} 条" if self.example_rows_skipped else "")
        )


class BillImportResult(_Model):
    """导入执行结果（V2.1 §5.5：跳过 / 替换 / 保留）。"""

    strategy: DuplicateStrategy = DuplicateStrategy.SKIP
    added: list[ElectricityBill] = Field(default_factory=list)
    replaced: list[ElectricityBill] = Field(default_factory=list)
    kept_both: list[ElectricityBill] = Field(default_factory=list)
    skipped_ids: list[str] = Field(default_factory=list)
    invalid_rows: list[int] = Field(default_factory=list)
    messages: list[str] = Field(default_factory=list)

    @property
    def added_count(self) -> int:
        return len(self.added)

    @property
    def total_written(self) -> int:
        """实际写库的条数（新增 + 替换 + 保留）。"""
        return len(self.added) + len(self.replaced) + len(self.kept_both)

    def summary_text(self) -> str:
        return (
            f"导入完成：新增 {len(self.added)} 条、替换 {len(self.replaced)} 条、"
            f"保留重复 {len(self.kept_both)} 条、跳过 {len(self.skipped_ids)} 条、"
            f"无效行 {len(self.invalid_rows)} 行"
        )


# --------------------------------------------------------------------------- #
# 单行 → 账单
# --------------------------------------------------------------------------- #
def build_bill_from_row(
    row: dict[str, Any],
    mapping: ColumnMapping,
    *,
    sheet_name: str,
    row_number: int,
    project_id: str,
    file_name: str,
    default_source: BillSourceType = BillSourceType.EXCEL,
    tolerance: BillTolerance | None = None,
) -> tuple[ElectricityBill | None, list[str], list[str]]:
    """把一行表格解析成账单（V2.1 §2.1、§5.5）。

    :return: ``(账单或 None, 错误列表, 提示列表)``。

        * 错误列表非空 → 该行**无效**，``bill`` 为 ``None``；
        * 提示列表只是告知（例如来源列填了别的值但已按 Excel 记录），不影响有效性。
    """
    errors: list[str] = []
    notices: list[str] = []

    def cell(field_name: str) -> Any:
        header = mapping.header_of(field_name)
        return row.get(header) if header else None

    period_start = _as_date(
        cell("billing_period_start"), sheet=sheet_name, row=row_number,
        field_name="billing_period_start", errors=errors,
    )
    period_end = _as_date(
        cell("billing_period_end"), sheet=sheet_name, row=row_number,
        field_name="billing_period_end", errors=errors,
    )
    if errors:
        return None, errors, notices

    month = _as_month(
        cell("billing_month"), sheet=sheet_name, row=row_number, field_name="billing_month", errors=errors
    )
    tariff = _as_enum(
        cell("tariff_structure"), TARIFF_STRUCTURE_TEXT, sheet=sheet_name, row=row_number,
        field_name="tariff_structure", errors=errors,
    )
    declared_source = _as_enum(
        cell("source_type"), BILL_SOURCE_TEXT, sheet=sheet_name, row=row_number,
        field_name="source_type", errors=errors,
    )

    numeric_fields = (
        "contract_capacity_kva",
        "billing_demand_kw",
        "energy_total_kwh",
        "energy_sharp_kwh",
        "energy_peak_kwh",
        "energy_flat_kwh",
        "energy_valley_kwh",
        "energy_offpeak_kwh",
        "energy_charge_yuan",
        "market_purchase_charge_yuan",
        "transmission_distribution_charge_yuan",
        "line_loss_charge_yuan",
        "system_operation_charge_yuan",
        "government_fund_charge_yuan",
        "basic_capacity_charge_yuan",
        "demand_charge_yuan",
        "power_factor_adjustment_yuan",
        "other_charge_yuan",
        "vat_yuan",
        "adjustment_charge_yuan",
        "bill_total_yuan",
    )
    values: dict[str, float | None] = {
        name: _as_float(cell(name), sheet=sheet_name, row=row_number, field_name=name, errors=errors)
        for name in numeric_fields
    }

    if errors:
        return None, errors, notices

    if period_start is None or period_end is None:  # pragma: no cover - 由上面的错误分支保证
        return None, [f"{_loc(sheet_name, row_number)}：账期缺失，无法生成账单"], notices

    meter_id = _as_text(cell("meter_id"))
    notes = _as_text(cell("notes"))

    # 来源诚实性（§2.1）：本行来自 Excel，来源就记 excel；表中的其他标注如实保留到备注
    if declared_source is not None and declared_source is not default_source:
        notices.append(
            f"{_loc(sheet_name, row_number, 'source_type')}：表中标注来源为『{declared_source.label}』，"
            f"但该行来自 Excel 导入，已按『{default_source.label}』记录，原标注保留在备注中"
        )
        notes = f"{notes or ''}（表中标注来源：{declared_source.label}）".strip()

    kwargs: dict[str, Any] = {
        "bill_id": make_bill_id(period_start, period_end, meter_id),        "project_id": project_id,
        "billing_period_start": period_start,
        "billing_period_end": period_end,
        "meter_id": meter_id,
        "customer_name": _as_text(cell("customer_name")),
        "voltage_level": _as_text(cell("voltage_level")),
        "tariff_structure": tariff if tariff is not None else TariffStructure.UNKNOWN,
        "source_type": default_source,
        "source_file_name": file_name,
        "source_row_number": row_number,
        "notes": notes,
        **values,
    }
    if month:
        kwargs["billing_month"] = month

    try:
        bill = ElectricityBill(**kwargs)
    except PydanticValidationError as exc:
        for err in exc.errors():
            location = err.get("loc", ("",))
            field_name = str(location[0]) if location else ""
            message = str(err.get("msg", "")).removeprefix("Value error, ")
            errors.append(f"{_loc(sheet_name, row_number, field_name)}：{message}")
        return None, errors, notices

    bill = apply_quality(bill, tolerance=tolerance)
    # 字段级硬问题（如负值）：作为行错误返回 → 该行"无效"，不参与导入
    for code_message in bill.quality_messages:
        if code_message.startswith("[B01]") or code_message.startswith("[B02]"):
            errors.append(f"{_loc(sheet_name, row_number)}：{code_message}")
    return bill, errors, notices


# --------------------------------------------------------------------------- #
# 预览
# --------------------------------------------------------------------------- #
def _row_is_example(row: dict[str, Any], marker_header: str | None) -> bool:
    if not marker_header:
        return False
    value = row.get(marker_header)
    if _is_blank(value):
        return False
    normalized = normalize_header(value)
    return "示例" in normalized or "example" in normalized


def preview_bill_import(
    path: str | Path,
    *,
    sheet: str | None = None,
    column_mapping: dict[str, str] | None = None,
    skip_example_rows: bool = True,
    existing_bills: Iterable[ElectricityBill] = (),
    project_id: str = "",
    tolerance: BillTolerance | None = None,
) -> BillImportPreview:
    """读取账单文件并生成**导入预览**（V2.1 §5.3、§5.5、§9.1）。

    :param sheet: 工作表名；``None`` = 模板默认的"月账单"（不存在时退回第一张表）
    :param column_mapping: 用户手工指定的列映射（``字段 → 表头``）
    :param skip_example_rows: 是否跳过标记为"示例行"的行（**默认跳过**，§5.3）
    :param existing_bills: 已存在的账单，用于检测"重复导入"
    :param project_id: 写入账单的 ``project_id``
    :raises ValidationError: 文件/工作表/必需列层面的致命错误（中文提示）
    """
    target = Path(path)
    existing_list = list(existing_bills)  # 允许传入生成器：只迭代一次
    sheets = list_sheets(target)
    if sheet is None:
        sheet_name = DEFAULT_SHEET if DEFAULT_SHEET in sheets else sheets[0]
    elif sheet in sheets:
        sheet_name = sheet
    else:
        raise ValidationError(
            f"工作表「{sheet}」不存在；文件「{target.name}」包含的工作表：{'、'.join(sheets)}",
            field="bill.import.sheet",
        )

    numbered_rows = read_table_from_sheet(target, sheet_name, with_row_numbers=True)
    numbered_rows = list(numbered_rows)  # type: ignore[arg-type]
    headers = list(numbered_rows[0][1].keys())
    mapping = resolve_bill_columns(headers, explicit=column_mapping)
    marker_header = next(
        (h for h in headers if normalize_header(h) == normalize_header(MARKER_COLUMN)), None
    )

    messages: list[str] = list(mapping.messages)
    if marker_header is None and skip_example_rows:
        messages.append(
            f"未找到『{MARKER_COLUMN}』列，无法识别示例行，已按全部数据行处理"
            "（模板自带示例行请务必删除或清空后再导入）"
        )
    ignored_headers = [
        h for h in mapping.unmapped_headers if marker_header is None or h != marker_header
    ]
    if ignored_headers:
        messages.append("未识别的列（已忽略）：" + "、".join(ignored_headers[:10]))

    preview = BillImportPreview(
        file_name=target.name,
        sheet_name=sheet_name,
        available_sheets=sheets,
        column_mapping=dict(mapping.mapping),
        unmapped_headers=list(mapping.unmapped_headers),
        existing_bill_ids=sorted({bill.bill_id for bill in existing_list}),
        messages=messages,
    )

    seen_keys: dict[str, str] = {}
    for existing in existing_list:
        seen_keys[existing.duplicate_key] = existing.bill_id

    example_skipped = 0
    for row_number, raw_row in numbered_rows:
        if skip_example_rows and _row_is_example(raw_row, marker_header):
            example_skipped += 1
            continue

        raw_text = {
            str(key): ("" if _is_blank(value) else str(value)) for key, value in raw_row.items()
        }
        bill, row_errors, notices = build_bill_from_row(
            raw_row,
            mapping,
            sheet_name=sheet_name,
            row_number=row_number,
            project_id=project_id,
            file_name=target.name,
            tolerance=tolerance,
        )

        row_messages: list[str] = [*row_errors, *notices]
        duplicate_of: str | None = None

        if bill is None or row_errors:
            status = BillQualityStatus.INVALID
            if not row_errors:
                row_messages.append(f"{_loc(sheet_name, row_number)}：该行无法解析为账单，已跳过")
        else:
            status = bill.quality_status
            key = bill.duplicate_key
            if key in seen_keys:
                duplicate_of = seen_keys[key]
                row_messages.append(
                    f"{_loc(sheet_name, row_number)}：疑似重复账单——与『{duplicate_of}』的"
                    "项目、账期、计量点组合相同，请选择跳过 / 替换 / 保留"
                )
            else:
                seen_keys[key] = f"文件第 {row_number} 行（{bill.bill_id}）"
            if bill.quality_messages:
                row_messages.extend(
                    f"{_loc(sheet_name, row_number)}：{message}" for message in bill.quality_messages
                )

        preview.total_rows += 1
        if duplicate_of:
            preview.duplicate_count += 1
        if status is BillQualityStatus.INVALID:
            preview.invalid_count += 1
        elif status is BillQualityStatus.WARNING:
            preview.warning_count += 1
        else:
            preview.valid_count += 1

        preview.rows.append(
            BillImportRow(
                row_number=row_number,
                status=status,
                bill=bill,
                duplicate_of=duplicate_of,
                messages=row_messages,
                raw=raw_text,
            )
        )

    preview.example_rows_skipped = example_skipped

    # 「分时电价」工作表：V2.1 不消费，有数据时明确提示（不静默忽略）
    if SHEET_TARIFF in sheets:
        try:
            tariff_rows = read_table_from_sheet(target, SHEET_TARIFF)
        except ValidationError:
            tariff_rows = []
        if tariff_rows:
            preview.messages.append(
                f"模板中『{SHEET_TARIFF}』工作表有 {len(tariff_rows)} 行数据；"
                "V2.1 不保存账单实际分时单价（该能力在 V2.3 的电价计划中实现，见规格书 §2.3、阶段 5），"
                "本次导入不会读取该表"
            )

    logger.info("账单导入预览：%s", preview.counts_text())
    return preview


# --------------------------------------------------------------------------- #
# 执行导入（跳过 / 替换 / 保留）
# --------------------------------------------------------------------------- #
def _unique_id(base_id: str, taken: set[str]) -> str:
    """为"两条都保留"生成不冲突的账单 ID（``X`` → ``X-2`` → ``X-3``…）。"""
    if base_id not in taken:
        return base_id
    index = 2
    while f"{base_id}-{index}" in taken:
        index += 1
    return f"{base_id}-{index}"


def apply_bill_import(
    preview: BillImportPreview, *, strategy: DuplicateStrategy = DuplicateStrategy.SKIP
) -> BillImportResult:
    """按策略把预览结果落实为"新增 / 替换 / 保留 / 跳过"清单（V2.1 §5.5）。

    纯函数：**不接触项目对象**，由 :class:`cenep.application.bill_service.BillService`
    负责把结果合并进项目并保存，保证"重复导入不重复写入"。
    """
    result = BillImportResult(strategy=strategy)
    # 已有账单编号先占位：KEEP_BOTH 时新账单必须换一个不冲突的编号，否则会与旧账单撞号
    taken: set[str] = set(preview.existing_bill_ids)

    for row in preview.rows:
        if not row.importable or row.bill is None:
            result.invalid_rows.append(row.row_number)
            continue

        bill = row.bill
        if row.duplicate_of is None:
            unique = bill.model_copy(update={"bill_id": _unique_id(bill.bill_id, taken)})
            taken.add(unique.bill_id)
            result.added.append(unique)
            continue

        if strategy is DuplicateStrategy.SKIP:
            result.skipped_ids.append(bill.bill_id)
        elif strategy is DuplicateStrategy.REPLACE:
            taken.add(bill.bill_id)
            result.replaced.append(bill)
        else:  # KEEP_BOTH
            unique = bill.model_copy(update={"bill_id": _unique_id(bill.bill_id, taken)})
            taken.add(unique.bill_id)
            result.kept_both.append(unique)

    result.messages.append(result.summary_text())
    logger.info("账单导入执行：%s", result.summary_text())
    return result


# --------------------------------------------------------------------------- #
# 模板生成（§5.3）
# --------------------------------------------------------------------------- #
def _example_row_values() -> list[dict[str, Any]]:
    """两行**自洽的**示例数据（标记为示例行，默认不导入）。

    刻意不出现任何"单价"：示例金额只体现数量关系，**不代表任何现行电价**
    （V2.1 §0.2、§4.1：不得预填示例电价当真实现行电价）。
    """
    first: dict[str, Any] = {
        "": EXAMPLE_MARKER,
        "billing_period_start": "2026-01-01",
        "billing_period_end": "2026-01-31",
        "billing_month": "2026-01",
        "meter_id": "M001",
        "customer_name": "示例企业（请替换为真实数据）",
        "voltage_level": "10kV",
        "tariff_structure": "单一制",
        "energy_total_kwh": 100000,
        "energy_sharp_kwh": 10000,
        "energy_peak_kwh": 30000,
        "energy_flat_kwh": 40000,
        "energy_valley_kwh": 20000,
        "energy_charge_yuan": 65000,
        "market_purchase_charge_yuan": 45000,
        "transmission_distribution_charge_yuan": 15000,
        "line_loss_charge_yuan": 1000,
        "system_operation_charge_yuan": 500,
        "government_fund_charge_yuan": 3500,
        "bill_total_yuan": 65000,
        "notes": "示例行：分时电量合计 = 总电量；电度分项合计 = 电度电费合计 = 账单总额",
    }
    second: dict[str, Any] = {
        "": EXAMPLE_MARKER,
        "billing_period_start": "2025-12-01",
        "billing_period_end": "2025-12-31",
        "billing_month": "2025-12",
        "meter_id": "M001",
        "customer_name": "示例企业（请替换为真实数据）",
        "voltage_level": "10kV",
        "tariff_structure": "两部制",
        "contract_capacity_kva": 630,
        "billing_demand_kw": 450,
        "energy_total_kwh": 80000,
        "energy_sharp_kwh": 8000,
        "energy_peak_kwh": 24000,
        "energy_flat_kwh": 32000,
        "energy_valley_kwh": 16000,
        "energy_charge_yuan": 52000,
        "market_purchase_charge_yuan": 36000,
        "transmission_distribution_charge_yuan": 12000,
        "line_loss_charge_yuan": 800,
        "system_operation_charge_yuan": 400,
        "government_fund_charge_yuan": 2800,
        "bill_total_yuan": 52000,
        "notes": "示例行：两部制账单；示例未提供基本电费分项（留空 = 账单未提供，不是 0）",
    }
    return [first, second]


def _sheet_rows(example: dict[str, Any]) -> list[Any]:
    return [example.get(column.field, "") for column in BILL_COLUMNS]


def _write_template(workbook_path: Path | None) -> Any:
    """构建模板工作簿；``workbook_path=None`` 时只返回内存工作簿。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover - 依赖缺失时给出可读提示
        raise ValidationError(
            "生成账单导入模板需要 openpyxl，请先安装该依赖", field="bill.template.openpyxl"
        ) from exc

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="4472C4")
    example_font = Font(italic=True, color="808080")
    wrap = Alignment(wrap_text=True, vertical="top")
    examples = _example_row_values()

    # ---------------- 填写说明 ----------------
    help_sheet = workbook.active
    help_sheet.title = SHEET_HELP
    help_sheet.append(["CENEP 电费账单导入模板 —— 填写说明"])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.append([])
    for line in (
        "1) 只需填写『月账单』工作表。",
        "2) 『分时电价』工作表在 V2.1 不参与导入（留待 V2.3 电价计划），请勿在此填写示例价格。",
        "3) 示例行在『数据标记』列写了『示例行』，导入时默认自动跳过；请填写真实数据并删除示例行。",
        "4) 单位统一：金额人民币元、电量 kWh、功率 kW、容量 kVA；单元格内不要写单位文字。",
        "5) 账单未提供的分项请留空（留空 = 未提供）；填 0 表示账单上确实是 0，两者含义不同。",
        "6) 本模板不预填、也不代表任何现行电价数值；实际单价一律以用户账单/合同为准（V2.1 §0.2、§4.1）。",
    ):
        help_sheet.append([line])
    help_sheet.append([])
    field_header_row = help_sheet.max_row + 1
    help_sheet.append(["字段", "单位", "是否必填", "填写说明"])
    for cell in help_sheet[field_header_row]:
        cell.font = header_font
        cell.fill = header_fill
    for column in BILL_COLUMNS:
        help_sheet.append(
            [column.header, column.unit, "必填" if column.required else "选填", column.description]
        )
    help_sheet.column_dimensions["A"].width = 30
    help_sheet.column_dimensions["B"].width = 12
    help_sheet.column_dimensions["C"].width = 12
    help_sheet.column_dimensions["D"].width = 70

    # ---------------- 月账单 ----------------
    bills_sheet = workbook.create_sheet(SHEET_BILLS)
    bills_sheet.append([column.header for column in BILL_COLUMNS])
    for cell in bills_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    bills_sheet.freeze_panes = "A2"
    for example in examples:
        bills_sheet.append(_sheet_rows(example))
    for row in bills_sheet.iter_rows(min_row=2, max_row=1 + len(examples)):
        for cell in row:
            cell.font = example_font
    for index, column in enumerate(BILL_COLUMNS, start=1):
        bills_sheet.column_dimensions[get_column_letter(index)].width = max(
            14, min(26, len(column.header) * 1.6)
        )

    # 枚举列下拉选项（openpyxl DataValidation，不依赖宏）
    enum_options = {
        "tariff_structure": ["单一制", "两部制", "未知"],
        "voltage_level": list(VOLTAGE_LEVEL_OPTIONS),
        "source_type": ["手动录入", "Excel导入", "估算"],
    }
    headers_in_order = [column.header for column in BILL_COLUMNS]
    last_row = 1 + len(examples) + 200
    for field_name, options in enum_options.items():
        header = next((c.header for c in BILL_COLUMNS if c.field == field_name), None)
        if header is None:
            continue
        letter = get_column_letter(headers_in_order.index(header) + 1)
        validation = DataValidation(
            type="list", formula1=f'"{",".join(options)}"', allow_blank=True, showErrorMessage=True
        )
        validation.errorTitle = "取值不在允许范围内"
        validation.error = "请从下拉列表中选择"
        bills_sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last_row}")

    # ---------------- 分时电价（只建表头，V2.1 不导入） ----------------
    tariff_sheet = workbook.create_sheet(SHEET_TARIFF)
    tariff_sheet.append([header for header, _ in TARIFF_SHEET_COLUMNS])
    for cell in tariff_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for index, (header, _description) in enumerate(TARIFF_SHEET_COLUMNS, start=1):
        tariff_sheet.column_dimensions[get_column_letter(index)].width = max(
            16, min(30, len(header) * 1.6)
        )
    tariff_sheet.freeze_panes = "A2"

    # ---------------- 数据字典 ----------------
    dict_sheet = workbook.create_sheet(SHEET_DICT)
    dict_sheet.append(["字段名（程序用）", "模板列名", "单位", "允许值 / 说明"])
    for cell in dict_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for column in BILL_COLUMNS:
        if column.field == "tariff_structure":
            allowed = "单一制 = single_part；两部制 = two_part；未知 = unknown"
        elif column.field == "source_type":
            allowed = "手动录入 = manual；Excel导入 = excel；估算 = estimated（导入时一律记为 excel）"
        elif column.field == "billing_month":
            allowed = "YYYY-MM，留空按账期起始日推导"
        else:
            allowed = column.description
        dict_sheet.append([column.field or "（辅助列）", column.header, column.unit, allowed])
    dict_sheet.append([])
    for field_name, header, unit, note in (
        ("energy_offpeak_kwh", "深谷电量(kWh)", "kWh", "深谷 ≠ 低谷：低谷是 energy_valley_kwh（V2.1 §2.1）"),
        ("billing_demand_kw", "账单计费需量(kW)", "kW", "账单计费需量，不是负荷曲线最大值（V2.1 §2.1）"),
        (
            "energy_charge_yuan",
            "电度电费合计(元)",
            "元",
            "与代理购电/输配/线损/系统运行/基金 5 个明细列是合计与明细的关系，不可同时计入账单总额",
        ),
    ):
        dict_sheet.append([field_name, header, unit, note])
    dict_sheet.append([])
    dict_sheet.append(["（工作表）", SHEET_TARIFF, "", "V2.1 不导入；列说明见下"])
    for header, description in TARIFF_SHEET_COLUMNS:
        dict_sheet.append([f"{SHEET_TARIFF}.{header}", header, "", description])
    for index, width in enumerate((34, 28, 10, 70), start=1):
        dict_sheet.column_dimensions[get_column_letter(index)].width = width

    if workbook_path is not None:
        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(workbook_path)
        logger.info("生成账单导入模板：%s", workbook_path)
    return workbook


def build_bill_template(target: str | Path) -> Path:
    """生成 ``CENEP_电费账单导入模板.xlsx``（V2.1 §5.3）。

    :param target: 目标文件路径；若传**已存在的目录**，则在目录下使用模板默认文件名。
    :return: 实际写入的路径
    """
    path = Path(target)
    if path.is_dir():
        path = path / TEMPLATE_FILE_NAME
    elif path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    _write_template(path)
    return path


def bill_template_bytes() -> bytes:
    """模板文件的字节内容，供界面"下载模板"直接保存（V2.1 §5.3、阶段 2）。"""
    workbook = _write_template(None)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
