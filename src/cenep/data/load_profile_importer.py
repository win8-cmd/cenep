"""高频负荷文件导入：列映射、时间间隔识别、质量检查与导入预览（V2.2 §6.1、§6.4、§2.2、§9.2）。

六步流程（与 V2.1 的账单导入保持一致的操作范式）
----------------------------------------------

.. code-block:: text

    ① 选文件   list_sheets()                       列出工作表
    ② 选工作表 read_table_from_sheet() / read_time_blocks()
    ③ 映射列   resolve_load_columns() + resolve_value_kind()
    ④ 预览     preview_load_import() / preview_block_import()  → LoadImportPreview
    ⑤ 校验     check_load_quality()（V01–V14 复用既有校验器 + 本阶段新增 V19–V24）
    ⑥ 确认导入 apply_load_import()                  → HighFrequencyLoadDataset

支持两种常见的交付形态
----------------------
1. **长表**（一行一个时间点）：``时间列 + 功率列(kW)`` 或 ``时间列 + 间隔电量列(kWh)``（§2.2）。
2. **宽表**（"时刻行 × 日期列"矩阵，国内工程资料里非常常见，东风本田真实资料即为此形态）：

   .. code-block:: text

        R1  日期序列号        B..AF = 当月各天
        R2  00:00  数值 数值 数值 …
        R3  00:15  数值 数值 数值 …
        …
        共 96 行 = 一天的 96 个 15 分钟时段

   :func:`read_time_blocks` 自动识别这类矩阵，并把连续的时刻行切成**数据块**；
   工作表里可以叠放多块（用电量 / 光伏发电量 / 消纳电量），每块单独选择导入。

本阶段新增的校验规则（编号顺延，见 ``DATA_IMPORT_SPEC.md`` 第 6 节"新增规则顺延"要求）
--------------------------------------------------------------------------------------
.. list-table::
   :header-rows: 1

   * - 编号
     - 判定
     - 级别
   * - V19
     - 时间间隔无法判定（不足 2 点、间隔杂乱）或识别到的间隔不在支持范围内且未显式确认
     - 拒载（中文错误）
   * - V20
     - 时间戳未落在间隔网格上（含系统性漂移、亚秒抖动）；已按声明的对齐策略归整并留痕
     - 告警
   * - V21
     - 时间戳携带时区／跨夏令时；已统一换算为 Asia/Shanghai（UTC+8，中国无夏令时）
     - 告警 / 提示
   * - V22
     - 数据年份与文件名／用户预期年份不一致（真实资料常见：文件名 2025、时间戳 2019）
     - 告警
   * - V23
     - 平年／闰年点数与日历不符（15 分钟 35040/35136、30 分钟 17520/17568、1 小时 8760/8784）
     - 告警
   * - V24
     - 数据覆盖率不足（默认 < 90%）；结果必须按"部分年度／代表日／年化估算"标注
     - 告警
   * - V25
     - 数据块自带的日期表头与工作表名月份不一致（真实资料中"复制上月表头未更新"很常见）；
     已按工作表名月份推导日期并留痕，可用 ``date_source="header"`` 强制按表头
     - 告警

单位与口径（§0.2、§3.2）
------------------------
* 电量 **kWh**、功率 **kW**；内部统一存"间隔电量 kWh"，功率是派生量；
* 源文件是功率还是电量**必须判定或显式声明**，判定不出时报中文错误，**绝不猜**；
* 时间语义统一为**左闭右开** ``[t, t + Δt)``，时间点为区间左端点；
* 时区固定 ``Asia/Shanghai``（UTC+8）。
"""

from __future__ import annotations

import logging
import math
import re
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from dataclasses import replace
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np
from pydantic import Field

from ..calculation.errors import ValidationError
from ..calculation.load_resample import resample_dataset
from ..calculation.timeseries_engine import build_time_axis, is_leap_year, points_per_year
from ..domain.base import _Model
from ..domain.enums import (
    LoadDataSourceType,
    LoadQualityStatus,
    LoadValueKind,
    Resolution,
)
from ..domain.load_data import (
    SUPPORTED_INTERVAL_MINUTES,
    TIMEZONE_DEFAULT,
    HighFrequencyLoadDataset,
    make_dataset_id,
)
from ..domain.timeseries import TimeSeriesPoint
from ..domain.timeseries_results import DataQualityIssue, DataQualityScore
from .bill_importer import ColumnMapping, _match_strength
from .importer import list_sheets, normalize_header, parse_timestamp, read_table_from_sheet
from .quality import score_quality
from .validator import _detect_unsorted, detect_missing, validate_series

logger = logging.getLogger(__name__)

__all__ = [
    "COVERAGE_WARNING_RATIO",
    "GRID_ALIGN_TOLERANCE_SECONDS",
    "INTERVAL_TOLERANCE_SECONDS",
    "LOAD_COLUMNS",
    "LOAD_COLUMN_ALIASES",
    "AlignmentResult",
    "IntervalDetection",
    "LoadColumn",
    "LoadImportPreview",
    "LoadImportRow",
    "TimeBlock",
    "align_timestamps",
    "apply_load_import",
    "build_points_from_block",
    "check_load_quality",
    "detect_interval",
    "detect_value_kind_from_header",
    "expected_year_from_file_name",
    "list_load_blocks",
    "load_block_catalog",
    "parse_load_timestamp",
    "preview_block_import",
    "preview_load_import",
    "read_time_blocks",
    "resolve_block_dates",
    "resolve_load_columns",
    "resolve_value_kind",
    "scan_timezone_issues",
    "sheet_month_hint",
]

# --------------------------------------------------------------------------- #
# 常量
# --------------------------------------------------------------------------- #
#: 相邻间隔与"主间隔"的允许偏差（秒）。真实资料的时间戳来自 Excel 浮点日期序列号，
#: 常见 0.005~0.5 s 的抖动，必须容差而不能判定为"不规则间隔"。
INTERVAL_TOLERANCE_SECONDS = 1.0

#: 绝对网格对齐容差（秒）：时间戳与"自当日 00:00 起 Δt 整数倍"的最大偏差。
GRID_ALIGN_TOLERANCE_SECONDS = 30.0

#: 行序重建要求"规则间隔"的最小占比
SEQUENTIAL_MIN_REGULAR_RATIO = 0.99

#: 覆盖率告警阈值（§2.2：覆盖不足不得冒充完整年度实测）
COVERAGE_WARNING_RATIO = 0.90

#: 预览默认返回的原始行数（完整数据仍保留在 ``LoadImportPreview.points``）
DEFAULT_PREVIEW_ROWS = 50

#: 一个宽表数据块内"一天应有的时段数"（15 分钟 96、30 分钟 48、1 小时 24、5 分钟 288）
_INTERVALS_PER_DAY: tuple[int, ...] = (24, 48, 96, 288)

#: 中国标准时间（UTC+8，1991 年后无夏令时；V2.2 §2.2 只支持 Asia/Shanghai）
CST = timezone(timedelta(hours=8))

#: 时间戳列别名（归一化后比较，见 :func:`cenep.data.importer.normalize_header`）
_TIMESTAMP_ALIASES: tuple[str, ...] = (
    "timestamp",
    "datetime",
    "time",
    "date",
    "时间",
    "时间戳",
    "时刻",
    "日期",
    "日期时间",
    "日期和时间",
    "采集时间",
    "记录时间",
    "数据时间",
)

#: 电量列别名（kWh）
_ENERGY_ALIASES: tuple[str, ...] = (
    "kwh",
    "电量",
    "用电量",
    "负荷电量",
    "间隔电量",
    "有功电量",
    "总加有功电能",
    "每小时用电量",
    "每日每小时用电量",
    "负荷",
    "用电",
    "消耗电量",
)

#: 功率列别名（kW）
_POWER_ALIASES: tuple[str, ...] = (
    "kw",
    "功率",
    "有功功率",
    "负荷功率",
    "平均功率",
    "瞬时功率",
    "需量",
    "power",
)


@dataclass(frozen=True)
class LoadColumn:
    """负荷模板列定义（表头、模型字段、单位、必填）。"""

    field: str
    header: str
    unit: str = ""
    required: bool = False
    description: str = ""
    aliases: tuple[str, ...] = dataclass_field(default_factory=tuple)

    @property
    def label(self) -> str:
        return self.header


#: 负荷模板的规范列（§2.2：至少支持"时间戳 + 功率 kW"与"时间戳 + 间隔电量 kWh"）
LOAD_COLUMNS: tuple[LoadColumn, ...] = (
    LoadColumn(
        "timestamp", "时间", "", True,
        "时间戳（区间左端点，左闭右开）。例：2025-01-01 00:15",
        _TIMESTAMP_ALIASES,
    ),
    LoadColumn(
        "energy_kwh", "间隔电量(kWh)", "kWh", False,
        "该时间间隔内的电量（kWh/间隔）。与功率列二选一",
        _ENERGY_ALIASES,
    ),
    LoadColumn(
        "power_kw", "间隔平均功率(kW)", "kW", False,
        "该时间间隔的平均功率（kW）。与电量列二选一；系统会按 Δt 换算成电量",
        _POWER_ALIASES,
    ),
)

#: 字段 → 别名（列映射用；顺序即优先级）
LOAD_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    column.field: (column.aliases or (column.field,)) for column in LOAD_COLUMNS
}

#: 模板列名 → :class:`TimeSeriesPoint` 字段名
_POINT_FIELD: dict[str, str] = {
    "energy_kwh": "load_kwh",
    "power_kw": "load_kwh",
    "timestamp": "timestamp",
}

_YEAR_IN_NAME = re.compile(r"(19|20)\d{2}")


# --------------------------------------------------------------------------- #
# ③ 列映射与数值口径
# --------------------------------------------------------------------------- #
def resolve_load_columns(
    headers: Iterable[Any],
    *,
    explicit: dict[str, str] | None = None,
) -> ColumnMapping:
    """把文件表头映射到规范列（``timestamp`` / ``energy_kwh`` / ``power_kw``）。

    规则与账单导入一致（复用 :func:`cenep.data.bill_importer._match_strength` 的打分实现，
    不复制第二套匹配逻辑）：**用户显式映射优先 > 精确匹配 > 最长别名包含匹配**。

    :raises ValidationError: 无表头、缺时间列、电量列与功率列都缺失（中文提示并列出实际表头）
    """
    header_list = [str(h) for h in headers if h is not None and str(h).strip()]
    if not header_list:
        raise ValidationError(
            "工作表没有可用表头，无法识别负荷数据的列；"
            "若文件是『时刻行 × 日期列』的宽表，系统会尝试按数据块导入",
            field="load.import.header",
        )

    mapping = ColumnMapping()
    used: set[str] = set()

    for field_name, header in (explicit or {}).items():
        if header not in header_list:
            raise ValidationError(
                f"指定的列「{header}」在文件中不存在；可用表头：{'、'.join(header_list[:10])}",
                field=f"load.import.{field_name}",
            )
        mapping.mapping[field_name] = header
        used.add(header)

    for field_name in ("timestamp", "energy_kwh", "power_kw"):
        if field_name in mapping.mapping:
            continue
        aliases = LOAD_COLUMN_ALIASES[field_name]
        scored = [
            (*_match_strength(normalize_header(header), aliases), header)
            for header in header_list
            if header not in used
        ]
        scored = [item for item in scored if item[0] > 0]
        if not scored:
            continue
        scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
        best = scored[0]
        mapping.mapping[field_name] = best[2]
        used.add(best[2])
        others = [item[2] for item in scored[1:] if item[0] == best[0] and item[1] == best[1]]
        if others:
            mapping.messages.append(
                f"列「{'、'.join(others)}」与「{best[2]}」都可能是"
                f"「{LOAD_COLUMNS[[c.field for c in LOAD_COLUMNS].index(field_name)].header}」，"
                f"已按「{best[2]}」识别；如需改判请在导入界面手工指定列映射"
            )

    mapping.unmapped_headers = [h for h in header_list if h not in used]

    if "timestamp" not in mapping.mapping:
        raise ValidationError(
            "负荷表缺少时间列；请指定一列时间戳（支持 时间/时间戳/时刻/日期和时间/datetime）"
            f"；实际表头：{'、'.join(header_list[:10])}",
            field="load.import.timestamp",
        )
    if "energy_kwh" not in mapping.mapping and "power_kw" not in mapping.mapping:
        raise ValidationError(
            "负荷表缺少数值列；请指定一列『间隔电量(kWh)』或『间隔平均功率(kW)』"
            f"；实际表头：{'、'.join(header_list[:10])}",
            field="load.import.value",
        )
    logger.debug("负荷列映射：%s", mapping.mapping)
    return mapping


def detect_value_kind_from_header(header: str | None) -> LoadValueKind | None:
    """按表头（或数据块名称）文字推断数值口径（kW / kWh）；判断不出返回 ``None``。

    判定顺序（先把"功率"判掉，避免 ``用电功率`` 这类同时含两类词的名称被误判）：

    1. 含 ``功率`` / ``需量``，或以 ``kw`` 结尾（``kwh`` 不算，它以 h 结尾）→ 功率 kW；
    2. 含 ``kwh`` / ``电量`` / ``电能`` / ``用电``，或含 ``度`` 且不含 ``温``（排除温度）→ 电量 kWh；
    3. 其余 → ``None``：调用方必须报中文错误，**不得猜测**（V2.2 §6.1）。
    """
    if header is None:
        return None
    norm = normalize_header(header)
    if not norm:
        return None
    if "功率" in norm or "需量" in norm or norm.endswith("kw"):
        return LoadValueKind.POWER_KW
    if (
        "kwh" in norm
        or "电量" in norm
        or "电能" in norm
        or "用电" in norm
        or ("度" in norm and "温" not in norm)
    ):
        return LoadValueKind.INTERVAL_ENERGY_KWH
    return None


def resolve_value_kind(
    header: str | None,
    declared: LoadValueKind | None = None,
    *,
    where: str = "",
) -> tuple[LoadValueKind, list[str]]:
    """确定数值列的口径（功率 kW 还是间隔电量 kWh），返回 ``(口径, 提示列表)``。

    * 用户显式声明 > 表头单位；两者都有且矛盾时报中文错误（不让用户"顺手点错"悄悄通过）；
    * 都判断不出时**报中文错误**，绝不默认成某一种（§9.2、§0.2）。
    """
    notices: list[str] = []
    from_header = detect_value_kind_from_header(header)

    if declared is not None:
        if from_header is not None and from_header is not declared:
            raise ValidationError(
                f"{where}表头「{header}」的单位看起来是「{from_header.label}」，"
                f"但界面选择的是「{declared.label}」；二者矛盾时不得继续导入，"
                f"请核对原始数据后重新选择（V2.2 §3.2：功率与电量的换算相差一个 Δt，"
                f"选错会造成 4 倍／0.25 倍的系统性误差）",
                field="load.import.value_kind",
            )
        return declared, notices

    if from_header is None:
        raise ValidationError(
            f"无法判定{where}数值列「{header or '（未指定）'}」的单位是"
            f"『间隔平均功率 kW』还是『间隔电量 kWh』；"
            f"请在导入界面显式指定数值口径（V2.2 §6.1 要求不得猜测口径）",
            field="load.import.value_kind",
        )
    notices.append(
        f"数值列「{header}」未显式指定口径，已按表头单位判定为「{from_header.label}」；"
        f"如与实际不符请显式指定"
    )
    return from_header, notices


# --------------------------------------------------------------------------- #
# 时间戳解析（时区 / 夏令时）
# --------------------------------------------------------------------------- #
def _round_to_second(value: datetime) -> datetime:
    """四舍五入到整秒（**不是**截断）。

    为什么必须四舍五入：真实资料的时间戳来自 Excel 浮点日期序列号，常见
    ``01:59:59.990``、``02:59:59.985`` 这类"差几毫秒不到整点"的值。
    若按截断处理会变成 ``01:59:59``，相邻间隔随之在 3599/3600 秒之间抖动，
    间隔识别与对齐都会被误判为"不规整"（实测：截断后主间隔占比 58.8%，
    四舍五入后 99.9%）。四舍五入既保留真实时点，又不改动超过 0.5 秒的信息。
    """
    base = value.replace(microsecond=0)
    return base + timedelta(seconds=1) if value.microsecond >= 500_000 else base


def parse_load_timestamp(value: Any) -> tuple[datetime | None, str | None]:
    """解析时间列，返回 ``(无时区的本地时间, 中文提示或 None)``。

    与 :func:`cenep.data.importer.parse_timestamp` 的差异**只在时区与亚秒处理**：

    * 带时区（``2025-01-01T00:00:00+09:00``、``...Z``）的时间戳统一换算为
      ``Asia/Shanghai``（UTC+8），并给出一条中文提示；
    * 亚秒部分**四舍五入到整秒**（见 :func:`_round_to_second`），
      避免 Excel 浮点序列号把"差几毫秒不到整点"的值截断成前一秒；
    * 其余格式（含 Excel 日期序列号）完全复用既有解析器，行为不变。
    """
    if isinstance(value, datetime) and value.tzinfo is not None:
        converted = _round_to_second(value.astimezone(CST).replace(tzinfo=None))
        offset = value.utcoffset()
        hours = offset.total_seconds() / 3600.0 if offset else 0.0
        note = (
            f"时间戳 {value.isoformat()} 带时区（UTC{hours:+.2f}），"
            f"已换算为 Asia/Shanghai（UTC+8）：{converted:%Y-%m-%d %H:%M:%S}"
        )
        return converted, note

    if isinstance(value, str) and value.strip():
        text = value.strip()
        if text.endswith("Z") or re.search(r"[+-]\d{2}:?\d{2}$", text):
            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            except ValueError:
                parsed = None
            if parsed is not None and parsed.tzinfo is not None:
                converted = _round_to_second(parsed.astimezone(CST).replace(tzinfo=None))
                return converted, (
                    f"时间戳 {text} 带时区，已换算为 Asia/Shanghai（UTC+8）："
                    f"{converted:%Y-%m-%d %H:%M}"
                )

    if isinstance(value, datetime):
        return _round_to_second(value.replace(tzinfo=None)), None

    stamp = parse_timestamp(value)
    if stamp is not None:
        return stamp.replace(tzinfo=None), None
    return None, None


def scan_timezone_issues(raw_values: Sequence[Any]) -> list[DataQualityIssue]:
    """V21 时区／夏令时检查（§2.2、§9.2）。

    判定（只报告，不阻断）：

    * 存在**多个不同 UTC 偏移** → 判定为跨夏令时（DST）切换；
    * 存在偏移不是 +08:00 的带时区时间戳 → 提示已换算；
    * 全部无时区 → 提示已按 ``Asia/Shanghai`` 解释（口径留痕）。
    """
    offsets: dict[float, int] = {}
    aware = 0
    for value in raw_values:
        if isinstance(value, datetime) and value.tzinfo is not None:
            offset = value.utcoffset()
            hours = round((offset.total_seconds() / 3600.0) if offset else 0.0, 3)
            offsets[hours] = offsets.get(hours, 0) + 1
            aware += 1
        elif isinstance(value, str) and value.strip():
            text = value.strip()
            if text.endswith("Z") or re.search(r"[+-]\d{2}:?\d{2}$", text):
                try:
                    parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
                except ValueError:
                    continue
                if parsed.tzinfo is not None:
                    offset = parsed.utcoffset()
                    hours = round((offset.total_seconds() / 3600.0) if offset else 0.0, 3)
                    offsets[hours] = offsets.get(hours, 0) + 1
                    aware += 1

    issues: list[DataQualityIssue] = []
    if not aware:
        if raw_values:
            issues.append(
                DataQualityIssue(
                    level="INFO",
                    category="source",
                    message=(
                        "[V21] 时间戳未携带时区信息，已统一按 Asia/Shanghai（UTC+8）解释；"
                        "中国自 1991 年后不实行夏令时，全年恒为 UTC+8（V2.2 §2.2）"
                    ),
                    count=len(raw_values),
                    samples=[],
                )
            )
        return issues

    if len(offsets) > 1:
        listed = "、".join(f"UTC{offset:+.2f}（{count} 点）" for offset, count in sorted(offsets.items()))
        issues.append(
            DataQualityIssue(
                level="WARNING",
                category="continuity",
                message=(
                    f"[V21] 检测到夏令时（DST）切换：数据含 {len(offsets)} 个不同 UTC 偏移"
                    f"（{listed}）。已全部换算为 Asia/Shanghai（UTC+8）；"
                    f"换算后请复核跨切换时刻的间隔是否仍为整数倍（§9.2 时区/DST 用例）"
                ),
                count=aware,
                samples=[],
            )
        )
    elif next(iter(offsets)) != 8.0:
        offset = next(iter(offsets))
        issues.append(
            DataQualityIssue(
                level="WARNING",
                category="continuity",
                message=(
                    f"[V21] 数据带时区（UTC{offset:+.2f}），已换算为 Asia/Shanghai（UTC+8）；"
                    f"请确认原始时区与生产时区一致（V2.2 §6.4：负荷与光伏必须同一时区）"
                ),
                count=aware,
                samples=[],
            )
        )
    else:
        issues.append(
            DataQualityIssue(
                level="INFO",
                category="source",
                message=(
                    f"[V21] 数据带时区 UTC+08:00（{aware} 点），与 Asia/Shanghai 一致，未做换算"
                ),
                count=aware,
                samples=[],
            )
        )
    return issues


# --------------------------------------------------------------------------- #
# ②' 宽表数据块（"时刻行 × 日期列"）
# --------------------------------------------------------------------------- #
@dataclass
class TimeBlock:
    """一个"时刻行 × 日期列"数据块（如某月 96 点用电量矩阵）。"""

    sheet_name: str
    label: str
    header_row: int
    first_data_row: int
    times: list[time]
    dates: list[date]
    values: list[list[float]]

    @property
    def row_count(self) -> int:
        return len(self.times)

    @property
    def day_count(self) -> int:
        return len(self.dates)

    @property
    def intervals_per_day(self) -> int:
        return self.row_count

    @property
    def interval_minutes(self) -> int:
        """按"一天多少行"推得的间隔（96 行 → 15 分钟）。"""
        if self.row_count <= 0:
            return 0
        return int(round(24 * 60 / self.row_count))

    @property
    def total_energy_kwh(self) -> float:
        total = 0.0
        for row in self.values:
            for value in row:
                if np.isfinite(value):
                    total += float(value)
        return total

    @property
    def cell_count(self) -> int:
        return self.row_count * self.day_count

    def describe(self) -> str:
        return (
            f"「{self.label}」（工作表「{self.sheet_name}」第 {self.first_data_row}–"
            f"{self.first_data_row + self.row_count - 1} 行）：{self.day_count} 天 × "
            f"{self.row_count} 个 {self.interval_minutes} 分钟时段 = {self.cell_count} 点"
        )


def _as_time_of_day(value: Any) -> time | None:
    """列 A 的单元格 → 当日时刻。"""
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%H:%M:%S", "%H:%M", "%H时%M分"):
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return None


def _as_date(value: Any) -> date | None:
    """表头单元格 → 日期（支持 ``date``/``datetime``、Excel 序列号、常见日期字符串）。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        if 20_000.0 <= serial <= 60_000.0:  # 1954-09 ~ 2064 年，足够覆盖工程资料
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date()
        return None
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y-%m", "%Y/%m"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _sheet_rows(path: str | Path, sheet: str | None) -> tuple[str, list[tuple]]:
    """读整张工作表为二维元组（保留物理行号，表头行号也在此确定）。"""
    from openpyxl import load_workbook  # 局部导入：与 importer.py 一致，避免模块级依赖

    target = Path(path)
    if not target.exists():
        raise ValidationError(f"文件不存在：{target}", field="load.import.path")
    suffix = target.suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        raise ValidationError(
            f"宽表（时刻行 × 日期列）导入只支持 .xlsx；"
            f"文件「{target.name}」的扩展名为「{suffix}」，请另存为 .xlsx",
            field="load.import.suffix",
        )
    try:
        workbook = load_workbook(target, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - 统一翻译为中文提示
        raise ValidationError(
            f"文件「{target.name}」不是有效的 Excel 工作簿（可能已损坏或被占用）：{exc}",
            field="load.import.corrupt",
        ) from exc
    try:
        names = list(workbook.sheetnames)
        if sheet is None:
            sheet_name = names[0]
        elif sheet in names:
            sheet_name = sheet
        else:
            raise ValidationError(
                f"工作表「{sheet}」不存在；文件「{target.name}」包含的工作表：{'、'.join(names)}",
                field="load.import.sheet",
            )
        rows = [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()
    return sheet_name, rows


def read_time_blocks(path: str | Path, sheet: str | None = None) -> list[TimeBlock]:
    """识别工作表中的"时刻行 × 日期列"数据块（可为多块，如用电量/光伏/消纳）。

    识别规则（文档即契约）：

    1. 从第 1 列找出**连续的时刻行**（值形如 ``00:00``、``00:15``…）；
    2. 该组长度必须是 ``24/48/96/288`` 之一（即"一天"的整时段数）；
    3. 组内第一行的**上一行**必须能在第 2 列起找到至少一个日期（或 Excel 日期序列号）；
    4. 块名取"表头行的上一行第 1 列"的文字（如 ``光伏发电量``），没有则用工作表名。

    :returns: 数据块列表（按出现顺序）；没有识别到任何块时返回空列表（由调用方决定报错文案）
    """
    sheet_name, rows = _sheet_rows(path, sheet)
    blocks: list[TimeBlock] = []
    index = 0
    while index < len(rows):
        if _as_time_of_day(rows[index][0] if rows[index] else None) is None:
            index += 1
            continue
        start = index
        while index < len(rows) and _as_time_of_day(rows[index][0] if rows[index] else None) is not None:
            index += 1
        end = index  # [start, end)
        length = end - start
        if length not in _INTERVALS_PER_DAY:
            logger.debug(
                "跳过时刻行组（%s 第 %d–%d 行）：长度 %d 不是一天整时段数 %s",
                sheet_name, start + 1, end, length, _INTERVALS_PER_DAY,
            )
            continue
        if start - 1 < 0:
            continue
        header = rows[start - 1]
        dates: list[date] = []
        for cell in header[1:]:
            parsed = _as_date(cell)
            if parsed is None:
                break  # 日期列必须连续，遇到空列即认为该块到此为止
            dates.append(parsed)
        if not dates:
            continue

        label = ""
        if start - 2 >= 0 and rows[start - 2]:
            text = rows[start - 2][0]
            if isinstance(text, str) and text.strip():
                label = text.strip()
        if not label:
            label = f"{sheet_name}#{len(blocks) + 1}"

        values: list[list[float]] = []
        for row in rows[start:end]:
            line: list[float] = []
            for offset in range(len(dates)):
                raw = row[1 + offset] if 1 + offset < len(row) else None
                if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                    line.append(float(raw))
                else:
                    line.append(float("nan"))
            values.append(line)

        times = [
            _as_time_of_day(rows[start + offset][0]) or time(0, 0)
            for offset in range(length)
        ]
        blocks.append(
            TimeBlock(
                sheet_name=sheet_name,
                label=label,
                header_row=start,  # 1-based 行号（rows[start-1] 是表头）
                first_data_row=start + 1,
                times=times,
                dates=dates,
                values=values,
            )
        )
        logger.debug("识别数据块：%s", blocks[-1].describe())
    return blocks


def load_block_catalog(path: str | Path) -> dict[str, list[str]]:
    """列出文件里所有工作表的全部数据块名（供导入界面第 ②/③ 步选择）。"""
    catalog: dict[str, list[str]] = {}
    for sheet in list_sheets(path):
        try:
            blocks = read_time_blocks(path, sheet)
        except ValidationError:
            continue
        if blocks:
            catalog[sheet] = [block.label for block in blocks]
    return catalog


def list_load_blocks(path: str | Path, sheet: str | None = None) -> list[TimeBlock]:
    """同 :func:`read_time_blocks` 的别名，语义更明确（供界面调用）。"""
    return read_time_blocks(path, sheet)


def build_points_from_block(
    block: TimeBlock,
    *,
    value_field: str = "load_kwh",
    dates: Sequence[date] | None = None,
) -> list[TimeSeriesPoint]:
    """把数据块展开成**逐时间点**序列（列优先：先按时间，后按日期）。

    :param value_field: ``TimeSeriesPoint`` 的字段名；负荷用 ``load_kwh``（默认），
        光伏用 ``pv_generation_kwh``。这样 PV 曲线也复用同一套读取与质量检查，
        不需要第二套导入器（V2.2 §6.4：光伏与负荷必须同一时间轴）。
    :param dates: 覆盖块自带的日期表头（用于"表头未更新"的真实资料，见
        :func:`resolve_block_dates`）；``None`` 时使用 ``block.dates``。
    """
    if value_field not in ("load_kwh", "pv_generation_kwh"):
        raise ValidationError(
            f"不支持把数据块写入字段「{value_field}」，仅支持负荷（load_kwh）与"
            f"光伏（pv_generation_kwh）",
            field="load.import.value_field",
        )
    day_list = list(dates) if dates is not None else list(block.dates)
    if len(day_list) != block.day_count:
        raise ValidationError(
            f"数据集日期列数（{len(day_list)}）与数据块日期列数（{block.day_count}）不一致，"
            f"无法按列位置对齐（V2.2 §6.4）",
            field="load.import.dates",
        )
    points: list[TimeSeriesPoint] = []
    for day_index, day in enumerate(day_list):
        for row_index, moment in enumerate(block.times):
            stamp = datetime(day.year, day.month, day.day, moment.hour, moment.minute)
            points.append(
                TimeSeriesPoint(
                    timestamp=stamp,
                    **{value_field: float(block.values[row_index][day_index])},
                )
            )
    points.sort(key=lambda point: point.timestamp)
    return points


#: 工作表名里的月份：``1月用电数据``、``2025年3月负荷`` 都能提取
_SHEET_MONTH = re.compile(r"(\d{1,2})\s*月")


def sheet_month_hint(sheet_name: str) -> int | None:
    """从工作表名提取月份（``11月用电数据`` → 11）；提取不到返回 ``None``。"""
    match = _SHEET_MONTH.search(sheet_name or "")
    if not match:
        return None
    month = int(match.group(1))
    return month if 1 <= month <= 12 else None


def trim_trailing_padding(
    block: TimeBlock,
    *,
    month: int | None = None,
    year: int | None = None,
) -> tuple[TimeBlock, int]:
    """裁掉数据块**尾部整列无正值**的模板遗留列，返回 ``(裁剪后的块, 裁剪列数)``。

    真实资料里极常见：把"31 天"的模板整块复制到月表，2 月多出 3 列、11 月多出 1 列，
    多出的列要么全空、要么全 0（东风本田 2 月光伏块的第 29–31 列即全 0）。
    不裁掉它们会凭空造出 2 月 29–31 日这种不存在的日期，或让相邻月份时间戳重复。

    裁剪规则（**只裁尾部、只裁无正值列，绝不改数值**）：

    * 已知 ``month`` / ``year`` 且列数**多于**该月天数 → 多出的列必须整列无正值，
      裁剪到"该月天数"为止；若多出的列里仍有正值，抛中文错误要求人工核对
      （不猜哪几列属于本月）；
    * 已知月份且列数不多于该月天数 → **不裁**（该月最后一天真值为 0 属正常数据）；
    * 月份未知 → 尾部所有无正值列一律裁掉，并在调用方留痕。
    """
    count = block.day_count
    if count == 0:
        return block, 0
    has_positive = [
        any(
            math.isfinite(float(row[index])) and float(row[index]) > 0.0
            for row in block.values
        )
        for index in range(count)
    ]

    if month is not None and year is not None:
        expected = _days_in_month(year, month)
        if count > expected:
            extra = count - expected
            trailing_positive = [
                index for index in range(expected, count) if has_positive[index]
            ]
            if trailing_positive:
                raise ValidationError(
                    f"数据块「{block.label}」（工作表「{block.sheet_name}」）有 {count} 列，"
                    f"而 {year} 年 {month} 月只有 {expected} 天；多出的 {extra} 列"
                    f"（第 {'、'.join(str(i + 1) for i in trailing_positive)} 列）仍有正值，"
                    f"无法判断哪些列属于本月，请人工核对工作表后删除多余列再导入（V2.2 §6.4）",
                    field="load.import.dates",
                )
            keep = expected
            return (
                replace(
                    block,
                    dates=block.dates[:keep],
                    values=[row[:keep] for row in block.values],
                ),
                extra,
            )
        return block, 0

    drop = 0
    while drop < count and not has_positive[count - 1 - drop]:
        drop += 1
    if drop == 0:
        return block, 0
    keep = count - drop
    return (
        replace(block, dates=block.dates[:keep], values=[row[:keep] for row in block.values]),
        drop,
    )


def resolve_block_dates(
    block: TimeBlock,
    *,
    date_source: str = "auto",
    expected_year: int | None = None,
) -> tuple[list[date], DataQualityIssue | None]:
    """确定数据块每一天的日期，返回 ``(日期列表, 中文告警或 None)``。

    真实工程资料里极常见的一种缺陷：**光伏/消纳块的日期表头是从上月工作表整块复制来的，
    没有随月份更新**（东风本田 15 分钟资料：2–12 月光伏块表头全部仍写 1 月；
    11–12 月用电块表头仍写 10 月），若直接采信表头，多个月份的曲线会被压进同一个月
    并产生大量重复时间戳。因此本函数提供三种口径，**默认 auto 且逐块留痕**，
    绝不静默改日期：

    ==================  ==================================================================
    ``header``          完全采信数据块自带的日期表头（最保守）
    ``sheet_month``     按**工作表名里的月份** + 列序号推导日期（1 日、2 日…）；
                        年份取表头日期的年份，其次 ``expected_year``
    ``auto``（默认）     表头日期的月份与"工作表名月份"一致 → 用表头；不一致 → 用工作表名
                        推导并给出 WARNING 级问题（[V25]，列出不一致的日期范围）
    ==================  ==================================================================

    :raises ValidationError: ``date_source`` 取值非法，或工作表名月份与列数不匹配（如 2 月有 31 列）
    """
    if date_source not in ("auto", "header", "sheet_month"):
        raise ValidationError(
            f"不支持的日期口径「{date_source}」；可用 auto / header / sheet_month",
            field="load.import.date_source",
        )
    if date_source == "header":
        return list(block.dates), None

    month = sheet_month_hint(block.sheet_name)
    header_years = {d.year for d in block.dates}
    year = (
        next(iter(header_years))
        if len(header_years) == 1
        else (expected_year or (block.dates[0].year if block.dates else None))
    )
    if month is None or year is None or not block.dates:
        if date_source == "sheet_month":
            raise ValidationError(
                f"工作表「{block.sheet_name}」无法按工作表名推导日期（月份={'未知' if month is None else month}、"
                f"年份={'未知' if year is None else year}）；请改用 date_source='header' 或修正工作表名",
                field="load.import.date_source",
            )
        return list(block.dates), None

    try:
        derived = [
            date(year, month, day) for day in range(1, block.day_count + 1)
        ]
    except ValueError as exc:
        raise ValidationError(
            f"工作表「{block.sheet_name}」按 {year} 年 {month} 月推导日期失败："
            f"该月只有 {_days_in_month(year, month)} 天，而数据块有 {block.day_count} 列；"
            f"请检查工作表名月份与数据列数是否匹配（V2.2 §6.4）",
            field="load.import.dates",
        ) from exc

    if date_source == "sheet_month":
        return derived, None

    # auto：表头月份与工作表名月份一致 → 采信表头
    if all(d.month == month for d in block.dates):
        return list(block.dates), None

    issue = DataQualityIssue(
        level="WARNING",
        category="continuity",
        message=(
            f"[V25] 数据块「{block.label}」（工作表「{block.sheet_name}」）自带的日期表头为 "
            f"{block.dates[0]:%Y-%m-%d}–{block.dates[-1]:%Y-%m-%d}（{block.dates[0].month} 月），"
            f"与工作表名的 {month} 月不一致，疑似复制上月表头未更新；"
            f"已按工作表名月份推导日期 {derived[0]:%Y-%m-%d}–{derived[-1]:%Y-%m-%d}。"
            f"如确需按表头日期导入，请把日期口径改为 header（V2.2 §6.4）"
        ),
        count=block.day_count,
        samples=[
            block.dates[0].strftime("%Y-%m-%d"),
            derived[0].strftime("%Y-%m-%d"),
        ],
    )
    return derived, issue


def _days_in_month(year: int, month: int) -> int:
    """某年某月的天数。"""
    if month == 12:
        return 31
    return (date(year, month + 1, 1) - date(year, month, 1)).days


# --------------------------------------------------------------------------- #
# 时间间隔识别与时间戳对齐
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class IntervalDetection:
    """时间间隔识别结果（V2.2 §6.1、§9.2）。

    ``regular_ratio`` = 与主间隔之差**在容差内**的相邻间隔占比（1 − 异常占比）。
    之所以不用"众数占比"判定规整性：真实资料的时间戳来自 Excel 浮点日期序列号，
    逐个时间戳四舍五入到整秒后，相邻间隔会在 ``3600`` 与 ``3601`` 秒之间跳，
    众数占比可能只有 ~59%，但**全部**都在 ±1 秒容差内，属于规整间隔。
    """

    interval_minutes: int | None
    resolution: Resolution | None
    regular: bool
    median_seconds: float
    dominant_seconds: float
    dominant_ratio: float
    regular_ratio: float
    irregular_count: int
    irregular_samples: tuple[str, ...]
    note: str

    def describe(self) -> str:
        if self.interval_minutes is None:
            return f"未能判定时间间隔：{self.note}"
        return (
            f"识别到 {self.interval_minutes} 分钟间隔（主间隔 {self.dominant_seconds:.0f} 秒，"
            f"规整占比 {self.regular_ratio:.2%}，中位数 {self.median_seconds:.3f} 秒）"
            + (f"；{self.irregular_count} 处间隔异常" if self.irregular_count else "")
        )


def detect_interval(
    timestamps: Sequence[datetime],
    *,
    tolerance_seconds: float = INTERVAL_TOLERANCE_SECONDS,
) -> IntervalDetection:
    """识别时间间隔（众数 + 容差），**不猜**：判不出来时 ``interval_minutes`` 为 ``None``。

    判定口径：

    * 相邻时间差取**众数**（按秒取整计数）作为"主间隔"，避免个别异常间隔带偏中位数；
    * 与主间隔之差 **≤ ``tolerance_seconds``（默认 1 秒）** 的都算"规整间隔"；
    * 规整间隔占比 ≥ 99%、且主间隔是**整分钟** → 认定为规则间隔；
    * 主间隔不是整分钟（如 45 秒）或规整占比不足 → 判为"间隔不规整"，返回 ``None``。
    """
    if len(timestamps) < 2:
        return IntervalDetection(
            None, None, False, 0.0, 0.0, 0.0, 0.0, 0,
            (),
            f"时间点不足 2 个（实际 {len(timestamps)} 个），无法判定间隔",
        )

    stamps = np.asarray(list(timestamps), dtype="datetime64[ns]")
    deltas = np.diff(stamps).astype("timedelta64[ns]").astype(np.int64) / 1e9
    positive = deltas[deltas > 0.0]
    if positive.size == 0:
        return IntervalDetection(
            None, None, False, 0.0, 0.0, 0.0, 0.0, len(deltas), (),
            "所有相邻时间差均为 0 或负值（时间列可能全部重复或倒序）",
        )

    median = float(np.median(positive))
    rounded = np.round(deltas).astype(np.int64)
    values, counts = np.unique(rounded, return_counts=True)
    order = np.lexsort((values, -counts))
    dominant_seconds = float(values[order[0]])
    dominant_ratio = float(counts[order[0]] / deltas.size)

    # 重复（Δ=0）与倒序（Δ<0）由 V02/V05 单独报告，**不计入**规整性比率：
    # 否则一个"只多了一行重复"的正常文件会被判成"间隔无法判定"而拒载（§5.4 各司其职）。
    non_positive = int(np.sum(deltas <= 0.0))
    # 缺口（Δ = k×主间隔，k ≥ 2 的整数倍，容差按倍数放大）属于"缺时间点"，
    # 由 V03/V04 报告，同样不计入"间隔不规整"。否则少一行数据就会导致整份文件无法判定间隔。
    with np.errstate(divide="ignore", invalid="ignore"):
        multiples = np.round(deltas / dominant_seconds) if dominant_seconds > 0 else np.zeros_like(deltas)
    gaps = (
        (deltas > 0.0)
        & (multiples >= 2.0)
        & (np.abs(deltas - multiples * dominant_seconds) <= tolerance_seconds * multiples)
    )
    irregular = np.flatnonzero(
        (deltas > 0.0)
        & (~gaps)
        & (np.abs(deltas - dominant_seconds) > tolerance_seconds)
    )
    gap_count = int(np.sum(gaps))
    samples: list[str] = []
    for idx in irregular[:5]:
        samples.append(
            f"{np.datetime_as_string(stamps[int(idx)], unit='m').replace('T', ' ')} → "
            f"{np.datetime_as_string(stamps[int(idx) + 1], unit='m').replace('T', ' ')}"
            f"（{deltas[int(idx)]:.3f} 秒）"
        )
    irregular_count = int(irregular.size)
    regular_ratio = 1.0 - irregular_count / positive.size

    minute_ok = abs(dominant_seconds - round(dominant_seconds / 60.0) * 60.0) <= tolerance_seconds
    regular = regular_ratio >= SEQUENTIAL_MIN_REGULAR_RATIO and minute_ok and dominant_seconds > 0.0
    interval_minutes = int(round(dominant_seconds / 60.0)) if regular else None
    extra_note = ""
    if non_positive:
        extra_note += f"；另有 {non_positive} 处重复/倒序时间戳（由 V02/V05 报告）"
    if gap_count:
        extra_note += f"；{gap_count} 处缺口（缺失时间点，由 V03/V04 报告）"
    note = (
        f"主间隔 {dominant_seconds:.0f} 秒（规整占比 {regular_ratio:.2%}）{extra_note}"
        if regular
        else (
            f"主间隔 {dominant_seconds:.3f} 秒不是整分钟，无法判定规整间隔{extra_note}"
            if not minute_ok
            else (
                f"主间隔 {dominant_seconds:.0f} 秒，但只有 {regular_ratio:.2%} 的相邻间隔"
                f"落在 ±{tolerance_seconds:g} 秒容差内（{irregular_count} 处异常），"
                f"间隔不规整{extra_note}"
            )
        )
    )
    return IntervalDetection(
        interval_minutes=interval_minutes,
        resolution=(
            Resolution.from_interval_minutes(interval_minutes)
            if interval_minutes is not None
            else None
        ),
        regular=regular,
        median_seconds=median,
        dominant_seconds=dominant_seconds,
        dominant_ratio=dominant_ratio,
        regular_ratio=regular_ratio,
        irregular_count=irregular_count,
        irregular_samples=tuple(samples),
        note=note,
    )


def require_interval(
    detection: IntervalDetection,
    *,
    declared_minutes: int | None = None,
    allow_nonstandard: bool = False,
    where: str = "",
) -> int:
    """把识别结果落实为可用间隔；无法判定或超出支持范围时报**中文错误**（V19）。

    :param declared_minutes: 用户在界面显式指定的间隔；与识别结果明显冲突时报错
    :param allow_nonstandard: 是否允许 15/30/60 之外的**规则**间隔（如 5 分钟）
    :raises ValidationError: 无法判定、非整分钟、与界面指定冲突、非标准间隔未确认
    """
    prefix = f"{where}：" if where else ""
    if declared_minutes is not None:
        declared = int(declared_minutes)
        if declared <= 0:
            raise ValidationError(
                f"{prefix}指定的时间间隔必须是正整数分钟，实际为 {declared_minutes}",
                field="load.import.interval",
            )
        if detection.interval_minutes is not None and detection.interval_minutes != declared:
            # 识别时长与指定时长的差异超过容差即视为冲突
            if abs(detection.dominant_seconds - declared * 60.0) > INTERVAL_TOLERANCE_SECONDS:
                raise ValidationError(
                    f"{prefix}界面指定的间隔为 {declared} 分钟，但时间列实际"
                    f"{detection.describe()}；二者矛盾，请核对原始数据或改回自动识别"
                    f"（V2.2 §9.2：不得未经确认按猜测的间隔重采样）",
                    field="load.import.interval",
                )
        return declared

    if detection.interval_minutes is None:
        raise ValidationError(
            f"{prefix}无法判定时间间隔：{detection.note}。"
            f"系统不会猜测间隔（V2.2 §6.1）；请检查时间列是否成规律，"
            f"或在导入界面显式指定时间间隔（15/30/60 分钟）后重试",
            field="load.import.interval",
        )
    minutes = detection.interval_minutes
    if minutes not in SUPPORTED_INTERVAL_MINUTES and not allow_nonstandard:
        raise ValidationError(
            f"{prefix}识别到 {minutes} 分钟间隔，超出 V2.2 一等支持的"
            f"{'/'.join(str(m) for m in SUPPORTED_INTERVAL_MINUTES)} 分钟范围。"
            f"如确需按该间隔导入，请在导入界面确认『允许非标准间隔』；"
            f"否则请先把数据聚合到 15/30/60 分钟（V2.2 §6.1、§6.4）",
            field="load.import.interval",
        )
    return minutes


@dataclass(frozen=True)
class AlignmentResult:
    """时间戳对齐结果（V2.2 §6.4：对齐规则必须明确并留痕）。"""

    strategy: str
    timestamps: tuple[datetime, ...]
    max_offset_seconds: float
    aligned_count: int
    notes: tuple[str, ...] = ()

    def describe(self) -> str:
        labels = {
            "grid": "绝对网格对齐（自当日 00:00 起 Δt 的整数倍）",
            "sequential": "行序重建（按行号以 Δt 递推，抗漂移）",
            "raw": "未对齐（保留原始时间戳）",
        }
        return (
            f"{labels.get(self.strategy, self.strategy)}：{self.aligned_count} 点，"
            f"最大调整 {self.max_offset_seconds:.3f} 秒"
        )


def align_timestamps(
    timestamps: Sequence[datetime],
    interval_minutes: int,
    *,
    strategy: str = "auto",
    grid_tolerance_seconds: float = GRID_ALIGN_TOLERANCE_SECONDS,
) -> AlignmentResult:
    """把时间戳归整到规整网格，返回对齐结果（含最大调整量与中文留痕）。

    三种策略：

    ==============  ==================================================================
    ``grid``        每个时间戳对齐到**绝对网格**（当日 00:00 起 Δt 的整数倍）。
                    要求所有点偏移 ≤ ``grid_tolerance_seconds``；否则该策略不适用。
    ``sequential``  按**行序**重建 ``t_first + i × Δt``。适用于"时间戳由公式生成、
                    逐行累计漂移"的真实资料（东风本田逐时用电量文件即如此：
                    相邻间隔 3600.495 秒，累计偏移接近 1 小时）。要求行序单调、
                    且非规整间隔占比 ≤ 1%。
    ``auto``        先试 ``grid``，不满足再试 ``sequential``；都不满足则返回
                    ``strategy="raw"`` 且 **不做任何改动**，由调用方报错（V20）。
    ==============  ==================================================================

    :raises ValidationError: ``interval_minutes`` 非正
    """
    if interval_minutes <= 0:
        raise ValidationError(
            f"对齐时间戳需要正的间隔分钟数，实际为 {interval_minutes}",
            field="load.import.align",
        )
    stamps = list(timestamps)
    if not stamps:
        return AlignmentResult("raw", (), 0.0, 0, ("数据为空，无需对齐",))
    step = timedelta(minutes=interval_minutes)
    if len(stamps) == 1:
        return AlignmentResult("grid", (stamps[0],), 0.0, 1, ())

    def _grid() -> AlignmentResult:
        base = datetime(stamps[0].year, stamps[0].month, stamps[0].day)
        out: list[datetime] = []
        max_offset = 0.0
        for stamp in stamps:
            index = round((stamp - base).total_seconds() / (interval_minutes * 60.0))
            snapped = base + step * index
            offset = abs((stamp - snapped).total_seconds())
            max_offset = max(max_offset, offset)
            out.append(snapped)
        return AlignmentResult("grid", tuple(out), max_offset, len(out))

    def _sequential() -> AlignmentResult:
        first = stamps[0]
        out = [first + step * index for index in range(len(stamps))]
        max_offset = max(abs((stamp - snapped).total_seconds()) for stamp, snapped in zip(stamps, out))
        return AlignmentResult("sequential", tuple(out), max_offset, len(out))

    grid_result = _grid()
    sequential_result = _sequential()

    if strategy == "grid":
        return grid_result
    if strategy == "sequential":
        return sequential_result
    if strategy != "auto":
        raise ValidationError(
            f"不支持的时间戳对齐策略「{strategy}」；可用 auto / grid / sequential",
            field="load.import.align",
        )

    if grid_result.max_offset_seconds <= grid_tolerance_seconds:
        return AlignmentResult(
            grid_result.strategy,
            grid_result.timestamps,
            grid_result.max_offset_seconds,
            grid_result.aligned_count,
            (f"时间戳与绝对网格最大偏差 {grid_result.max_offset_seconds:.3f} 秒，已按网格归整",),
        )

    # grid 不适用时，只有"行序规则"才允许重建
    detection = detect_interval(stamps)
    regular_ok = (
        detection.regular
        and detection.interval_minutes == interval_minutes
        and detection.regular_ratio >= SEQUENTIAL_MIN_REGULAR_RATIO
    )
    monotonic = all(a < b for a, b in zip(stamps, stamps[1:]))
    if regular_ok and monotonic:
        return AlignmentResult(
            sequential_result.strategy,
            sequential_result.timestamps,
            sequential_result.max_offset_seconds,
            sequential_result.aligned_count,
            (
                f"时间戳与绝对网格最大偏差 {grid_result.max_offset_seconds:.1f} 秒"
                f"（超过 {grid_tolerance_seconds:g} 秒容差），已改用**行序重建**："
                f"以首点 {stamps[0]:%Y-%m-%d %H:%M} 为基准按 {interval_minutes} 分钟递推。"
                f"原始时间戳存在系统性漂移，报告必须标注该处理（V2.2 §6.4）",
            ),
        )

    return AlignmentResult(
        "raw",
        tuple(stamps),
        grid_result.max_offset_seconds,
        len(stamps),
        (
            f"时间戳既未落在绝对网格上（最大偏差 {grid_result.max_offset_seconds:.1f} 秒），"
            f"也无法按行序重建（主间隔不规整或时间列非单调），未做对齐",
        ),
    )


# --------------------------------------------------------------------------- #
# ⑤ 质量检查（复用 V01–V14，新增 V19–V24）
# --------------------------------------------------------------------------- #
def expected_year_from_file_name(file_name: str) -> int | None:
    """从文件名提取预期年份（如 ``2025年每天的用电量.xlsx`` → 2025）。"""
    matches = _YEAR_IN_NAME.findall(file_name or "")
    if not matches:
        return None
    found = re.findall(r"(?:19|20)\d{2}", file_name)
    return int(found[0]) if found else None


def _dominant_year(points: Sequence[TimeSeriesPoint]) -> int | None:
    years: dict[int, int] = {}
    for point in points:
        years[point.timestamp.year] = years.get(point.timestamp.year, 0) + 1
    if not years:
        return None
    return max(years.items(), key=lambda item: item[1])[0]


def check_load_quality(
    points: Sequence[TimeSeriesPoint],
    *,
    interval_minutes: int,
    kind: str = "load",
    expected_year: int | None = None,
    file_name: str = "",
    alignment: AlignmentResult | None = None,
    timezone_issues: Sequence[DataQualityIssue] = (),
    coverage_warning_ratio: float = COVERAGE_WARNING_RATIO,
    annual_reference_kwh: float | None = None,
    capacity_kwp: float | None = None,
    extra_issues: Sequence[DataQualityIssue] = (),
) -> list[DataQualityIssue]:
    """负荷曲线质量检查（§2.2、§9.2）：复用既有 V01–V14，并追加 V19–V25。

    返回顺序固定（同一文件重复导入的报错条数与顺序完全一致，§5.4）：

    ``V19（间隔） → V20（对齐） → V21（时区） → V22（年份） → V23（平闰年点数）
    → V01…V14（既有校验器） → V24（覆盖率） → V25（读取阶段问题，如日期表头未更新）``
    """
    issues: list[DataQualityIssue] = []
    resolution = Resolution.from_interval_minutes(interval_minutes)

    if not points:
        issues.append(
            DataQualityIssue(
                level="ERROR",
                category="completeness",
                message="[V19] 没有可用的数据点：时间列或数值列未能解析出任何一行",
                count=0,
                samples=[],
            )
        )
        return issues
    if resolution is None:
        issues.append(
            DataQualityIssue(
                level="ERROR",
                category="continuity",
                message=(
                    f"[V19] 间隔 {interval_minutes} 分钟没有对应的内部分辨率，"
                    f"V2.2 支持 15/30/60 分钟（以及日、月）"
                ),
                count=1,
                samples=[],
            )
        )
        return issues

    # --- V20 对齐/漂移 ---
    if alignment is not None:
        if alignment.strategy == "raw":
            issues.append(
                DataQualityIssue(
                    level="ERROR",
                    category="continuity",
                    message=(
                        f"[V20] 时间戳无法归整到 {interval_minutes} 分钟网格："
                        f"{'；'.join(alignment.notes) or '未说明原因'}。"
                        f"请修正原始时间列后重试（V2.2 §6.1：不得未经确认自动重采样）"
                    ),
                    count=len(alignment.timestamps),
                    samples=[],
                )
            )
        elif alignment.max_offset_seconds > 1.0:
            detail = "；".join(alignment.notes) if alignment.notes else alignment.describe()
            issues.append(
                DataQualityIssue(
                    level="WARNING",
                    category="continuity",
                    message=(
                        f"[V20] {len(alignment.timestamps)} 个时间戳与规整网格不一致，"
                        f"最大偏差 {alignment.max_offset_seconds:.3f} 秒：{detail}"
                    ),
                    count=len(alignment.timestamps),
                    samples=[],
                )
            )

    # --- V21 时区 / 夏令时 ---
    issues.extend(timezone_issues)

    # --- V22 年份一致性 ---
    data_year = _dominant_year(points)
    if expected_year is not None and data_year is not None and data_year != expected_year:
        issues.append(
            DataQualityIssue(
                level="WARNING",
                category="source",
                message=(
                    f"[V22] 数据年份为 {data_year} 年，与文件名/预期年份 {expected_year} 年不一致"
                    + (f"（文件「{file_name}」）" if file_name else "")
                    + "。系统不会自动改年；请确认数据实际年份后再用于年度消纳分析"
                    "（真实资料中常见『文件名写 2025、时间戳为 2019』）"
                ),
                count=len(points),
                samples=[points[0].timestamp.strftime("%Y-%m-%d %H:%M")],
            )
        )

    # --- V23 平年/闰年点数 ---
    if data_year is not None and len({p.timestamp.year for p in points}) == 1:
        expected_points = points_per_year(data_year, resolution)
        leap = is_leap_year(data_year)
        other = points_per_year(data_year + (1 if not leap else -1), resolution)
        if len(points) != expected_points and len(points) == other:
            issues.append(
                DataQualityIssue(
                    level="WARNING",
                    category="completeness",
                    message=(
                        f"[V23] {data_year} 年（{'闰年' if leap else '平年'}）"
                        f"{resolution.label}数据的应有 {expected_points} 点，实得 {len(points)} 点"
                        f"（等于{'闰年' if not leap else '平年'}点数）；"
                        f"平闰年点数不得硬裁剪或混用（V2.2 §2.2）"
                    ),
                    count=abs(expected_points - len(points)),
                    samples=[],
                )
            )

    # --- V01…V14：复用既有校验器（时间轴 + 数值 + 交叉）---
    axis = build_time_axis(data_year or expected_year or datetime.now().year, resolution)
    issues.extend(
        validate_series(
            list(points),
            axis,
            kind=kind,
            capacity_kwp=capacity_kwp,
            annual_reference_kwh=annual_reference_kwh,
        )
    )

    # --- V24 覆盖率 ---
    missing_indices, _ = detect_missing(list(points), axis)
    expected_count = len(axis.timestamps)
    coverage = (expected_count - len(missing_indices)) / expected_count if expected_count else 0.0
    if coverage < coverage_warning_ratio:
        issues.append(
            DataQualityIssue(
                level="WARNING",
                category="completeness",
                message=(
                    f"[V24] {axis.year} 年 {resolution.label}数据覆盖率仅 {coverage:.2%}"
                    f"（阈值 {coverage_warning_ratio:.0%}，缺失 {len(missing_indices)} 点）。"
                    f"该曲线只能作为『部分年度／代表日／年化估算』使用，"
                    f"不得冒充完整年度实测（V2.2 §2.2、§6.4）"
                ),
                count=len(missing_indices),
                samples=[
                    axis.timestamps[i].strftime("%Y-%m-%d %H:%M")
                    for i in missing_indices[:5]
                ],
            )
        )

    logger.info(
        "负荷质量检查完成：kind=%s 点数=%d 间隔=%d 分钟 问题=%d（ERROR %d / WARNING %d / INFO %d）",
        kind,
        len(points),
        interval_minutes,
        len(issues),
        sum(issue.level == "ERROR" for issue in issues),
        sum(issue.level == "WARNING" for issue in issues),
        sum(issue.level == "INFO" for issue in issues),
    )

    # --- V25 读取阶段问题（如日期表头未更新）---
    issues.extend(extra_issues)
    return issues


def status_from_issues(issues: Sequence[DataQualityIssue]) -> LoadQualityStatus:
    """问题清单 → 质量状态（任一 ERROR 即无效；判定不由得分决定）。"""
    if any(issue.level == "ERROR" for issue in issues):
        return LoadQualityStatus.INVALID
    if any(issue.level == "WARNING" for issue in issues):
        return LoadQualityStatus.WARNING
    return LoadQualityStatus.VALID


# --------------------------------------------------------------------------- #
# ④ 导入预览模型
# --------------------------------------------------------------------------- #
class LoadImportRow(_Model):
    """导入预览中的一行原始数据（与源文件行号一致，便于用户定位）。"""

    row_number: int = Field(default=0, ge=0, description="源文件行号（宽表为展开后的序号）")
    timestamp_text: str = ""
    value_text: str = ""
    status: LoadQualityStatus = LoadQualityStatus.VALID
    messages: list[str] = Field(default_factory=list)


class LoadImportPreview(_Model):
    """负荷导入预览（六步流程的第 ④/⑤ 步产物，供界面展示与用户确认）。"""

    file_name: str = ""
    sheet_name: str = ""
    available_sheets: list[str] = Field(default_factory=list)
    layout: str = Field(default="table", description="table = 长表；time_block = 时刻行×日期列")
    block_label: str = ""
    available_blocks: list[str] = Field(default_factory=list)

    column_mapping: dict[str, str] = Field(default_factory=dict)
    unmapped_headers: list[str] = Field(default_factory=list)
    timestamp_column: str = ""
    value_column: str = ""
    value_kind: LoadValueKind = LoadValueKind.INTERVAL_ENERGY_KWH

    interval_minutes: int = 0
    resolution: Resolution | None = None
    interval_note: str = ""
    alignment_strategy: str = ""
    alignment_note: str = ""
    max_offset_seconds: float = 0.0
    expected_year: int | None = None
    data_year: int | None = None

    total_rows: int = 0
    valid_count: int = 0
    warning_count: int = 0
    invalid_count: int = 0
    duplicate_intervals: int = 0
    missing_intervals: int = 0
    irregular_intervals: int = 0
    outlier_count: int = 0

    coverage_ratio: float = Field(default=0.0, ge=0.0, le=1.0)
    point_count: int = 0
    annual_energy_kwh: float = 0.0
    peak_power_kw: float = 0.0

    quality_status: LoadQualityStatus = LoadQualityStatus.VALID
    quality_issues: list[DataQualityIssue] = Field(default_factory=list)
    quality_score: DataQualityScore | None = None
    messages: list[str] = Field(default_factory=list)

    points: list[TimeSeriesPoint] = Field(
        default_factory=list, description="已解析、已对齐的时间点（第 ⑥ 步直接使用）"
    )
    rows: list[LoadImportRow] = Field(default_factory=list)
    source_type: LoadDataSourceType = LoadDataSourceType.HIGH_FREQUENCY_IMPORT

    @property
    def importable(self) -> bool:
        """是否可确认导入（无效数据默认不可导入）。"""
        return self.quality_status is not LoadQualityStatus.INVALID and self.point_count > 0

    def counts_text(self) -> str:
        """一行中文统计（§5.3 对账单的要求同样适用于负荷）。"""
        return (
            f"{'宽表' if self.layout == 'time_block' else '长表'}「{self.block_label or self.sheet_name}」"
            f"识别到 {self.total_rows} 行 / {self.point_count} 个时间点（间隔 "
            f"{self.interval_minutes or '未判定'} 分钟）：有效 {self.valid_count}、"
            f"警告 {self.warning_count}、无效 {self.invalid_count}；"
            f"覆盖率 {self.coverage_ratio:.2%}，质量「{self.quality_status.label}」"
        )

    def issues_text(self) -> str:
        """全部质量问题的中文拼接（供界面/日志直接显示）。"""
        return "\n".join(issue.message for issue in self.quality_issues)


# --------------------------------------------------------------------------- #
# ④/⑤ 预览构建
# --------------------------------------------------------------------------- #
def _build_preview(
    *,
    file_name: str,
    sheet_name: str,
    available_sheets: Sequence[str],
    layout: str,
    block_label: str,
    available_blocks: Sequence[str],
    column_mapping: dict[str, str],
    unmapped_headers: Sequence[str],
    timestamp_column: str,
    value_column: str,
    value_kind: LoadValueKind,
    points: list[TimeSeriesPoint],
    raw_timestamps: Sequence[Any],
    total_rows: int,
    expected_year: int | None,
    alignment: AlignmentResult,
    interval_minutes: int,
    messages: list[str],
    rows: list[LoadImportRow],
    coverage_warning_ratio: float,
    annual_reference_kwh: float | None = None,
    extra_issues: Sequence[DataQualityIssue] = (),
    source_type: LoadDataSourceType = LoadDataSourceType.HIGH_FREQUENCY_IMPORT,
) -> LoadImportPreview:
    """组装预览（统一走这里，保证长表与宽表的质量口径完全一致）。

    :param extra_issues: 读取阶段已发现的问题（如 [V25] 日期表头未更新），
        统一并入质量清单，保证"同一份文件 → 同一份问题清单"（§5.4）。
    """
    timezone_issues = scan_timezone_issues(raw_timestamps)
    issues = check_load_quality(
        points,
        interval_minutes=interval_minutes,
        expected_year=expected_year,
        file_name=file_name,
        alignment=alignment,
        timezone_issues=timezone_issues,
        coverage_warning_ratio=coverage_warning_ratio,
        annual_reference_kwh=annual_reference_kwh,
        extra_issues=extra_issues,
    )

    resolution = Resolution.from_interval_minutes(interval_minutes)
    axis = build_time_axis(
        _dominant_year(points) or expected_year or datetime.now().year,
        resolution or Resolution.HOURLY,
    )
    missing_indices, _ = detect_missing(points, axis)
    expected_count = len(axis.timestamps)
    coverage = (expected_count - len(missing_indices)) / expected_count if expected_count else 0.0

    duplicates = next((i for i in issues if i.message.startswith("[V02]")), None)
    irregular = next((i for i in issues if i.message.startswith("[V04]")), None)
    outliers = [i for i in issues if i.category == "outlier"]

    status = status_from_issues(issues)
    score = score_quality(
        points,
        axis,
        issues=issues,
        source_type=source_type.parameter_source,
        kind="load",
    )
    energies = [float(p.load_kwh) for p in points if np.isfinite(p.load_kwh)]
    delta_hours = interval_minutes / 60.0
    peak = (max(energies) / delta_hours) if energies and delta_hours > 0 else 0.0

    preview = LoadImportPreview(
        file_name=file_name,
        sheet_name=sheet_name,
        available_sheets=list(available_sheets),
        layout=layout,
        block_label=block_label,
        available_blocks=list(available_blocks),
        column_mapping=dict(column_mapping),
        unmapped_headers=list(unmapped_headers),
        timestamp_column=timestamp_column,
        value_column=value_column,
        value_kind=value_kind,
        interval_minutes=interval_minutes,
        resolution=resolution,
        interval_note=alignment.describe(),
        alignment_strategy=alignment.strategy,
        alignment_note="；".join(alignment.notes),
        max_offset_seconds=float(alignment.max_offset_seconds),
        expected_year=expected_year,
        data_year=_dominant_year(points),
        total_rows=total_rows,
        valid_count=sum(1 for row in rows if row.status is LoadQualityStatus.VALID),
        warning_count=sum(1 for row in rows if row.status is LoadQualityStatus.WARNING),
        invalid_count=sum(1 for row in rows if row.status is LoadQualityStatus.INVALID),
        duplicate_intervals=int(duplicates.count) if duplicates else 0,
        missing_intervals=len(missing_indices),
        irregular_intervals=int(irregular.count) if irregular else 0,
        outlier_count=sum(int(issue.count) for issue in outliers),
        coverage_ratio=coverage,
        point_count=len(points),
        annual_energy_kwh=float(sum(energies)),
        peak_power_kw=float(peak),
        quality_status=status,
        quality_issues=issues,
        quality_score=score,
        messages=list(messages),
        points=points,
        rows=rows,
    )
    logger.info("负荷导入预览：%s", preview.counts_text())
    return preview


def preview_load_import(
    path: str | Path,
    *,
    sheet: str | None = None,
    header_row: int = 1,
    column_mapping: dict[str, str] | None = None,
    value_kind: LoadValueKind | None = None,
    interval_minutes: int | None = None,
    allow_nonstandard_interval: bool = False,
    expected_year: int | None = None,
    alignment_strategy: str = "auto",
    date_source: str = "auto",
    coverage_warning_ratio: float = COVERAGE_WARNING_RATIO,
    max_preview_rows: int = DEFAULT_PREVIEW_ROWS,
    annual_reference_kwh: float | None = None,
    block_index: int = 0,
) -> LoadImportPreview:
    """读取负荷文件并生成导入预览（长表优先，宽表自动回退）。

    :param sheet: 工作表名；``None`` = 第一张工作表
    :param header_row: 长表的表头所在行（真实资料常在第 2 行）
    :param column_mapping: 用户手工指定的列映射（``字段 → 表头``）
    :param value_kind: 数值口径（kW / kWh）；``None`` 时按表头单位判定，判不出报中文错误
    :param interval_minutes: 用户显式指定的间隔；``None`` 时自动识别
    :param allow_nonstandard_interval: 是否允许 15/30/60 之外的规则间隔
    :param expected_year: 预期数据年份；``None`` 时从文件名推断（V22 年份一致性检查）
    :param alignment_strategy: ``auto`` / ``grid`` / ``sequential``
    :param date_source: 仅宽表生效的日期口径 ``auto`` / ``header`` / ``sheet_month``
    :param block_index: 自动回退到宽表时使用第几个数据块（默认第 1 个）
    :raises ValidationError: 文件/工作表/列/间隔层面的致命问题（全部中文提示）
    """
    target = Path(path)
    sheets = list_sheets(target)
    if sheet is None:
        sheet_name = sheets[0]
    elif sheet in sheets:
        sheet_name = sheet
    else:
        raise ValidationError(
            f"工作表「{sheet}」不存在；文件「{target.name}」包含的工作表：{'、'.join(sheets)}",
            field="load.import.sheet",
        )
    if expected_year is None:
        expected_year = expected_year_from_file_name(target.name)

    # 先按"长表"读；读不出（宽表没有可用表头）或表头不像负荷表时，回退到宽表数据块
    # 注意：read_table_from_sheet 对"表头行全空/只有第一张表"等情况会抛中文错误，
    # 这些情况同样可能是宽表，因此统一降级到数据块识别，由数据块识别给出最终结论。
    try:
        numbered = list(
            read_table_from_sheet(target, sheet_name, header_row=header_row, with_row_numbers=True)
        )
    except ValidationError:
        numbered = []

    if numbered:
        headers = list(numbered[0][1].keys())
        if _looks_like_load_table(headers):
            return _preview_from_table(
                target,
                sheet_name,
                sheets,
                numbered,
                column_mapping=column_mapping,
                value_kind=value_kind,
                interval_minutes=interval_minutes,
                allow_nonstandard_interval=allow_nonstandard_interval,
                expected_year=expected_year,
                alignment_strategy=alignment_strategy,
                coverage_warning_ratio=coverage_warning_ratio,
                max_preview_rows=max_preview_rows,
                annual_reference_kwh=annual_reference_kwh,
            )

    blocks = read_time_blocks(target, sheet_name)
    if not blocks:
        raise ValidationError(
            f"工作表「{sheet_name}」既不是『时间列 + 数值列』的长表（缺少时间列或数值列），"
            f"也没有识别到『时刻行 × 日期列』的数据块。"
            f"可用列名示例：时间/时间戳/时刻、负荷/用电量(kWh)、功率(kW)；"
            f"宽表需第 1 列为时刻（00:00、00:15…）且表头行含日期（V2.2 §2.2）",
            field="load.import.layout",
        )
    return preview_block_import(
        target,
        sheet=sheet_name,
        block_index=block_index,
        value_field="load_kwh",
        value_kind=value_kind,
        interval_minutes=interval_minutes,
        allow_nonstandard_interval=allow_nonstandard_interval,
        expected_year=expected_year,
        alignment_strategy=alignment_strategy,
        date_source=date_source,
        coverage_warning_ratio=coverage_warning_ratio,
        max_preview_rows=max_preview_rows,
        available_sheets=sheets,
    )


def _looks_like_load_table(headers: Sequence[Any]) -> bool:
    """表头是否像负荷长表（**至少**能认出时间列）。

    这里只要求"有时间列"：数值列的名称可能千奇百怪（``1MWp光伏出力``、``总有功电能``…），
    能否映射出数值列交给 :func:`resolve_load_columns` 给出可定位的中文错误，
    或由用户在界面显式指定列映射；而"宽表"的特征是**第 1 列就是时刻、表头行没有列名**，
    因此只要认得出时间列就按长表处理。
    """
    normalized = [normalize_header(h) for h in headers if h is not None and str(h).strip()]
    if not normalized:
        return False
    return any(any(alias in norm for alias in _TIMESTAMP_ALIASES) for norm in normalized)


def _preview_from_table(
    target: Path,
    sheet_name: str,
    sheets: Sequence[str],
    numbered: Sequence[tuple[int, dict[str, Any]]],
    *,
    column_mapping: dict[str, str] | None,
    value_kind: LoadValueKind | None,
    interval_minutes: int | None,
    allow_nonstandard_interval: bool,
    expected_year: int | None,
    alignment_strategy: str,
    coverage_warning_ratio: float,
    max_preview_rows: int,
    annual_reference_kwh: float | None,
) -> LoadImportPreview:
    """长表路径：``时间列 + 功率列／电量列``。"""
    headers = list(numbered[0][1].keys())
    mapping = resolve_load_columns(headers, explicit=column_mapping)
    timestamp_column = mapping.mapping["timestamp"]
    value_column = (
        mapping.mapping.get("energy_kwh") or mapping.mapping.get("power_kw") or ""
    )
    kind, notices = resolve_value_kind(value_column, value_kind, where="负荷表")
    messages: list[str] = list(mapping.messages) + notices
    if mapping.unmapped_headers:
        messages.append("未识别的列（已忽略）：" + "、".join(mapping.unmapped_headers[:10]))

    # 单次解析：行号 → (时间戳, 数值)；解析失败的行单独计数并给出可定位的中文问题
    raw_values: list[Any] = []
    rows: list[LoadImportRow] = []
    parsed: list[tuple[int, datetime, float]] = []
    invalid_rows = 0
    bad_value_rows = 0
    for row_number, raw_row in numbered:
        raw_stamp = raw_row.get(timestamp_column)
        raw_value = raw_row.get(value_column)
        raw_values.append(raw_stamp)
        stamp, tz_note = parse_load_timestamp(raw_stamp)
        row_messages: list[str] = []
        if tz_note:
            row_messages.append(tz_note)

        number = float("nan")
        if raw_value is not None and str(raw_value).strip():
            try:
                number = float(raw_value)
            except (TypeError, ValueError):
                bad_value_rows += 1
                row_messages.append(
                    f"工作表「{sheet_name}」第 {row_number} 行：数值「{raw_value}」不是数字，"
                    f"已按缺失处理（V11）；缺失不得按 0 计算（V2.2 §6.4）"
                )
        else:
            bad_value_rows += 1
            row_messages.append(
                f"工作表「{sheet_name}」第 {row_number} 行：数值列为空，已按缺失处理（V11）"
            )

        if stamp is None:
            invalid_rows += 1
            status = LoadQualityStatus.INVALID
            row_messages.append(
                f"工作表「{sheet_name}」第 {row_number} 行：时间「{raw_stamp}」无法解析，"
                f"请使用 YYYY-MM-DD HH:MM 或 Excel 日期格式（V01）"
            )
        else:
            status = (
                LoadQualityStatus.WARNING if not np.isfinite(number) else LoadQualityStatus.VALID
            )
            parsed.append((row_number, stamp, number))

        if len(rows) < max_preview_rows:
            rows.append(
                LoadImportRow(
                    row_number=row_number,
                    timestamp_text="" if raw_stamp is None else str(raw_stamp),
                    value_text="" if raw_value is None else str(raw_value),
                    status=status,
                    messages=row_messages,
                )
            )

    if not parsed:
        raise ValidationError(
            f"文件「{target.name}」工作表「{sheet_name}」的时间列「{timestamp_column}」"
            f"没有任何一行能解析成时间（无效行 {invalid_rows} 行），请检查时间格式（V01）",
            field="load.import.timestamp",
        )
    # 乱序必须在**排序前**记录：排序之后 V05 就再也看不到了（§9.2 要求报告乱序）
    out_of_order_issues: list[DataQualityIssue] = _detect_unsorted(
        [
            TimeSeriesPoint(timestamp=stamp, load_kwh=0.0)
            for _, stamp, _ in parsed
        ]
    )
    parsed.sort(key=lambda item: item[1])

    detection = detect_interval([item[1] for item in parsed])
    minutes = require_interval(
        detection,
        declared_minutes=interval_minutes,
        allow_nonstandard=allow_nonstandard_interval,
        where=f"工作表「{sheet_name}」",
    )
    alignment = align_timestamps(
        [item[1] for item in parsed], minutes, strategy=alignment_strategy
    )
    if alignment.strategy == "raw":
        raise ValidationError(
            f"工作表「{sheet_name}」的时间戳无法归整到 {minutes} 分钟网格："
            f"{'；'.join(alignment.notes)}。请修正原始时间列，或先把数据补齐到规整间隔"
            f"（V2.2 §6.1：不得未经确认自动重采样）",
            field="load.import.align",
        )

    delta_hours = minutes / 60.0
    points: list[TimeSeriesPoint] = []
    for (_, _, raw_number), stamp in zip(parsed, alignment.timestamps):
        energy = (
            raw_number
            if kind is LoadValueKind.INTERVAL_ENERGY_KWH
            else raw_number * delta_hours
        )
        points.append(TimeSeriesPoint(timestamp=stamp, load_kwh=float(energy)))

    if kind is LoadValueKind.POWER_KW:
        messages.append(
            f"数值列「{value_column}」按间隔平均功率 kW 处理：已按 E = P × Δt"
            f"（Δt = {delta_hours:g} h）换算为间隔电量；**未**对功率直接累加（V2.2 §3.2）"
        )
    else:
        messages.append(
            f"数值列「{value_column}」按间隔电量 kWh 处理：直接累加，**未**再乘 Δt（V2.2 §3.2）"
        )
    if detection.irregular_count:
        messages.append(
            f"时间列存在 {detection.irregular_count} 处非规整间隔，"
            f"示例：{'；'.join(detection.irregular_samples)}"
        )
    if invalid_rows:
        messages.append(f"有 {invalid_rows} 行时间无法解析，已从导入结果中排除（V01）")
    if bad_value_rows:
        messages.append(f"有 {bad_value_rows} 行数值为空或非数字，已按缺失处理（V11）")

    return _build_preview(
        file_name=target.name,
        sheet_name=sheet_name,
        available_sheets=sheets,
        layout="table",
        block_label="",
        available_blocks=[],
        column_mapping=mapping.mapping,
        unmapped_headers=mapping.unmapped_headers,
        timestamp_column=timestamp_column,
        value_column=value_column,
        value_kind=kind,
        points=points,
        raw_timestamps=raw_values,
        total_rows=len(numbered),
        expected_year=expected_year,
        alignment=alignment,
        interval_minutes=minutes,
        messages=messages,
        rows=rows,
        coverage_warning_ratio=coverage_warning_ratio,
        annual_reference_kwh=annual_reference_kwh,
        extra_issues=out_of_order_issues,
    )


def preview_block_import(
    path: str | Path,
    *,
    sheet: str | None = None,
    block_index: int = 0,
    all_sheets: bool | None = None,
    value_field: str = "load_kwh",
    value_kind: LoadValueKind | None = None,
    interval_minutes: int | None = None,
    allow_nonstandard_interval: bool = False,
    expected_year: int | None = None,
    alignment_strategy: str = "auto",
    date_source: str = "auto",
    coverage_warning_ratio: float = COVERAGE_WARNING_RATIO,
    max_preview_rows: int = DEFAULT_PREVIEW_ROWS,
    available_sheets: Sequence[str] | None = None,
) -> LoadImportPreview:
    """宽表路径：从"时刻行 × 日期列"数据块生成预览。

    :param block_index: 每张工作表取第几个数据块（0 = 第 1 块）
    :param all_sheets: ``True`` 时把**所有**含该序号数据块的工作表纵向拼接
        （真实资料 = 12 张月表，必须拼接才是全年曲线）；``None`` 时自动判断：
        文件里超过 1 张表含有该序号的数据块就拼接。
    :param value_field: 写入 ``TimeSeriesPoint`` 的字段名（默认负荷）
    :param date_source: 日期口径 ``auto`` / ``header`` / ``sheet_month``，见
        :func:`resolve_block_dates`（默认 auto：表头月份与工作表名月份不一致时
        按工作表名推导并逐块告警 [V25]，绝不静默改日期）
    """
    target = Path(path)
    sheets = list(available_sheets) if available_sheets is not None else list_sheets(target)
    if expected_year is None:
        expected_year = expected_year_from_file_name(target.name)

    if all_sheets is None:
        candidate_count = 0
        for name in sheets:
            try:
                if len(read_time_blocks(target, name)) > block_index:
                    candidate_count += 1
            except ValidationError:
                continue
        all_sheets = candidate_count > 1

    targets = sheets if all_sheets else [sheet or sheets[0]]
    blocks: list[TimeBlock] = []
    catalog: list[str] = []
    for name in targets:
        try:
            sheet_blocks = read_time_blocks(target, name)
        except ValidationError:
            continue
        catalog.extend(f"{name}｜{block.label}" for block in sheet_blocks)
        if len(sheet_blocks) > block_index:
            blocks.append(sheet_blocks[block_index])

    if not blocks:
        raise ValidationError(
            f"文件「{target.name}」没有找到第 {block_index + 1} 个『时刻行 × 日期列』数据块；"
            f"请确认工作表第 1 列是时刻（00:00、00:15…），且表头行为日期"
            f"（可用数据块：{'、'.join(catalog) if catalog else '无'}）",
            field="load.import.block",
        )

    interval_from_blocks = blocks[0].interval_minutes
    for block in blocks[1:]:
        if block.interval_minutes != interval_from_blocks:
            raise ValidationError(
                f"待拼接的数据块间隔不一致：{blocks[0].describe()} 为 "
                f"{interval_from_blocks} 分钟，而 {block.describe()} 为 {block.interval_minutes} 分钟；"
                f"不同间隔的数据不得直接拼接（V2.2 §6.4）",
                field="load.import.block",
            )

    label_kind = (
        value_kind
        if value_kind is not None
        else detect_value_kind_from_header(blocks[0].label)
    )
    messages: list[str] = []
    if value_kind is None:
        if label_kind is None:
            raise ValidationError(
                f"数据块「{blocks[0].label}」的名称里没有单位信息，无法判定数值是"
                f"『间隔平均功率 kW』还是『间隔电量 kWh』；请在导入界面显式指定"
                f"（V2.2 §6.1：不得猜测口径）",
                field="load.import.value_kind",
            )
        messages.append(
            f"数据块「{blocks[0].label}」未显式指定口径，已按名称判定为「{label_kind.label}」"
        )
    kind = label_kind or LoadValueKind.INTERVAL_ENERGY_KWH
    if value_kind is not None:
        from_label = detect_value_kind_from_header(blocks[0].label)
        if from_label is not None and from_label is not value_kind:
            messages.append(
                f"提示：数据块名称「{blocks[0].label}」暗示「{from_label.label}」，"
                f"而界面指定为「{kind.label}」；已按界面指定处理，请确认无误"
            )

    # 逐块确定日期（真实资料里日期表头可能整体未更新，见 resolve_block_dates）
    resolved: list[list[date]] = []
    resolved_blocks: list[TimeBlock] = []
    extra_issues: list[DataQualityIssue] = []
    for block in blocks:
        month = sheet_month_hint(block.sheet_name)
        header_years = {d.year for d in block.dates}
        year = (
            next(iter(header_years))
            if len(header_years) == 1
            else (expected_year or (block.dates[0].year if block.dates else None))
        )
        if date_source != "header" and month is not None and year is not None:
            trimmed, dropped = trim_trailing_padding(block, month=month, year=year)
            if dropped:
                messages.append(
                    f"数据块「{block.label}」（工作表「{block.sheet_name}」）尾部有 {dropped} 列"
                    f"整列无正值（模板遗留列），已裁剪，不参与计算；裁剪后 {trimmed.day_count} 列"
                )
            block = trimmed
        dates, issue = resolve_block_dates(
            block, date_source=date_source, expected_year=expected_year
        )
        resolved.append(dates)
        resolved_blocks.append(block)
        if issue is not None:
            extra_issues.append(issue)
    if extra_issues:
        messages.append(
            f"有 {len(extra_issues)} 个数据块自带的日期表头与工作表名月份不一致"
            f"（疑似复制上月表头未更新），已按工作表名月份推导日期"
            f"（详见质量清单 [V25]）；如需按表头日期导入，请把日期口径改为 header"
        )

    detection = detect_interval(
        [
            datetime(day.year, day.month, day.day, moment.hour, moment.minute)
            for dates in resolved
            for day in dates
            for moment in blocks[0].times
        ]
    )
    minutes = require_interval(
        detection,
        declared_minutes=interval_minutes,
        allow_nonstandard=allow_nonstandard_interval,
        where=f"数据块「{blocks[0].label}」",
    )

    points: list[TimeSeriesPoint] = []
    for block, dates in zip(resolved_blocks, resolved):
        points.extend(build_points_from_block(block, value_field=value_field, dates=dates))
    points.sort(key=lambda point: point.timestamp)
    if not points:
        raise ValidationError(
            f"数据块「{blocks[0].label}」没有任何有效数据点", field="load.import.block"
        )

    alignment = align_timestamps(
        [point.timestamp for point in points], minutes, strategy=alignment_strategy
    )
    if alignment.strategy == "raw":
        raise ValidationError(
            f"数据块「{blocks[0].label}」的时间戳无法归整到 {minutes} 分钟网格："
            f"{'；'.join(alignment.notes)}；请检查日期表头与时刻列（V2.2 §6.1）",
            field="load.import.align",
        )
    delta_hours = minutes / 60.0
    if kind is LoadValueKind.POWER_KW:
        for point, aligned in zip(points, alignment.timestamps):
            setattr(point, "timestamp", aligned)
            point.load_kwh = float(point.load_kwh) * delta_hours
        messages.append(
            f"数据块「{blocks[0].label}」按间隔平均功率 kW 处理：已按 E = P × Δt"
            f"（Δt = {delta_hours:g} h）换算为电量（V2.2 §3.2）"
        )
    else:
        for point, aligned in zip(points, alignment.timestamps):
            setattr(point, "timestamp", aligned)
        messages.append(
            f"数据块「{blocks[0].label}」按间隔电量 kWh 处理：直接使用，未乘 Δt（V2.2 §3.2）"
        )

    if len(blocks) > 1:
        messages.append(
            f"已拼接 {len(blocks)} 张工作表的数据块（第 {block_index + 1} 块）："
            + "、".join(f"{block.sheet_name}｜{block.label}" for block in blocks)
        )

    rows: list[LoadImportRow] = [
        LoadImportRow(
            row_number=index + 1,
            timestamp_text=point.timestamp.strftime("%Y-%m-%d %H:%M"),
            value_text=f"{point.load_kwh:.6f}",
            status=LoadQualityStatus.VALID
            if np.isfinite(point.load_kwh)
            else LoadQualityStatus.WARNING,
        )
        for index, point in enumerate(points[:max_preview_rows])
    ]
    nan_count = int(sum(1 for point in points if not np.isfinite(point.load_kwh)))
    if nan_count:
        messages.append(
            f"数据块中有 {nan_count} 个空单元格（已按缺失处理，V11）；"
            f"缺失不得默认按 0 计算（V2.2 §6.4）"
        )

    return _build_preview(
        file_name=target.name,
        sheet_name=blocks[0].sheet_name if len(blocks) == 1 else "＋".join(
            dict.fromkeys(block.sheet_name for block in blocks)
        ),
        available_sheets=sheets,
        layout="time_block",
        block_label=blocks[0].label,
        available_blocks=catalog,
        column_mapping={},
        unmapped_headers=[],
        timestamp_column="时间（第 1 列）",
        value_column=f"{blocks[0].label}（数据块）",
        value_kind=kind,
        points=points,
        raw_timestamps=[point.timestamp for point in points],
        total_rows=sum(block.cell_count for block in blocks),
        expected_year=expected_year,
        alignment=alignment,
        interval_minutes=minutes,
        messages=messages,
        rows=rows,
        coverage_warning_ratio=coverage_warning_ratio,
        extra_issues=extra_issues,
    )


# --------------------------------------------------------------------------- #
# ⑥ 确认导入
# --------------------------------------------------------------------------- #
def apply_load_import(
    preview: LoadImportPreview,
    *,
    project_id: str = "",
    allow_invalid: bool = False,
    source_type: LoadDataSourceType | None = None,
    resample_to_minutes: int | None = None,
) -> HighFrequencyLoadDataset:
    """把预览结果落实为 :class:`HighFrequencyLoadDataset`（第 ⑥ 步：用户确认后调用）。

    :param allow_invalid: 预览质量为"无效"时是否仍允许导入（默认阻断，§0.2：不得静默接受坏数据）
    :param source_type: 数据来源标签；默认沿用预览。**估算数据不得标成实测**（§0.2 红线），
        模型构造函数会再次拦截 ``source_type`` 与 ``estimated`` 不自洽的组合
    :param resample_to_minutes: 导入后立即重采样到目标间隔（复用
        :func:`cenep.calculation.load_resample.resample_dataset`，口径与假设全程留痕）
    :raises ValidationError: 预览无效且未显式允许、或重采样失败（中文提示）
    """
    resolved_source = source_type or preview.source_type
    if preview.quality_status is LoadQualityStatus.INVALID and not allow_invalid:
        errors = "；".join(
            issue.message for issue in preview.quality_issues if issue.level == "ERROR"
        )
        raise ValidationError(
            f"负荷文件「{preview.file_name}」存在 {sum(1 for i in preview.quality_issues if i.level == 'ERROR')} "
            f"项阻断性问题，已停止导入：{errors}。"
            f"请修正原始数据后重试；如确认数据可用，可在导入界面显式选择『忽略无效标记继续导入』",
            field="load.import.invalid",
        )
    if preview.point_count == 0:
        raise ValidationError(
            f"负荷文件「{preview.file_name}」没有解析出任何时间点，无法导入",
            field="load.import.empty",
        )

    dataset = HighFrequencyLoadDataset(
        profile_id=make_dataset_id(Path(preview.file_name).stem, preview.interval_minutes),
        project_id=project_id,
        name=f"{Path(preview.file_name).stem}（{preview.interval_minutes} 分钟）",
        source_type=resolved_source,
        value_kind=preview.value_kind,
        interval_minutes=preview.interval_minutes,
        resolution=preview.resolution or Resolution.HOURLY,
        timezone=TIMEZONE_DEFAULT,
        period_start=preview.points[0].timestamp,
        period_end=preview.points[-1].timestamp
        + timedelta(minutes=preview.interval_minutes),
        points=list(preview.points),
        annualized=False,
        estimated=not resolved_source.is_measured,
        coverage_ratio=preview.coverage_ratio,
        missing_intervals=preview.missing_intervals,
        duplicate_intervals=preview.duplicate_intervals,
        irregular_intervals=preview.irregular_intervals,
        quality_status=preview.quality_status,
        quality_messages=[issue.message for issue in preview.quality_issues],
        source_file_name=preview.file_name,
        source_sheet=preview.sheet_name,
        mapping_config=dict(preview.column_mapping),
        assumptions=list(preview.messages),
        created_at=datetime.now(),
    )
    if resample_to_minutes is not None:
        dataset, plan = resample_dataset(
            dataset, resample_to_minutes, project_id=project_id or dataset.project_id
        )
        logger.info("导入后重采样：%s", plan.describe())
    logger.info(
        "负荷数据集导入完成：%s；%s", dataset.profile_id, dataset.quality_summary_text()
    )
    return dataset
