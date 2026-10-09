"""自检模式（供打包后的 EXE 验证与日常巡检使用）。

用法::

    python -m cenep --selftest                  # 结果打印到标准输出
    python -m cenep --selftest report.json      # 结果写入文件（--windowed 打包后无控制台，用这个）

自检内容：三类黄金项目各跑一次完整计算，再各导出一次 Excel 与 PDF，全部成功则返回码 0。

**重要**：自检会在每个阶段结束时把报告写回文件（``stage`` 字段），
这样即使进程被硬崩溃或弹出异常对话框，也能从文件看出卡在哪一步。
"""

from __future__ import annotations

import json
import sys
import tempfile
import traceback
from pathlib import Path

from .domain.enums import (
    InvestmentMode,
    OpexMode,
    ProjectType,
    RepaymentMethod,
    RoofRentMode,
    TariffMode,
)
from .domain.models import (
    BasicInfo,
    FinancingConfig,
    InvestmentConfig,
    LoadConfig,
    OpexConfig,
    Project,
    PVConfig,
    StorageConfig,
    TariffConfig,
    TaxConfig,
)


def _golden_project(project_type: ProjectType, name: str) -> Project:
    """与 tests/conftest.py 的黄金案例参数保持一致。"""
    return Project(
        basic_info=BasicInfo(project_name=name, province="湖北", city="武汉市", project_type=project_type),
        load=LoadConfig(
            annual_load_kwh=1_200_000.0,
            working_days=300,
            daytime_load_ratio=0.6,
            nighttime_load_ratio=0.4,
        ),
        pv=PVConfig(
            pv_capacity_kwp=1000.0,
            roof_area_m2=8000.0,
            usable_roof_area_m2=6500.0,
            area_per_kwp=6.0,
            equivalent_hours=1100.0,
            performance_ratio=1.0,
            annual_degradation_rate=0.005,
            curtailment_rate=0.0,
            self_consumption_ratio=0.8,
        ),
        storage=StorageConfig(
            storage_power_kw=500.0,
            storage_energy_kwh=1000.0,
            round_trip_efficiency=0.88,
            annual_cycles=330.0,
            depth_of_discharge=0.9,
            annual_degradation_rate=0.02,
        ),
        tariff=TariffConfig(
            tariff_mode=TariffMode.TOU,
            peak_price=1.0,
            flat_price=0.7,
            valley_price=0.4,
            peak_ratio=0.3,
            flat_ratio=0.4,
            valley_ratio=0.3,
            export_price=0.35,
        ),
        investment=InvestmentConfig(
            mode=InvestmentMode.UNIT_PRICE, pv_capex_per_kw=3000.0, storage_capex_per_kwh=1000.0
        ),
        opex=OpexConfig(
            pv_opex=0.02,
            pv_opex_mode=OpexMode.RATIO_OF_CAPEX,
            storage_opex=0.02,
            storage_opex_mode=OpexMode.RATIO_OF_CAPEX,
            insurance=0.005,
            insurance_mode=OpexMode.RATIO_OF_CAPEX,
            management_cost=10000.0,
            management_mode=OpexMode.FIXED,
            other_opex=5000.0,
            other_opex_mode=OpexMode.FIXED,
            roof_rent_mode=RoofRentMode.FIXED,
        ),
        tax=TaxConfig(
            depreciation_years=20,
            residual_value_ratio=0.05,
            depreciable_capex_ratio=1.0,
            income_tax_rate=0.25,
        ),
        financing=FinancingConfig(
            enabled=True,
            debt_ratio=0.5,
            equity_ratio=0.5,
            interest_rate=0.04,
            loan_term=10,
            grace_period=0,
            repayment_method=RepaymentMethod.EQUAL_PRINCIPAL,
        ),
        analysis_period=25,
        discount_rate=0.08,
    )


def _dump(report: dict, output: str | None) -> None:
    """把当前报告写出去（每个阶段都调用，便于定位崩溃点）。"""
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if output:
        try:
            Path(output).write_text(text, encoding="utf-8")
        except OSError:
            pass
        return
    try:
        sys.stdout.write(text + "\n")
        sys.stdout.flush()
    except Exception:  # pragma: no cover - --windowed 下没有标准输出
        pass


def run_selftest(output: str | None = None) -> int:
    """执行自检，返回 0 表示全部通过。"""
    report: dict = {
        "ok": False,
        "stage": "start",
        "checks": [],
        "projects": [],
        "errors": [],
    }
    _dump(report, output)

    try:
        from openpyxl import load_workbook

        from .calculation.engine import calculation_engine
        from .domain.enums import ScenarioType, SensitivityVariable
        from .reports.excel_exporter import SHEET_NAMES, ExcelExporter
        from .reports.pdf_exporter import REPORT_SECTIONS, PdfExporter

        report["stage"] = "imports-ok"
        _dump(report, output)

        with tempfile.TemporaryDirectory() as tmp:
            out_dir = Path(tmp)
            for project_type in (
                ProjectType.COMMERCIAL_PV,
                ProjectType.COMMERCIAL_STORAGE,
                ProjectType.PV_STORAGE,
            ):
                report["stage"] = f"{project_type.value}:calculate"
                _dump(report, output)
                project = _golden_project(project_type, f"自检-{project_type.label}")
                result = calculation_engine.calculate(project)

                # 能量守恒（规范 §113）
                for row in result.annual_results:
                    lhs = row.pv_generation_kwh
                    rhs = row.pv_self_use_kwh + row.pv_to_storage_kwh + row.pv_export_kwh + row.pv_loss_kwh
                    assert abs(lhs - rhs) <= 1e-6, f"{project_type} 第 {row.year} 年电量不守恒"

                # 情景与敏感性（规范 §92、§95）
                assert [s.scenario for s in result.scenarios] == [
                    ScenarioType.BASE.value,
                    ScenarioType.CONSERVATIVE.value,
                    ScenarioType.OPTIMISTIC.value,
                ]
                if project_type.has_storage:
                    assert SensitivityVariable.STORAGE_CYCLES.value in {r.variable for r in result.sensitivity}
                if project_type.has_pv:
                    assert SensitivityVariable.GENERATION.value in {r.variable for r in result.sensitivity}

                report["stage"] = f"{project_type.value}:excel"
                _dump(report, output)
                excel_path = ExcelExporter().export(project, result, out_dir / f"{project_type.value}")

                report["stage"] = f"{project_type.value}:pdf"
                _dump(report, output)
                pdf_path = PdfExporter().export(project, result, out_dir / f"{project_type.value}")

                report["stage"] = f"{project_type.value}:verify"
                _dump(report, output)
                sheets = load_workbook(excel_path).sheetnames
                assert sheets == SHEET_NAMES, "Excel 工作表不符合规范（V1 §108 / V2 §67 / V2.1 §8.1）"
                assert pdf_path.stat().st_size > 5000, "PDF 输出过小"
                # V2 §66 由 15 章重组为 16 部分；V2.1 §8.1（阶段 2）再新增
                # 「账单事实与校验」章节，共 17 部分
                assert len(REPORT_SECTIONS) == 17
                # V2 §67 工作表由 13 张扩展为 24 张；V2.1 §8.1 再新增
                # 「账单原始数据」「账单校验」两张，共 26 张（V1 的 13 张全部保留）
                assert len(sheets) == 26

                report["projects"].append(
                    {
                        "type": project_type.value,
                        "label": project_type.label,
                        "total_capex": result.total_capex,
                        "project_irr": result.project_irr,
                        "project_npv": result.project_npv,
                        "static_payback": result.static_payback,
                        "lcoe": result.lcoe,
                        "lcos": result.lcos,
                        "excel_sheets": len(sheets),
                        "pdf_bytes": pdf_path.stat().st_size,
                    }
                )
                report["checks"].append(f"{project_type.label}：计算 + Excel + PDF 全部通过")
                _dump(report, output)

        report["ok"] = True
        report["stage"] = "done"
    except Exception as exc:  # noqa: BLE001 - 自检需要把失败原因完整带出来
        report["errors"].append(f"{type(exc).__name__}: {exc}")
        report["traceback"] = traceback.format_exc()
        report["stage"] = f"{report['stage']}:FAILED"
    finally:
        _dump(report, output)

    return 0 if report["ok"] else 1


__all__ = ["run_selftest"]
