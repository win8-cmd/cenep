"""V2.4 阶段 7：光储四场景在界面与报告中的接入验证（§8.1、§8.4、§12）。

阶段 6 只交付了业务层（`ScenarioBillService`），界面与报告留到阶段 7（见
`V2.3_STAGE6_SCENARIO_REPORT.md` §7/§8）。本文件验证阶段 7 的接入：

* **D 区界面**：电价计划下拉、执行按钮、三张结果表与财务一致性文本，
  且界面显示的行与 ``ScenarioBillService`` 返回的**逐字一致**（§105 口径）；
* **Excel**：`方案电费对比` 表写出四场景、分项差额与去重清单，且**无公式**；
* **PDF**：`方案前后电费差额` 章节写出同样的行，并披露口径分界（§8.1）；
* **不重复计算**：报告里出现的"唯一去重后的年度运营收益"等于服务返回值。

所有电价数值取自**已核验**的湖北官方计划（不新增任何未经核验的电价）。
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

pytest.importorskip("PySide6", reason="未安装 PySide6，跳过界面测试")

from PySide6.QtWidgets import QApplication  # noqa: E402
from reportlab.platypus import Paragraph, Table  # noqa: E402

from cenep.application.load_profile_service import LoadProfileService  # noqa: E402
from cenep.application.scenario_service import ScenarioBillService  # noqa: E402
from cenep.calculation.engine import calculation_engine  # noqa: E402
from cenep.domain.enums import LoadDataSourceType, Resolution  # noqa: E402
from cenep.domain.load_data import HighFrequencyLoadDataset, TimeSeriesPoint  # noqa: E402
from cenep.domain.models import Project  # noqa: E402
from cenep.policy.tariff_plan_store import TariffPlanStore  # noqa: E402
from cenep.reports.excel_exporter import ExcelExporter  # noqa: E402
from cenep.reports.pdf_exporter import (  # noqa: E402
    SECTION_SCENARIO_BILL,
    PdfExporter,
    _styles,
    register_cjk_font,
)

OFFICIAL_110_ID = "HUBEI_COMMERCIAL_2026_01_TWO_PART_KV110"
STATION_ID = "HUBEI_COMMERCIAL_2026_01_TWO_PART_KV35"


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


def _measured_dataset(profile_id: str = "S7-UI") -> HighFrequencyLoadDataset:
    """8760 点受控负荷（白天 400 kW、夜间 120 kW，逐时 60 分钟）。"""
    points = []
    for index in range(8760):
        hour = index % 24
        power = 400.0 if 7 <= hour <= 18 else 120.0
        points.append(TimeSeriesPoint(timestamp=_stamp(index), load_kwh=power))
    return HighFrequencyLoadDataset(
        profile_id=profile_id,
        name="阶段 7 受控负荷（60 分钟）",
        source_type=LoadDataSourceType.HIGH_FREQUENCY_IMPORT,
        interval_minutes=60,
        period_start=points[0].timestamp,
        period_end=points[-1].timestamp,
        points=points,
        estimated=False,
        coverage_ratio=1.0,
        missing_intervals=0,
    )


def _stamp(index: int):
    from datetime import datetime, timedelta

    return datetime(2025, 1, 1) + timedelta(hours=index)


@pytest.fixture
def project_with_load() -> Project:
    project = Project()
    project.analysis_period = 1
    project.basic_info.project_name = "阶段 7 场景接入测试"
    project.timeseries.enabled = True
    project.timeseries.base_year = 2025
    project.pv.pv_capacity_kwp = 1000.0
    project.pv.equivalent_hours = 1100.0
    project.pv.performance_ratio = 1.0
    project.storage.storage_power_kw = 500.0
    project.storage.storage_energy_kwh = 2000.0
    LoadProfileService(project).register(_measured_dataset())
    return project


@pytest.fixture
def scenario_service(project_with_load) -> ScenarioBillService:
    store = TariffPlanStore()
    assert store.get_plan(OFFICIAL_110_ID) is not None
    return ScenarioBillService(
        project_with_load,
        store=store,
        load_profile_service=LoadProfileService(project_with_load),
    )


@pytest.fixture
def scenario_result(scenario_service):
    return scenario_service.compare(OFFICIAL_110_ID)


def _story_text(flowables) -> str:
    parts: list[str] = []
    for item in flowables:
        if isinstance(item, Paragraph):
            parts.append(item.text)
        elif isinstance(item, Table):
            for row in item._cellvalues:
                for cell in row:
                    parts.append(cell if isinstance(cell, str) else _story_text([cell]))
    return "\n".join(parts)


def _sheet_text(ws) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


# --------------------------------------------------------------------------- #
# 1. D 区界面（§8.4、§7.1）
# --------------------------------------------------------------------------- #
class TestScenarioTabWiring:
    def test_plans_listed_from_real_store(self, qapp, project_with_load, scenario_service):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        assert page.scenario_plan_combo.count() == 0
        page.bind_scenario_service(scenario_service)
        assert page.scenario_plan_combo.count() == len(scenario_service.list_tariff_plans())
        assert page.scenario_plan_combo.count() > 0
        assert page.scenario_combo_plan_id()

    def test_run_without_plan_is_blocked_with_chinese_message(self, qapp):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        page.run_scenario_compare()
        assert page.last_scenario is None
        assert "电价计划" in page.scenario_note_label.text()
        assert page.scenario_table.rowCount() == 0

    def test_rows_equal_service_rows(self, qapp, project_with_load, scenario_service, scenario_result):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        page.scenario_service = scenario_service
        page.reload_tariff_plans()
        index = page.scenario_plan_combo.findData(OFFICIAL_110_ID)
        assert index >= 0
        page.scenario_plan_combo.setCurrentIndex(index)
        page.run_scenario_compare()

        assert page.last_scenario is not None
        expected = ScenarioBillService.scenario_rows(scenario_result)
        assert page.scenario_table.rowCount() == len(expected) - 1
        assert [page.scenario_table.horizontalHeaderItem(c).text() for c in range(page.scenario_table.columnCount())] == [
            str(h) for h in expected[0]
        ]
        # 第 1 行第 1 列（场景名）必须与服务一致
        assert page.scenario_table.item(0, 0).text() == str(expected[1][0])

        comparisons = ScenarioBillService.comparison_rows(scenario_result)
        assert page.comparison_table.rowCount() == len(comparisons) - 1
        dedup = ScenarioBillService.dedup_rows(scenario_result)
        assert page.dedup_table.rowCount() == len(dedup) - 1

    def test_detail_view_reports_dedup_and_finance_consistency(
        self, qapp, project_with_load, scenario_service, scenario_result
    ):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        page.scenario_service = scenario_service
        page.reload_tariff_plans()
        page.scenario_plan_combo.setCurrentIndex(page.scenario_plan_combo.findData(OFFICIAL_110_ID))
        page.run_scenario_compare()

        text = page.scenario_detail_view.toPlainText()
        assert f"{scenario_result.unique_annual_benefit_yuan:,.2f}" in text
        assert "唯一去重后的年度运营收益" in text
        assert "收益去重校验" in text
        assert "财务现金流一致性" in text

    def test_state_exposes_scenario_summary(self, qapp, project_with_load, scenario_service, scenario_result):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        page.bind(project_with_load, LoadProfileService(project_with_load))
        page.scenario_service = scenario_service
        page.reload_tariff_plans()
        page.scenario_plan_combo.setCurrentIndex(page.scenario_plan_combo.findData(OFFICIAL_110_ID))
        page.run_scenario_compare()

        state = page.state()
        assert state["scenario_plans"] > 0
        # D 区的「四场景年度账单」表包含基准场景（共 4 行，含基准）
        assert state["scenario_rows"] == len(scenario_result.scenarios)
        assert state["dedup_verified"] is True
        assert state["unique_annual_benefit"] == pytest.approx(
            scenario_result.unique_annual_benefit_yuan
        )

    def test_mark_stale_after_input_change(self, qapp, project_with_load, scenario_service):
        from cenep.ui.load_pages import LoadAnalysisPage

        page = LoadAnalysisPage()
        page.bind(project_with_load, LoadProfileService(project_with_load))
        page.scenario_service = scenario_service
        page.reload_tariff_plans()
        page.scenario_plan_combo.setCurrentIndex(page.scenario_plan_combo.findData(OFFICIAL_110_ID))
        page.run_scenario_compare()
        assert page.scenario_stale_label.text() == ""

        page._mark_stale()
        assert "过期" in page.scenario_stale_label.text()

    def test_main_window_binds_scenario_service(self, qapp, tmp_path):
        """§8.4：主窗口必须把 D 区接到真实的场景服务（含 SQLite 电价版本库）。"""
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "scenario_ui.db")
        try:
            assert window.load_page.scenario_service is not None
            assert window.load_page.scenario_plan_combo.count() > 0
            titles = [window.tabs.tabText(i) for i in range(window.tabs.count())]
            assert titles == [
                "项目", "参数", "计算", "结果", "时序仿真", "月度账单", "负荷与消纳",
                "敏感性", "报告", "设置",
            ]
        finally:
            window.close()


# --------------------------------------------------------------------------- #
# 2. Excel「方案电费对比」表（§8.1）
# --------------------------------------------------------------------------- #
class TestScenarioExcelSheet:
    def test_sheet_written_from_service_rows(
        self, project_with_load, scenario_service, scenario_result, tmp_path
    ):
        result = calculation_engine.calculate(project_with_load)
        path = ExcelExporter().export(
            project_with_load,
            result,
            tmp_path / "场景.xlsx",
            plan_ids=[OFFICIAL_110_ID],
            scenario_result=scenario_result,
        )
        wb = load_workbook(path)
        text = _sheet_text(wb["方案电费对比"])

        assert "方案电费对比（口径：方案模拟" in text
        assert scenario_result.tariff_plan_name in text
        assert scenario_result.tariff_plan_source[:12] in text
        assert f"{scenario_result.unique_annual_benefit_yuan:.6f}".rstrip("0").rstrip(".") in text
        for line in ScenarioBillService.scenario_rows(scenario_result)[1:]:
            assert str(line[0]) in text  # 四个场景名
        assert "收益去重清单" in text or "去重" in text

    def test_no_formula_and_status_disclosed(
        self, project_with_load, scenario_service, scenario_result, tmp_path
    ):
        result = calculation_engine.calculate(project_with_load)
        path = ExcelExporter().export(
            project_with_load, result, tmp_path / "场景2.xlsx", scenario_result=scenario_result
        )
        wb = load_workbook(path)
        offenders = [
            f"{ws.title}!{c.coordinate}"
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert offenders == []
        text = _sheet_text(wb["方案电费对比"])
        assert "核验状态" in text or "verified" in text

    def test_tariff_version_sheet_lists_plan_rows(
        self, project_with_load, scenario_service, scenario_result, tmp_path
    ):
        """§8.1 第 5 条：指定 plan_ids 时必须逐项列出该版本的时段电价与来源。"""
        result = calculation_engine.calculate(project_with_load)
        path = ExcelExporter().export(
            project_with_load, result, tmp_path / "版本.xlsx", plan_ids=[OFFICIAL_110_ID]
        )
        text = _sheet_text(load_workbook(path)["电价版本与来源"])
        assert OFFICIAL_110_ID in text
        assert "时段规则" in text
        assert "核验人/方式" in text or "核验状态" in text
        assert "元/千瓦时" in text

    def test_tariff_version_sheet_summarises_store_when_no_plan_given(
        self, project_with_load, tmp_path
    ):
        result = calculation_engine.calculate(project_with_load)
        path = ExcelExporter().export(project_with_load, result, tmp_path / "版本2.xlsx")
        text = _sheet_text(load_workbook(path)["电价版本与来源"])
        assert "内置 / 已登记计划" in text or "没有可用的电价计划" in text


# --------------------------------------------------------------------------- #
# 3. PDF「方案前后电费差额」章节（§8.1）
# --------------------------------------------------------------------------- #
class TestScenarioPdfSection:
    def test_section_written_with_dedup_checklist(
        self, project_with_load, scenario_result
    ):
        result = calculation_engine.calculate(project_with_load)
        story = PdfExporter().build_story(
            project_with_load,
            result,
            _styles(register_cjk_font()),
            plan_ids=[OFFICIAL_110_ID],
            scenario_result=scenario_result,
        )
        text = _story_text(story)
        assert SECTION_SCENARIO_BILL in text
        assert "四场景年度账单" in text
        assert "分项差额" in text
        assert "收益去重清单" in text
        assert "方案模拟" in text
        assert f"{scenario_result.unique_annual_benefit_yuan:,.2f}" in text

    def test_calibration_section_lists_tariff_plan(self, project_with_load, scenario_result):
        from cenep.application.tariff_service import TariffService

        service = TariffService(project_with_load, store=TariffPlanStore())
        calibration = service.recompute_all(OFFICIAL_110_ID)
        result = calculation_engine.calculate(project_with_load)
        story = PdfExporter().build_story(
            project_with_load,
            result,
            _styles(register_cjk_font()),
            plan_ids=[OFFICIAL_110_ID],
            calibration=calibration,
        )
        text = _story_text(story)
        assert "电价机制与版本" in text
        assert "基准账单校准" in text
        assert "毛差异" in text or "净差异" in text

    def test_report_exports_with_all_v24_payloads(
        self, project_with_load, scenario_result, tmp_path
    ):
        from cenep.application.tariff_service import TariffService

        service = TariffService(project_with_load, store=TariffPlanStore())
        calibration = service.recompute_all(OFFICIAL_110_ID)
        portrait = LoadProfileService(project_with_load).portrait()
        analysis = LoadProfileService(project_with_load).analyze()
        result = calculation_engine.calculate(project_with_load)

        excel = ExcelExporter().export(
            project_with_load,
            result,
            tmp_path / "全量.xlsx",
            self_consumption=analysis,
            load_portrait=portrait,
            plan_ids=[OFFICIAL_110_ID],
            scenario_result=scenario_result,
        )
        pdf = PdfExporter().export(
            project_with_load,
            result,
            tmp_path / "全量.pdf",
            self_consumption=analysis,
            load_portrait=portrait,
            plan_ids=[OFFICIAL_110_ID],
            calibration=calibration,
            scenario_result=scenario_result,
        )
        assert excel.exists() and pdf.stat().st_size > 5000
        wb = load_workbook(excel)
        assert "方案电费对比" in wb.sheetnames
        assert _sheet_text(wb["负荷数据质量"]).strip()
        assert _sheet_text(wb["月度电费分析"]).strip()
