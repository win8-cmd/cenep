"""V2.2 阶段 3：高频负荷导入、时间间隔识别与数据质量测试（规格书 §3.2、§6.1、§6.4、§9.2）。

覆盖两条主线
------------
1. **长表**（时间列 + 功率 kW / 间隔电量 kWh）：列映射、口径判定、间隔识别、
   15/30/60 分钟归一、闰年点数、缺失/重复/乱序/负值/异常值、时区与夏令时、
   年份不一致、覆盖率不足、平闰年点数不符、确定性；
2. **宽表**（"时刻行 × 日期列"矩阵）：数据块识别、多表拼接、日期表头未更新（V25）、
   尾部模板列裁剪、月份与列数矛盾、口径无法判定时拒绝导入。

全部测试数据在 ``tmp_path`` 内生成，**不依赖任何外部文件**（真实资料验证放在
``taiqu-storage/stage3_real_data_check.py``，其结构特征已在下面的用例中复现）。
"""

from __future__ import annotations

import csv
import math
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from cenep.calculation.errors import ValidationError
from cenep.calculation.load_resample import (
    downsample_energy,
    downsample_power,
    energy_from_power,
    power_from_energy,
    resample_dataset,
    resample_energy,
    resample_power,
    to_energy_kwh,
    to_power_kw,
    upsample_energy,
    upsample_power,
)
from cenep.calculation.timeseries_engine import build_time_axis, points_per_year
from cenep.data import (
    apply_load_import,
    build_points_from_block,
    check_load_quality,
    detect_interval,
    detect_value_kind_from_header,
    expected_year_from_file_name,
    load_block_catalog,
    preview_block_import,
    preview_load_import,
    read_time_blocks,
    resolve_load_columns,
    resolve_value_kind,
    scan_timezone_issues,
)
from cenep.domain.enums import (
    LoadDataSourceType,
    LoadQualityStatus,
    LoadValueKind,
    Resolution,
)
from cenep.domain.load_data import HighFrequencyLoadDataset, provides_measured_curve
from cenep.domain.timeseries import TimeSeriesPoint

ISO = "%Y-%m-%d %H:%M:%S"


# --------------------------------------------------------------------------- #
# 夹具构造
# --------------------------------------------------------------------------- #
def write_xlsx(path: Path, rows: list[list], sheet_name: str = "Sheet1") -> Path:
    """写一个只有一张工作表、逐行写入的 xlsx（首行即表头）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def write_wide_xlsx(
    path: Path,
    sheets: dict[str, list[list]],
) -> Path:
    """写一个多工作表 xlsx（用于"时刻行 × 日期列"宽表）。"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def day_stamps(day: date, minutes: int) -> list[datetime]:
    count = 24 * 60 // minutes
    return [
        datetime(day.year, day.month, day.day) + timedelta(minutes=minutes * index)
        for index in range(count)
    ]


def long_table_rows(
    start: datetime,
    count: int,
    minutes: int,
    *,
    header: tuple[str, str] = ("时间", "负荷(kWh)"),
    value: float = 10.0,
) -> list[list]:
    rows: list[list] = [list(header)]
    for index in range(count):
        rows.append(
            [(start + timedelta(minutes=minutes * index)).strftime(ISO), value]
        )
    return rows


def wide_block_rows(
    *,
    label: str | None,
    dates: list[date],
    minutes: int,
    value: float,
    day_header_dates: list[date] | None = None,
    zero_tail: int = 0,
) -> list[list]:
    """构造"时刻行 × 日期列"数据块（可选标签行 + 表头日期行 + 96/48/24 行时刻）。

    ``zero_tail``：末尾若干列填 0（复现真实资料"模板多出的列全为 0"）。
    """
    per_day = 24 * 60 // minutes
    header_dates = dates if day_header_dates is None else day_header_dates
    rows: list[list] = []
    if label is not None:
        rows.append([label])
    rows.append(["时间"] + [d.toordinal() - 693594 for d in header_dates])
    for index in range(per_day):
        moment = datetime(2000, 1, 1) + timedelta(minutes=minutes * index)
        line = [value] * len(header_dates)
        for offset in range(zero_tail):
            line[-1 - offset] = 0.0
        rows.append([moment.time(), *line])
    return rows


# --------------------------------------------------------------------------- #
# ① 列映射与数值口径
# --------------------------------------------------------------------------- #
class TestLoadColumnMapping:
    def test_maps_chinese_and_english_headers(self) -> None:
        mapping = resolve_load_columns(["时间戳", "间隔电量(kWh)", "备注"])
        assert mapping.mapping["timestamp"] == "时间戳"
        assert mapping.mapping["energy_kwh"] == "间隔电量(kWh)"
        assert mapping.unmapped_headers == ["备注"]

    def test_maps_power_column_and_ignores_other_columns(self) -> None:
        mapping = resolve_load_columns(["DATE TIME", "有功功率(kW)", "月", "日"])
        assert mapping.mapping["timestamp"] == "DATE TIME"
        assert mapping.mapping["power_kw"] == "有功功率(kW)"

    def test_explicit_mapping_wins(self) -> None:
        mapping = resolve_load_columns(
            ["A", "B"], explicit={"timestamp": "A", "energy_kwh": "B"}
        )
        assert mapping.mapping == {"timestamp": "A", "energy_kwh": "B"}

    def test_missing_timestamp_reports_chinese_error(self) -> None:
        with pytest.raises(ValidationError) as excinfo:
            resolve_load_columns(["负荷(kWh)", "备注"])
        assert "缺少时间列" in str(excinfo.value)

    def test_missing_value_column_reports_chinese_error(self) -> None:
        with pytest.raises(ValidationError) as excinfo:
            resolve_load_columns(["时间", "备注"])
        assert "缺少数值列" in str(excinfo.value)

    def test_explicit_mapping_of_unknown_header_reports_error(self) -> None:
        with pytest.raises(ValidationError) as excinfo:
            resolve_load_columns(["时间", "负荷"], explicit={"energy_kwh": "不存在"})
        assert "不存在" in str(excinfo.value)


class TestValueKind:
    @pytest.mark.parametrize(
        ("header", "expected"),
        [
            ("间隔电量(kWh)", LoadValueKind.INTERVAL_ENERGY_KWH),
            ("每日每小时用电量", LoadValueKind.INTERVAL_ENERGY_KWH),
            ("负荷(kW)", LoadValueKind.POWER_KW),
            ("有功功率", LoadValueKind.POWER_KW),
            ("用电功率(kW)", LoadValueKind.POWER_KW),
            ("温度(℃)", None),
            ("负荷", None),
            ("", None),
        ],
    )
    def test_detect_from_header(self, header: str, expected) -> None:
        assert detect_value_kind_from_header(header) is expected

    def test_declared_kind_is_used(self) -> None:
        kind, notices = resolve_value_kind(
            "负荷", LoadValueKind.POWER_KW, where="负荷表"
        )
        assert kind is LoadValueKind.POWER_KW
        assert notices == []

    def test_header_unit_is_used_with_notice(self) -> None:
        kind, notices = resolve_value_kind("每日每小时用电量", None, where="负荷表")
        assert kind is LoadValueKind.INTERVAL_ENERGY_KWH
        assert notices and "已按表头单位判定" in notices[0]

    def test_declared_conflicts_with_header_reports_error(self) -> None:
        with pytest.raises(ValidationError) as excinfo:
            resolve_value_kind("负荷(kWh)", LoadValueKind.POWER_KW, where="负荷表")
        assert "矛盾" in str(excinfo.value)

    def test_undetectable_kind_reports_error_not_guess(self) -> None:
        """§6.1：口径判不出来必须报中文错误，不得猜。"""
        with pytest.raises(ValidationError) as excinfo:
            resolve_value_kind("负荷", None, where="负荷表")
        assert "无法判定" in str(excinfo.value)
        assert "不得猜测" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# ② 时间间隔识别
# --------------------------------------------------------------------------- #
class TestIntervalDetection:
    @pytest.mark.parametrize("minutes", [15, 30, 60])
    def test_detects_standard_intervals(self, minutes: int) -> None:
        stamps = day_stamps(date(2025, 3, 1), minutes)
        detection = detect_interval(stamps)
        assert detection.interval_minutes == minutes
        assert detection.regular
        assert detection.resolution is Resolution.from_interval_minutes(minutes)

    def test_tolerates_sub_second_jitter(self) -> None:
        """真实资料（Excel 浮点序列号）会把整点写成 1:59:59.99，必须容差而非判为不规整。"""
        stamps = [
            datetime(2025, 1, 1, 0, 0, 0),
            datetime(2025, 1, 1, 1, 0, 0),
            datetime(2025, 1, 1, 1, 59, 59, 990000),
            datetime(2025, 1, 1, 2, 59, 59, 985000),
        ]
        detection = detect_interval(stamps)
        assert detection.interval_minutes == 60
        assert detection.regular_ratio == pytest.approx(1.0, abs=1e-9)

    def test_irregular_intervals_are_not_guessed(self) -> None:
        stamps = [
            datetime(2025, 1, 1, 0, 0),
            datetime(2025, 1, 1, 0, 37),
            datetime(2025, 1, 1, 2, 5),
            datetime(2025, 1, 1, 2, 11),
            datetime(2025, 1, 1, 5, 49),
        ]
        detection = detect_interval(stamps)
        assert detection.interval_minutes is None
        assert not detection.regular
        assert "不规整" in detection.note or "无法判定" in detection.note

    def test_single_point_returns_none(self) -> None:
        detection = detect_interval([datetime(2025, 1, 1)])
        assert detection.interval_minutes is None
        assert "不足 2 个" in detection.note

    def test_duplicate_timestamps_do_not_break_interval_detection(self) -> None:
        """重复时间戳由 V02 报告，不应把整份文件的间隔判成"无法判定"。"""
        stamps = [
            datetime(2025, 1, 1, 0, 0),
            datetime(2025, 1, 1, 1, 0),
            datetime(2025, 1, 1, 1, 0),
            datetime(2025, 1, 1, 2, 0),
        ]
        detection = detect_interval(stamps)
        assert detection.interval_minutes == 60
        assert detection.regular
        assert "重复/倒序" in detection.note

    def test_missing_points_are_treated_as_gaps_not_irregular(self) -> None:
        """缺时间点（Δ = k×主间隔）由 V03/V04 报告，不阻断间隔识别。"""
        stamps = [
            datetime(2025, 1, 1, 0, 0),
            datetime(2025, 1, 1, 1, 0),
            datetime(2025, 1, 1, 3, 0),  # 缺 02:00
            datetime(2025, 1, 1, 4, 0),
        ]
        detection = detect_interval(stamps)
        assert detection.interval_minutes == 60
        assert detection.regular
        assert "缺口" in detection.note


# --------------------------------------------------------------------------- #
# ③ 长表导入
# --------------------------------------------------------------------------- #
class TestLongTableImport:
    def test_15min_kwh_import(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "load15.xlsx",
            long_table_rows(datetime(2025, 1, 1), 96, 15, value=8.0),
        )
        preview = preview_load_import(path)
        assert preview.layout == "table"
        assert preview.interval_minutes == 15
        assert preview.value_kind is LoadValueKind.INTERVAL_ENERGY_KWH
        assert preview.point_count == 96
        assert preview.annual_energy_kwh == pytest.approx(96 * 8.0)
        assert preview.coverage_ratio == pytest.approx(96 / 35040)
        # 15 分钟只导入一天 → 覆盖率不足必须告警（§2.2）
        assert preview.quality_status is LoadQualityStatus.WARNING
        assert any("[V24]" in issue.message for issue in preview.quality_issues)

    def test_30min_power_import_multiplies_delta_t(self, tmp_path: Path) -> None:
        """§3.2：功率 → 电量必须乘 Δt（30 分钟 = 0.5 h），不得直接累加。"""
        path = write_xlsx(
            tmp_path / "load30_power.xlsx",
            long_table_rows(
                datetime(2025, 1, 1), 48, 30, header=("时间", "有功功率(kW)"), value=100.0
            ),
        )
        preview = preview_load_import(path)
        assert preview.interval_minutes == 30
        assert preview.resolution is Resolution.HALF_HOURLY
        assert preview.value_kind is LoadValueKind.POWER_KW
        assert preview.annual_energy_kwh == pytest.approx(48 * 100.0 * 0.5)
        assert any("E = P × Δt" in message for message in preview.messages)

    def test_60min_full_year_has_no_coverage_warning(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "load60.xlsx",
            long_table_rows(datetime(2025, 1, 1), 8760, 60, value=10.0),
        )
        preview = preview_load_import(path)
        assert preview.point_count == 8760
        assert preview.coverage_ratio == pytest.approx(1.0)
        assert not any("[V24]" in i.message for i in preview.quality_issues)
        assert preview.missing_intervals == 0

    def test_header_row_two_is_supported(self, tmp_path: Path) -> None:
        rows = [["某某项目负荷数据（第一行是标题）"], *long_table_rows(
            datetime(2025, 1, 1), 24, 60, value=5.0
        )]
        path = write_xlsx(tmp_path / "load_title.xlsx", rows)
        preview = preview_load_import(path, header_row=2)
        assert preview.point_count == 24

    def test_csv_is_supported(self, tmp_path: Path) -> None:
        path = tmp_path / "load.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(long_table_rows(datetime(2025, 1, 1), 24, 60, value=3.0))
        preview = preview_load_import(path)
        assert preview.point_count == 24
        assert preview.annual_energy_kwh == pytest.approx(72.0)

    def test_unsorted_timestamps_are_sorted(self, tmp_path: Path) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 24, 60, value=1.0)
        rows[1], rows[5] = rows[5], rows[1]  # 制造乱序
        path = write_xlsx(tmp_path / "unsorted.xlsx", rows)
        preview = preview_load_import(path)
        stamps = [point.timestamp for point in preview.points]
        assert stamps == sorted(stamps)
        assert any("[V05]" in issue.message for issue in preview.quality_issues)

    def test_duplicate_timestamps_block_import_by_default(self, tmp_path: Path) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 24, 60, value=1.0)
        rows.append([rows[1][0], 1.0])
        path = write_xlsx(tmp_path / "dup.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.quality_status is LoadQualityStatus.INVALID
        assert any("[V02]" in issue.message for issue in preview.quality_issues)
        with pytest.raises(ValidationError) as excinfo:
            apply_load_import(preview)
        assert "阻断性问题" in str(excinfo.value)
        dataset = apply_load_import(preview, allow_invalid=True)
        assert dataset.quality_status is LoadQualityStatus.INVALID

    def test_negative_value_is_reported(self, tmp_path: Path) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 24, 60, value=1.0)
        rows[3][1] = -5.0
        path = write_xlsx(tmp_path / "negative.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.quality_status is LoadQualityStatus.INVALID
        assert any("[V08]" in issue.message for issue in preview.quality_issues)

    def test_empty_value_becomes_missing_not_zero(self, tmp_path: Path) -> None:
        """§6.4：缺失不得默认按 0 处理。"""
        rows = long_table_rows(datetime(2025, 1, 1), 24, 60, value=1.0)
        rows[4][1] = None
        path = write_xlsx(tmp_path / "blank.xlsx", rows)
        preview = preview_load_import(path)
        assert any("[V11]" in issue.message for issue in preview.quality_issues)
        values = [point.load_kwh for point in preview.points]
        assert math.isnan(values[3])  # 第 4 个点保持缺失
        assert preview.annual_energy_kwh == pytest.approx(23 * 1.0)  # 缺失不计入合计

    def test_unparsable_timestamp_row_is_reported_and_excluded(
        self, tmp_path: Path
    ) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 24, 60, value=1.0)
        rows[2][0] = "不是时间"
        path = write_xlsx(tmp_path / "bad_ts.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.total_rows == 24
        assert preview.invalid_count == 1
        assert preview.point_count == 23
        assert any("无法解析" in row.messages[0] for row in preview.rows if row.messages)

    def test_no_parsable_timestamp_reports_chinese_error(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "all_bad.xlsx",
            [["时间", "负荷(kWh)"], ["x", 1.0], ["y", 2.0]],
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path)
        assert "没有任何一行能解析成时间" in str(excinfo.value)

    def test_interval_not_detectable_reports_chinese_error(self, tmp_path: Path) -> None:
        rows = [["时间", "负荷(kWh)"]]
        for moment in ("2025-01-01 00:00", "2025-01-01 00:37", "2025-01-01 02:05"):
            rows.append([moment, 1.0])
        path = write_xlsx(tmp_path / "irregular.xlsx", rows)
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path)
        assert "无法判定时间间隔" in str(excinfo.value)

    def test_declared_interval_resolves_jittered_data(self, tmp_path: Path) -> None:
        """时间列有秒级抖动时，用户显式指定间隔仍可导入（网格对齐 + V20 留痕）。"""
        rows = [
            ["时间", "负荷(kWh)"],
            ["2025-01-01 00:00:00", 1.0],
            ["2025-01-01 00:30:20", 1.0],
            ["2025-01-01 01:00:00", 1.0],
        ]
        path = write_xlsx(tmp_path / "jitter.xlsx", rows)
        preview = preview_load_import(path, interval_minutes=30)
        assert preview.interval_minutes == 30
        assert [p.timestamp.strftime("%H:%M") for p in preview.points] == [
            "00:00",
            "00:30",
            "01:00",
        ]
        assert any("[V20]" in issue.message for issue in preview.quality_issues)

    def test_declared_interval_conflict_reports_error(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "conflict.xlsx",
            long_table_rows(datetime(2025, 1, 1), 96, 15, value=1.0),
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path, interval_minutes=60)
        assert "矛盾" in str(excinfo.value)

    def test_nonstandard_interval_requires_confirmation(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "load5.xlsx",
            long_table_rows(datetime(2025, 1, 1), 288, 5, value=1.0),
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path)
        assert "超出 V2.2 一等支持" in str(excinfo.value)
        preview = preview_load_import(path, allow_nonstandard_interval=True)
        assert preview.interval_minutes == 5


# --------------------------------------------------------------------------- #
# ④ 闰年 / 年份 / 时区 / 覆盖率
# --------------------------------------------------------------------------- #
class TestLeapYearAndCalendar:
    def test_quarter_hourly_point_counts(self) -> None:
        assert points_per_year(2024, Resolution.QUARTER_HOURLY) == 35136
        assert points_per_year(2025, Resolution.QUARTER_HOURLY) == 35040
        assert points_per_year(2024, Resolution.HALF_HOURLY) == 17568
        assert points_per_year(2025, Resolution.HALF_HOURLY) == 17520

    def test_half_hourly_axis_has_correct_timestamps(self) -> None:
        axis = build_time_axis(2025, Resolution.HALF_HOURLY)
        assert axis.point_count == 17520
        assert axis.timestamps[1].minute == 30
        assert axis.timestamps[2].hour == 1
        assert axis.delta_hours == 0.5

    def test_leap_year_15min_not_reported_as_mismatch(self, tmp_path: Path) -> None:
        """2024 年 15 分钟应有 35136 点；写满则不得报 V23/V24。"""
        rows = [["时间", "负荷(kWh)"]]
        start = datetime(2024, 1, 1)
        for index in range(35136):
            rows.append([(start + timedelta(minutes=15 * index)).strftime(ISO), 1.0])
        path = write_xlsx(tmp_path / "2024_15min.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.point_count == 35136
        assert preview.coverage_ratio == pytest.approx(1.0)
        assert not any("[V23]" in i.message for i in preview.quality_issues)

    def test_non_leap_point_count_in_leap_year_reports_v23(
        self, tmp_path: Path
    ) -> None:
        """2024 年（闰年）只给 35040 点 = 平年点数 → V23 告警，不得静默裁剪。"""
        rows = [["时间", "负荷(kWh)"]]
        start = datetime(2024, 1, 1)
        for index in range(35040):
            rows.append([(start + timedelta(minutes=15 * index)).strftime(ISO), 1.0])
        path = write_xlsx(tmp_path / "2024_flat15min.xlsx", rows)
        preview = preview_load_import(path)
        assert any("[V23]" in issue.message for issue in preview.quality_issues)
        assert any("闰年" in issue.message for issue in preview.quality_issues)

    def test_year_mismatch_between_file_name_and_data(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "2025年每天的用电量.xlsx",
            long_table_rows(datetime(2019, 1, 1), 24, 60, value=1.0),
        )
        preview = preview_load_import(path)
        assert preview.expected_year == 2025
        assert preview.data_year == 2019
        assert any("[V22]" in issue.message for issue in preview.quality_issues)

    def test_expected_year_from_file_name(self) -> None:
        assert expected_year_from_file_name("2025年每天的用电量.xlsx") == 2025
        assert expected_year_from_file_name("load_2024_15min.csv") == 2024
        assert expected_year_from_file_name("负荷数据.xlsx") is None


class TestTimezone:
    def test_aware_timestamps_are_converted_to_utc8(self, tmp_path: Path) -> None:
        """带 +09:00 的时间戳必须换算为 UTC+8（时间减 1 小时）并留痕（§2.2）。"""
        rows = [["时间", "负荷(kWh)"]]
        for index in range(24):
            stamp = datetime(2025, 1, 1, tzinfo=timezone(timedelta(hours=9))) + timedelta(
                hours=index
            )
            rows.append([stamp.isoformat(), 1.0])
        path = write_xlsx(tmp_path / "tz9.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.points[0].timestamp == datetime(2024, 12, 31, 23, 0)
        assert any("[V21]" in issue.message for issue in preview.quality_issues)
        assert any(
            "UTC+9" in issue.message or "+09" in issue.message
            for issue in preview.quality_issues
        )

    def test_dst_switch_is_detected(self, tmp_path: Path) -> None:
        rows = [["时间", "负荷(kWh)"]]
        for index in range(6):
            offset = (
                timezone(timedelta(hours=1)) if index < 3 else timezone(timedelta(hours=2))
            )
            stamp = datetime(2025, 3, 30, index, tzinfo=offset)
            rows.append([stamp.isoformat(), 1.0])
        path = write_xlsx(tmp_path / "dst.xlsx", rows)
        preview = preview_load_import(path)
        assert any("夏令时" in issue.message for issue in preview.quality_issues)

    def test_dst_switch_detected_from_strings(self) -> None:
        issues = scan_timezone_issues(
            ["2025-03-30T01:00:00+01:00", "2025-03-30T03:00:00+02:00"]
        )
        assert any("夏令时" in issue.message for issue in issues)

    def test_naive_timestamps_get_utc8_note(self) -> None:
        issues = scan_timezone_issues([datetime(2025, 1, 1), datetime(2025, 1, 1, 1)])
        assert len(issues) == 1
        assert "[V21]" in issues[0].message
        assert "Asia/Shanghai" in issues[0].message

    def test_timezone_issue_for_consistent_utc8_aware_data(self) -> None:
        tz8 = timezone(timedelta(hours=8))
        issues = scan_timezone_issues([datetime(2025, 1, 1, tzinfo=tz8)])
        assert len(issues) == 1
        assert "未做换算" in issues[0].message


class TestDriftingTimestamps:
    """复现东风本田逐时文件的真实缺陷：时间戳由公式生成、逐行累计漂移。"""

    def test_sequential_alignment_recovers_clean_hourly_axis(
        self, tmp_path: Path
    ) -> None:
        rows = [["日期和时间", "每日每小时用电量"]]
        stamp = datetime(2019, 1, 1)
        for index in range(48):
            rows.append([stamp, 100.0])
            stamp = stamp + timedelta(seconds=3600.495)
        path = write_xlsx(tmp_path / "drift.xlsx", rows)
        preview = preview_load_import(path)
        assert preview.interval_minutes == 60
        assert preview.alignment_strategy in ("grid", "sequential")
        stamps = [point.timestamp for point in preview.points]
        assert stamps[0] == datetime(2019, 1, 1)
        assert all(
            (b - a) == timedelta(hours=1) for a, b in zip(stamps, stamps[1:])
        )
        assert preview.annual_energy_kwh == pytest.approx(4800.0)

    def test_large_drift_is_reported_not_silent(self, tmp_path: Path) -> None:
        rows = [["日期和时间", "每日每小时用电量"]]
        stamp = datetime(2019, 1, 1)
        for _ in range(300):
            rows.append([stamp, 10.0])
            stamp = stamp + timedelta(seconds=3600.5)
        path = write_xlsx(tmp_path / "drift2.xlsx", rows)
        preview = preview_load_import(path)
        assert any("[V20]" in issue.message for issue in preview.quality_issues)
        assert preview.alignment_strategy == "sequential"
        assert preview.max_offset_seconds > 1.0


# --------------------------------------------------------------------------- #
# ⑤ 宽表（时刻行 × 日期列）
# --------------------------------------------------------------------------- #
class TestWideTableBlocks:
    def test_reads_blocks_and_expands_points(self, tmp_path: Path) -> None:
        dates = [date(2025, 1, day) for day in range(1, 4)]
        path = write_wide_xlsx(
            tmp_path / "wide.xlsx",
            {
                "1月用电数据": wide_block_rows(
                    label=None, dates=dates, minutes=15, value=2.0
                )
            },
        )
        blocks = read_time_blocks(path, "1月用电数据")
        assert len(blocks) == 1
        block = blocks[0]
        assert block.day_count == 3
        assert block.row_count == 96
        assert block.interval_minutes == 15
        assert block.total_energy_kwh == pytest.approx(3 * 96 * 2.0)
        points = build_points_from_block(block)
        assert len(points) == 288
        assert points[0].timestamp == datetime(2025, 1, 1, 0, 0)
        assert points[0].load_kwh == pytest.approx(2.0)
        assert points[-1].timestamp == datetime(2025, 1, 3, 23, 45)

    def test_multiple_blocks_and_catalog(self, tmp_path: Path) -> None:
        dates = [date(2025, 1, day) for day in range(1, 3)]
        path = write_wide_xlsx(
            tmp_path / "wide2.xlsx",
            {
                "1月用电数据": [
                    *wide_block_rows(label=None, dates=dates, minutes=15, value=3.0),
                    *wide_block_rows(label="光伏发电量", dates=dates, minutes=15, value=1.0),
                ]
            },
        )
        catalog = load_block_catalog(path)
        assert catalog["1月用电数据"] == ["1月用电数据#1", "光伏发电量"]
        pv = preview_block_import(
            path, sheet="1月用电数据", block_index=1, all_sheets=False
        )
        assert pv.block_label == "光伏发电量"
        assert pv.value_kind is LoadValueKind.INTERVAL_ENERGY_KWH
        assert pv.annual_energy_kwh == pytest.approx(2 * 96 * 1.0)

    def test_all_sheets_are_concatenated(self, tmp_path: Path) -> None:
        jan = [date(2025, 1, day) for day in range(1, 32)]
        feb = [date(2025, 2, day) for day in range(1, 29)]
        path = write_wide_xlsx(
            tmp_path / "wide3.xlsx",
            {
                "1月用电数据": wide_block_rows(label=None, dates=jan, minutes=15, value=1.0),
                "2月用电数据": wide_block_rows(label=None, dates=feb, minutes=15, value=1.0),
            },
        )
        preview = preview_block_import(path, all_sheets=True)
        assert preview.point_count == (31 + 28) * 96
        stamps = [point.timestamp for point in preview.points]
        assert stamps == sorted(stamps)
        assert len(set(stamps)) == len(stamps)
        assert any("已拼接 2 张工作表" in message for message in preview.messages)

    def test_stale_date_header_is_reported_and_replaced(self, tmp_path: Path) -> None:
        """复现真实缺陷：2 月工作表的光伏块仍写着 1 月的 31 列日期表头（尾部 3 列全 0）。"""
        feb = [date(2025, 2, day) for day in range(1, 29)]
        stale = [date(2025, 1, day) for day in range(1, 32)]
        path = write_wide_xlsx(
            tmp_path / "stale.xlsx",
            {
                "1月用电数据": wide_block_rows(
                    label=None, dates=[date(2025, 1, d) for d in range(1, 32)],
                    minutes=15, value=1.0,
                ),
                "2月用电数据": [
                    *wide_block_rows(label=None, dates=feb, minutes=15, value=1.0),
                    *wide_block_rows(
                        label="光伏发电量", dates=feb, minutes=15, value=0.5,
                        day_header_dates=stale, zero_tail=3,
                    ),
                ],
            },
        )
        preview = preview_block_import(
            path, sheet="2月用电数据", block_index=1, all_sheets=False
        )
        assert preview.point_count == 28 * 96
        assert preview.points[0].timestamp == datetime(2025, 2, 1)
        assert any("[V25]" in issue.message for issue in preview.quality_issues)
        assert any("疑似复制上月表头未更新" in m for m in preview.messages)

    def test_trailing_empty_columns_are_trimmed(self, tmp_path: Path) -> None:
        """真实缺陷：11 月工作表沿用 10 月的 31 列表头，第 31 列整列为空。"""
        nov = [date(2025, 11, day) for day in range(1, 31)]
        rows = wide_block_rows(label=None, dates=nov, minutes=15, value=1.0)
        # 表头多出一列（10 月 31 日），数据列整列为空
        rows[0] = [*rows[0], date(2025, 10, 31).toordinal() - 693594]
        for index in range(1, len(rows)):
            rows[index] = [*rows[index], None]
        path = write_wide_xlsx(tmp_path / "trim.xlsx", {"11月用电数据": rows})
        preview = preview_block_import(
            path, sheet="11月用电数据", all_sheets=False
        )
        assert preview.point_count == 30 * 96
        assert preview.points[0].timestamp == datetime(2025, 11, 1)
        assert preview.points[-1].timestamp == datetime(2025, 11, 30, 23, 45)
        assert any("已裁剪" in message for message in preview.messages)

    def test_month_with_too_many_positive_columns_reports_error(
        self, tmp_path: Path
    ) -> None:
        """2 月出现 31 列且多出的列有正值 → 报中文错误，不猜哪几列属于本月。"""
        stale = [date(2025, 1, day) for day in range(1, 32)]
        path = write_wide_xlsx(
            tmp_path / "conflict_month.xlsx",
            {
                "2月用电数据": [
                    *wide_block_rows(label=None, dates=[date(2025, 2, d) for d in range(1, 29)],
                                     minutes=15, value=1.0),
                    *wide_block_rows(
                        label="光伏发电量", dates=[date(2025, 2, 1)] * 31, minutes=15,
                        value=1.0, day_header_dates=stale,
                    ),
                ]
            },
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_block_import(
                path, sheet="2月用电数据", block_index=1, all_sheets=False
            )
        assert "只有 28 天" in str(excinfo.value)

    def test_block_without_unit_requires_explicit_kind(self, tmp_path: Path) -> None:
        dates = [date(2025, 1, 1)]
        path = write_wide_xlsx(
            tmp_path / "nounit.xlsx",
            {"数据块": wide_block_rows(label="第一块", dates=dates, minutes=60, value=1.0)},
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_block_import(path, sheet="数据块", all_sheets=False)
        assert "无法判定数值是" in str(excinfo.value)
        preview = preview_block_import(
            path,
            sheet="数据块",
            all_sheets=False,
            value_kind=LoadValueKind.POWER_KW,
        )
        # 1 小时粒度：Δt = 1 h，功率与电量数值相同，但口径必须记录为功率
        assert preview.value_kind is LoadValueKind.POWER_KW

    def test_sheet_not_found_reports_chinese_error(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "x.xlsx", long_table_rows(datetime(2025, 1, 1), 24, 60)
        )
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path, sheet="不存在")
        assert "不存在" in str(excinfo.value)

    def test_layout_error_lists_expected_columns(self, tmp_path: Path) -> None:
        path = write_xlsx(tmp_path / "plain.xlsx", [["甲", "乙"], [1, 2]])
        with pytest.raises(ValidationError) as excinfo:
            preview_load_import(path)
        assert "时间列" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# ⑥ 数据集与来源标签（§0.2 红线）
# --------------------------------------------------------------------------- #
class TestLegacyResolutionInference:
    """V2 既有 ``import_load_file`` 的分辨率推断：15 分钟不变，30 分钟修正为 HALF_HOURLY。"""

    def test_15min_still_inferred_as_quarter_hourly(self, tmp_path: Path) -> None:
        from cenep.data import import_load_file

        path = write_xlsx(
            tmp_path / "legacy15.xlsx",
            long_table_rows(datetime(2025, 1, 1), 96, 15, value=1.0),
        )
        profile = import_load_file(path)
        assert profile.resolution is Resolution.QUARTER_HOURLY

    def test_30min_now_inferred_as_half_hourly(self, tmp_path: Path) -> None:
        from cenep.data import import_load_file

        path = write_xlsx(
            tmp_path / "legacy30.xlsx",
            long_table_rows(datetime(2025, 1, 1), 48, 30, value=1.0),
        )
        profile = import_load_file(path)
        assert profile.resolution is Resolution.HALF_HOURLY
        assert profile.resolution.delta_hours == pytest.approx(0.5)


class TestDatasetAndProvenance:
    def _dataset(self, tmp_path: Path) -> HighFrequencyLoadDataset:
        path = write_xlsx(
            tmp_path / "load60_full.xlsx",
            long_table_rows(datetime(2025, 1, 1), 8760, 60, value=10.0),
        )
        return apply_load_import(preview_load_import(path), project_id="P1")

    def test_dataset_basics(self, tmp_path: Path) -> None:
        dataset = self._dataset(tmp_path)
        assert dataset.interval_minutes == 60
        assert dataset.resolution is Resolution.HOURLY
        assert dataset.point_count == 8760
        assert dataset.annual_energy_kwh == pytest.approx(87600.0)
        assert dataset.peak_power_kw == pytest.approx(10.0)
        assert dataset.average_power_kw == pytest.approx(10.0)
        assert dataset.load_factor == pytest.approx(1.0)
        assert dataset.coverage_ratio == pytest.approx(1.0)
        assert provides_measured_curve(dataset)
        assert "实测" in dataset.provenance_text

    def test_energy_and_power_views_are_consistent(self, tmp_path: Path) -> None:
        path = write_xlsx(
            tmp_path / "load15_power.xlsx",
            long_table_rows(
                datetime(2025, 1, 1), 96, 15, header=("时间", "功率(kW)"), value=40.0
            ),
        )
        dataset = apply_load_import(preview_load_import(path), allow_invalid=True)
        assert dataset.value_kind is LoadValueKind.POWER_KW
        assert dataset.values_kw[0] == pytest.approx(40.0)
        assert dataset.interval_energy_kwh[0] == pytest.approx(10.0)
        assert dataset.annual_energy_kwh == pytest.approx(96 * 10.0)
        assert any("E = P × Δt" in note for note in dataset.assumptions)

    def test_estimated_dataset_cannot_claim_measured(self) -> None:
        """§0.2 红线：估算负荷不得伪装成实测高频曲线。"""
        with pytest.raises(ValueError) as excinfo:
            HighFrequencyLoadDataset(
                interval_minutes=15,
                resolution=Resolution.QUARTER_HOURLY,
                source_type=LoadDataSourceType.MONTHLY_BILL_ESTIMATE,
                estimated=False,
            )
        assert "estimated 必须为 True" in str(excinfo.value)

        with pytest.raises(ValueError) as excinfo:
            HighFrequencyLoadDataset(
                interval_minutes=15,
                resolution=Resolution.QUARTER_HOURLY,
                source_type=LoadDataSourceType.HIGH_FREQUENCY_IMPORT,
                estimated=True,
            )
        assert "不得同时标记为估算数据" in str(excinfo.value)

    def test_estimated_dataset_is_labelled_and_scored_lower(self) -> None:
        dataset = HighFrequencyLoadDataset(
            profile_id="est",
            interval_minutes=15,
            resolution=Resolution.QUARTER_HOURLY,
            source_type=LoadDataSourceType.MONTHLY_BILL_ESTIMATE,
            estimated=True,
            points=[
                TimeSeriesPoint(timestamp=datetime(2025, 1, 1, 0, 0), load_kwh=1.0)
            ],
        )
        assert dataset.estimated
        assert not provides_measured_curve(dataset)
        assert "不是实测" in dataset.source_type.report_badge
        assert "估算" in dataset.provenance_text

    def test_resolution_interval_mismatch_is_rejected(self) -> None:
        with pytest.raises(ValueError) as excinfo:
            HighFrequencyLoadDataset(
                interval_minutes=15, resolution=Resolution.HOURLY
            )
        assert "分辨率与间隔必须一致" in str(excinfo.value)

    def test_unsupported_timezone_is_rejected(self) -> None:
        with pytest.raises(ValueError) as excinfo:
            HighFrequencyLoadDataset(
                interval_minutes=60,
                resolution=Resolution.HOURLY,
                timezone="UTC",
            )
        assert "时区只支持" in str(excinfo.value)

    def test_to_load_profile_reuses_v2_persistence(self, tmp_path: Path) -> None:
        dataset = self._dataset(tmp_path)
        profile = dataset.to_load_profile()
        assert profile.resolution is Resolution.HOURLY
        assert profile.annual_energy == pytest.approx(dataset.annual_energy_kwh)
        assert len(profile.points) == dataset.point_count

    def test_resample_after_import_records_assumptions(self, tmp_path: Path) -> None:
        rows = [["时间", "负荷(kWh)"]]
        start = datetime(2025, 1, 1)
        for index in range(96):
            rows.append([(start + timedelta(minutes=15 * index)).strftime(ISO), 10.0])
        path = write_xlsx(tmp_path / "resample.xlsx", rows)
        preview = preview_load_import(path)
        dataset = apply_load_import(preview, resample_to_minutes=60)
        assert dataset.interval_minutes == 60
        assert dataset.point_count == 24
        assert dataset.annual_energy_kwh == pytest.approx(960.0)
        assert any("15 分钟 → 60 分钟" in note for note in dataset.assumptions)


# --------------------------------------------------------------------------- #
# ⑦ 质量检查的确定性
# --------------------------------------------------------------------------- #
class TestQualityDeterminism:
    def test_issue_order_and_text_are_stable(self, tmp_path: Path) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 96, 15, value=1.0)
        rows[5][1] = -1.0
        rows.append([rows[1][0], 1.0])
        path = write_xlsx(tmp_path / "dirty.xlsx", rows)
        first = preview_load_import(path)
        second = preview_load_import(path)
        assert [i.message for i in first.quality_issues] == [
            i.message for i in second.quality_issues
        ]
        assert [i.level for i in first.quality_issues] == [
            i.level for i in second.quality_issues
        ]

    def test_quality_codes_cover_new_rules(self, tmp_path: Path) -> None:
        rows = long_table_rows(datetime(2025, 1, 1), 96, 15, value=1.0)
        path = write_xlsx(tmp_path / "2024_half.xlsx", rows)
        preview = preview_load_import(path, expected_year=2024)
        text = "\n".join(i.message for i in preview.quality_issues)
        assert "[V22]" in text  # 年份不一致
        assert "[V24]" in text  # 覆盖率不足
        assert "[V03]" in text  # 时间点缺失（复用既有校验器）

    def test_check_load_quality_reuses_existing_codes(self) -> None:
        points = [
            TimeSeriesPoint(timestamp=ts, load_kwh=1.0)
            for ts in day_stamps(date(2025, 1, 1), 60)
        ]
        issues = check_load_quality(points, interval_minutes=60)
        text = "\n".join(i.message for i in issues)
        assert "[V24]" in text  # 覆盖率不足（只导入一天）
        assert "[V03]" in text  # 缺失（既有校验器）

    def test_status_from_issues(self) -> None:
        from cenep.data.load_profile_importer import status_from_issues
        from cenep.domain.timeseries_results import DataQualityIssue

        assert status_from_issues([]) is LoadQualityStatus.VALID
        assert (
            status_from_issues([DataQualityIssue(level="INFO", message="x")])
            is LoadQualityStatus.VALID
        )
        assert (
            status_from_issues([DataQualityIssue(level="WARNING", message="x")])
            is LoadQualityStatus.WARNING
        )
        assert (
            status_from_issues([DataQualityIssue(level="ERROR", message="x")])
            is LoadQualityStatus.INVALID
        )
