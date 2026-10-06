"""方案比较与寻优写入唯一结果对象的接线测试（V2 §43–§48、§62、§86、§105）。

背景（交付缺口）
----------------
规范 §62 要求 :class:`~cenep.domain.results.CalculationResult` 携带
``scenario_results`` 与 ``optimization_results``；Phase 8 交付的
:mod:`cenep.calculation.scenario_engine` 与 :mod:`cenep.optimization` 只有函数入口，
没有任何地方把结果写进唯一结果对象，于是 GUI 的「方案比较」图
（:class:`~cenep.ui.charts.ScenarioBarChart`）与 Excel 的「方案比较」「方案寻优」
两张表永远显示"未执行"。本文件锁定修复后的接线契约：

* §43/§44：时序启用时**默认**产出三种标准组合（仅光伏 / 仅储能 / 光伏+储能）；
* §105：「当前方案」一条的 IRR / NPV / LCOE / LCOS 与 ``result`` 上的同名字段**逐位一致**
  （它不是近似值，而是直接取自本次唯一计算结果）；
* §45–§48：寻优**默认不执行**（开销大），仅当 ``timeseries.optimization_enabled``
  为真时执行，且 ``best_run_id`` 指向的候选确为全部可行候选中的最优；
* §1.1：``timeseries.enabled == False`` 的 V1 项目完全不产生 V2 结果，
  ``result.scenarios`` / ``result.sensitivity`` 行为不变；
* §86：方案比较 / 寻优失败时只记录 warning + 中文说明，绝不让主计算失败。

测试统一用 :func:`test_scenario_engine.make_project` 在黄金案例（§85）上补全时序配置，
并把计算期压到 10 年以控制测试时长（精确复核的代价与时长相乘）。
"""

from __future__ import annotations

import logging

import pytest

from cenep import optimization as optimization_package
from cenep.application.calculation_service import (
    CURRENT_SCENARIO_LABEL,
    CURRENT_SCENARIO_NAME,
    OPTIMIZATION_SKIPPED_KEYWORD,
    SCENARIO_SKIPPED_KEYWORD,
    CalculationService,
)
from cenep.calculation import scenario_engine as se
from cenep.calculation.engine import calculation_engine
from cenep.domain.enums import ProjectType
from cenep.domain.models import Project
from cenep.domain.results import CalculationResult

from test_scenario_engine import make_project  # noqa: E402  （pytest 会把 tests/ 加入 sys.path）


def _has_chinese(text: str) -> bool:
    """文本是否含中文（规范 §132：面向用户的说明一律中文）。"""
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


@pytest.fixture(scope="module")
def golden_ts_project() -> Project:
    """V2 黄金案例（§85 光储项目）+ 有效时序配置（10 年，控制测试时长）。"""
    from conftest import _base_project  # 黄金案例项目工厂（tests/conftest.py）

    return make_project(_base_project(ProjectType.PV_STORAGE), period=10)


@pytest.fixture(scope="module")
def default_result(golden_ts_project: Project) -> CalculationResult:
    """默认配置（未开启寻优）下的计算结果，供多个用例共享（只算一次）。"""
    return CalculationService().calculate(golden_ts_project)


# --------------------------------------------------------------------------- #
# ① 默认产出方案比较（V2 §43、§44、§62）
# --------------------------------------------------------------------------- #
class TestScenarioComparisonWiring:
    def test_default_result_carries_standard_variants(self, default_result):
        """时序启用的项目算完后必须自带三种标准组合（≥3 条）。"""
        scenarios = default_result.scenario_results
        assert len(scenarios) >= 3, "时序启用时方案比较必须默认产出，不得留空"

        labels = "、".join(item.label for item in scenarios)
        for keyword in ("仅光伏", "仅储能", "光伏+储能"):
            assert keyword in labels, f"方案比较缺少「{keyword}」组合：{labels}"

        for item in scenarios:
            assert item.label, "每条方案都必须有中文标签"
            assert isinstance(item.total_capex, float)
            assert item.project_irr is not None, f"{item.label} 缺少项目 IRR"
            assert isinstance(item.project_npv, float)

    def test_current_scenario_matches_main_result_bitwise(self, default_result):
        """§105 口径一致：与当前项目同配置的那一条必须与主结果**逐位一致**。"""
        current = next(
            item
            for item in default_result.scenario_results
            if item.name == CURRENT_SCENARIO_NAME
        )
        assert current.label == CURRENT_SCENARIO_LABEL
        assert current.is_baseline is False
        # 逐位一致（不是近似）：以下断言一律用 == 而不是 pytest.approx
        assert current.project_irr == default_result.project_irr
        assert current.project_npv == default_result.project_npv
        assert current.equity_irr == default_result.equity_irr
        assert current.static_payback == default_result.static_payback
        assert current.discounted_payback == default_result.discounted_payback
        assert current.lcoe == default_result.lcoe
        assert current.lcos == default_result.lcos
        assert current.total_capex == default_result.total_capex
        assert current.pv_capacity_kwp == default_result.pv_capacity_kwp
        assert current.storage_energy_kwh == default_result.storage_energy_kwh

    def test_excel_and_gui_consumers_see_scenarios(self, default_result):
        """GUI 图表与 Excel 两张表的数据源非空（缺口描述里的"永远未执行"）。"""
        assert default_result.scenario_results, "GUI「方案比较」图的数据源不得为空"
        # 方案比较的每一条都要能算出柱状图所需的指标（NPV / IRR / 回收期）
        assert any(item.project_npv for item in default_result.scenario_results)

    def test_pv_only_project_still_gets_comparison(self):
        """仅光伏项目：「仅储能」组合退化为 §42 基准方案，必须跳过而非整体失败。"""
        from conftest import _base_project

        project = make_project(_base_project(ProjectType.COMMERCIAL_PV), period=10)
        project.storage.storage_energy_kwh = 0.0
        project.storage.storage_power_kw = 0.0

        result = CalculationService().calculate(project)
        assert result.project_irr is not None
        assert len(result.scenario_results) >= 3, "仅光伏项目也必须给出方案比较"
        notes = "\n".join(result.notes)
        assert "未参与比较" in notes
        assert "仅储能" in notes and "基准方案" in notes


# --------------------------------------------------------------------------- #
# ② notes 披露：做了哪些比较、近似口径、寻优如何开启（V2 §86、§105）
# --------------------------------------------------------------------------- #
class TestNotesDisclosure:
    def test_notes_explain_comparison_and_approximation(self, default_result):
        text = "\n".join(default_result.notes)
        assert _has_chinese(text)
        assert "方案比较已执行" in text
        for keyword in ("仅光伏", "仅储能", "光伏+储能"):
            assert keyword in text, f"notes 未说明做了哪种方案比较：缺少「{keyword}」"
        assert "线性外推" in text, "notes 必须说明排名用的是首年 + 线性外推的近似"
        assert "首年" in text
        assert "近似" in text
        assert "精确复核" in text, "notes 必须说明最优候选是精确复核（§48）"
        assert "逐位一致" in text, "notes 必须说明「当前方案」与主结果逐位一致（§105）"

    def test_notes_explain_how_to_enable_optimization(self, default_result):
        text = "\n".join(default_result.notes)
        assert OPTIMIZATION_SKIPPED_KEYWORD in text
        assert "timeseries.optimization_enabled" in text, "notes 必须给出开启寻优的配置项"
        assert "启用方案寻优" in text, "notes 必须给出界面上的开启位置"
        assert _has_chinese(text)


# --------------------------------------------------------------------------- #
# ③ 寻优默认不跑（V2 §45–§48、§86：开销大，必须显式开启）
# --------------------------------------------------------------------------- #
class TestOptimizationDisabledByDefault:
    def test_default_does_not_run_optimization(self, default_result):
        assert default_result.optimization_results is None
        # 寻优开关不影响方案比较：比较仍须产出，GUI / Excel 不得回到"未执行"
        assert default_result.scenario_results

    def test_new_field_defaults_to_false(self):
        """开关字段本身必须是 False（不得擅自默认开启，V2 §86 性能红线）。"""
        assert Project().timeseries.optimization_enabled is False


# --------------------------------------------------------------------------- #
# ④ 开启后能跑通，且最优候选确为最优（V2 §45–§48）
# --------------------------------------------------------------------------- #
class TestOptimizationWhenEnabled:
    def test_enabled_optimization_produces_best_candidate(self, golden_ts_project):
        project = golden_ts_project.model_copy(deep=True)
        project.timeseries.optimization_enabled = True

        result = CalculationService().calculate(project)
        optimization = result.optimization_results
        assert optimization is not None
        assert optimization.candidates, "寻优必须给出全部候选（§48 禁止黑盒）"
        assert optimization.best_run_id
        assert optimization.best_scenario is not None

        best = next(
            item for item in optimization.candidates if item.run_id == optimization.best_run_id
        )
        feasible = [item for item in optimization.candidates if item.feasible]
        assert feasible, "至少要有一个可行候选"
        assert best.objective_value == max(item.objective_value for item in feasible), (
            "best_run_id 指向的候选其目标函数值必须确为全部可行候选中的最优"
        )
        assert any(se.VERIFY_NOTE in item.note for item in optimization.candidates), (
            "最优候选必须做精确复核（§48、§86 两阶段评估）"
        )
        notes = "\n".join(result.notes)
        assert "方案寻优已执行" in notes
        assert _has_chinese(notes)

    def test_excel_sheets_no_longer_say_not_executed(self, golden_ts_project, tmp_path):
        """端到端：Excel「方案比较」「方案寻优」两张表不再输出"未执行"占位说明。"""
        from openpyxl import load_workbook

        from cenep.reports.excel_exporter import ExcelExporter

        project = golden_ts_project.model_copy(deep=True)
        project.timeseries.optimization_enabled = True
        result = CalculationService().calculate(project)

        path = ExcelExporter().export(project, result, tmp_path / "wiring.xlsx")
        workbook = load_workbook(path)

        def _text(sheet_name: str) -> str:
            return "\n".join(
                str(cell.value)
                for row in workbook[sheet_name].iter_rows()
                for cell in row
                if cell.value is not None
            )

        compare = _text("方案比较")
        assert "未执行方案比较" not in compare
        assert CURRENT_SCENARIO_LABEL in compare
        assert "仅光伏" in compare and "仅储能" in compare

        optimization_sheet = _text("方案寻优")
        assert "未执行方案寻优" not in optimization_sheet
        assert result.optimization_results is not None
        assert result.optimization_results.best_run_id in optimization_sheet


# --------------------------------------------------------------------------- #
# ⑤ V1 不受影响（V2 §1.1 兼容性红线）
# --------------------------------------------------------------------------- #
class TestV1Unaffected:
    def test_v1_project_keeps_v2_fields_empty(self, golden_pv_storage):
        project = golden_pv_storage.model_copy(deep=True)
        assert project.timeseries.enabled is False

        result = CalculationService().calculate(project)
        assert result.scenario_results == [], "未启用时序时不得执行 V2 方案比较"
        assert result.optimization_results is None, "未启用时序时不得执行 V2 寻优"
        assert result.time_series_results is None
        assert result.baseline_results is None
        assert result.energy_balance is None
        assert "方案比较" not in "\n".join(result.notes)

    def test_v1_scenarios_and_sensitivity_unchanged(self, golden_pv_storage):
        """V1 的 ``result.scenarios`` / ``result.sensitivity`` 行为逐位不变。"""
        result = CalculationService().calculate(golden_pv_storage.model_copy(deep=True))
        baseline = calculation_engine.calculate(golden_pv_storage.model_copy(deep=True))

        assert len(result.scenarios) == 3
        assert [item.model_dump() for item in result.scenarios] == [
            item.model_dump() for item in baseline.scenarios
        ]
        assert [item.model_dump() for item in result.sensitivity] == [
            item.model_dump() for item in baseline.sensitivity
        ]


# --------------------------------------------------------------------------- #
# ⑥ 失败不炸：方案比较 / 寻优出错只记录 warning + 中文说明（V2 §86）
# --------------------------------------------------------------------------- #
class TestFailureIsolation:
    def test_comparison_failure_is_swallowed(self, golden_ts_project, monkeypatch, caplog):
        def _boom(*args, **kwargs):  # pragma: no cover - 只用于制造失败
            raise RuntimeError("模拟方案比较失败（接线测试）")

        monkeypatch.setattr(se, "compare_scenarios", _boom)
        with caplog.at_level(logging.WARNING, logger="cenep"):
            result = CalculationService().calculate(golden_ts_project)

        # 主计算必须照常返回
        assert result.project_irr is not None
        assert result.project_npv != 0.0
        assert result.scenario_results == []
        assert result.optimization_results is None

        notes = "\n".join(result.notes)
        assert SCENARIO_SKIPPED_KEYWORD in notes
        assert "模拟方案比较失败" in notes
        assert _has_chinese(notes)

        warnings = [record for record in caplog.records if record.levelno >= logging.WARNING]
        assert any("方案比较" in record.getMessage() for record in warnings), (
            "失败必须记录 warning 日志（不得静默）"
        )

    def test_optimization_failure_is_swallowed(self, golden_ts_project, monkeypatch, caplog):
        project = golden_ts_project.model_copy(deep=True)
        project.timeseries.optimization_enabled = True

        def _boom(*args, **kwargs):  # pragma: no cover - 只用于制造失败
            raise RuntimeError("模拟寻优失败（接线测试）")

        monkeypatch.setattr(optimization_package, "optimize", _boom)
        with caplog.at_level(logging.WARNING, logger="cenep"):
            result = CalculationService().calculate(project)

        assert result.optimization_results is None
        assert result.scenario_results, "寻优失败不得影响方案比较"
        notes = "\n".join(result.notes)
        assert OPTIMIZATION_SKIPPED_KEYWORD in notes
        assert "模拟寻优失败" in notes
        warnings = [record for record in caplog.records if record.levelno >= logging.WARNING]
        assert any("方案寻优" in record.getMessage() for record in warnings)
