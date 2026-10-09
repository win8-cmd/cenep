"""V2 报表升级测试（V2 §66、§67、§105、§108；V2.1 §8.1；V2.2 §6.3；V2.4 §8.1）。

覆盖三件事：

1. **Excel 表清单**（§67 + V2.1 §8.1 + V2.2 §6.3 + V2.4 §8.1）：
   V1 的 13 张全部保留 + V2 新增 11 张 + V2.1 新增 2 张 + V2.2 新增 1 张
   + V2.4 新增 5 张，顺序固定；表数**一律引用 ``len(SHEET_NAMES)``**，不硬编码；
2. **PDF 章节**（§66 + V2.1 §8.1 + V2.2 §6.3 + V2.4 §8.1）：含 6 个时序章节、账单与校准章节、
   负荷与消纳章节、质量等级章节与 3 张图表；章节号由 ``section_title`` 生成，不硬编码；
3. **结果一致性**（§105）：Excel / PDF 展示的数字必须来自同一个 ``CalculationResult``，
   且两者都包含 V2 新增的关键指标。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab.graphics.shapes import Drawing
from reportlab.platypus import Paragraph, Table

from cenep.application.calculation_service import CalculationService
from cenep.application.project_service import ProjectService
from cenep.calculation.engine import calculation_engine
from cenep.reports.excel_exporter import SHEET_NAMES, ExcelExporter
from cenep.reports.pdf_exporter import (
    DISCLAIMER,
    REPORT_SECTIONS,
    PdfExporter,
    _styles,
    register_cjk_font,
    section_title,
)

pytest.importorskip("test_v2_integration", reason="需要 V2 Golden Case 构造器")
from test_v2_integration import build_project  # noqa: E402

V2_SHEETS = (
    "储能与调度", "负荷曲线", "光伏曲线", "分时电价", "8760时序仿真",
    "能量平衡", "年度汇总", "收益分解", "方案比较", "方案寻优", "数据质量",
)
#: V2.1 §8.1（阶段 2）新增的两张账单表
V21_SHEETS = ("账单原始数据", "账单校验")
#: V2.4 §8.1（阶段 7）新增的五张表
V24_SHEETS = ("月度电费分析", "负荷数据质量", "电价版本与来源", "方案电费对比", "计算假设与警告")
V1_SHEETS = (
    "项目概况", "基础参数", "技术参数", "电价参数", "投资参数", "运维参数",
    "融资参数", "年度现金流", "财务指标", "敏感性分析", "情景分析",
    "政策依据", "参数来源",
)
V2_PDF_SECTIONS = ("负荷分析", "PV时序分析", "储能SOC分析", "能源流", "电费分析", "储能收益")
#: V2.4 §8.1 新增的 6 个章节
V24_PDF_SECTIONS = (
    "数据来源与质量等级", "基准账单校准", "月度电费分析",
    "方案前后电费差额", "经济性指标与风险提示", "关键假设、未建模项与免责声明",
)


def _service() -> CalculationService:
    return CalculationService(project_service=ProjectService(autosave_enabled=False))


@pytest.fixture(scope="module")
def v2_case():
    """V2 Golden Case（规范 §85）与其计算结果。"""
    project = build_project()
    return project, _service().calculate(project)


@pytest.fixture
def v1_case(golden_pv_storage):
    """V1 项目（未启用时序）与其计算结果（函数作用域，因 golden fixture 为函数作用域）。"""
    return golden_pv_storage, calculation_engine.calculate(golden_pv_storage)


def _export_excel(project, result, path: Path):
    return load_workbook(ExcelExporter().export(project, result, path))


def _sheet_text(ws) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def _numbers(ws) -> list[float]:
    return [
        float(c.value)
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
    ]


def _collect_text(flowables) -> str:
    parts: list[str] = []
    for item in flowables:
        if isinstance(item, Paragraph):
            parts.append(item.text)
        elif isinstance(item, Table):
            for row in item._cellvalues:
                for cell in row:
                    parts.append(cell if isinstance(cell, str) else _collect_text([cell]))
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# 任务 A：Excel（V2 §67）
# --------------------------------------------------------------------------- #
class TestExcelSheetSet:
    def test_sheet_count_follows_declaration(self):
        """V2.4 §8.1：表数**不得硬编码**，一律引用 ``SHEET_NAMES``。

        历史：V2 的 24 张 + V2.1 账单 2 张 = 26；V2.2 §6.3 加「消纳率分析」= 27；
        V2.4 §8.1 再加 5 张 = 32。
        """
        assert len(SHEET_NAMES) == 32
        assert list(SHEET_NAMES[-len(V24_SHEETS):]) == list(V24_SHEETS)

    def test_v2_case_sheet_order(self, v2_case, tmp_path):
        project, result = v2_case
        wb = _export_excel(project, result, tmp_path / "v2.xlsx")
        assert wb.sheetnames == SHEET_NAMES

    def test_v1_case_sheet_order(self, v1_case, tmp_path):
        """V1 项目也必须生成全部工作表（新增表输出占位说明）。"""
        project, result = v1_case
        wb = _export_excel(project, result, tmp_path / "v1.xlsx")
        assert wb.sheetnames == SHEET_NAMES

    def test_all_required_sheets_present(self, v2_case, tmp_path):
        project, result = v2_case
        names = set(_export_excel(project, result, tmp_path / "v2.xlsx").sheetnames)
        for name in (*V1_SHEETS, *V2_SHEETS, *V21_SHEETS, *V24_SHEETS):
            assert name in names, f"缺少工作表：{name}"

    @pytest.mark.parametrize("sheet", V21_SHEETS)
    def test_no_bill_project_gets_placeholder_not_error(self, v1_case, tmp_path, sheet):
        """V2.1 §8.1/§8.2：无账单时两张账单表照常生成，写明未录入与录入方法，不缺表不报错。"""
        project, result = v1_case
        wb = _export_excel(project, result, tmp_path / "v1.xlsx")
        text = _sheet_text(wb[sheet])
        assert text.strip(), f"{sheet} 内容为空"
        assert "本项目尚未录入电费账单" in text
        assert "导入" in text and "录入方法" in text

    @pytest.mark.parametrize("sheet", V24_SHEETS)
    def test_v24_sheets_never_missing(self, v1_case, tmp_path, sheet):
        """V2.4 §8.1/§8.2：5 张新表在 V1 项目（无账单、无负荷、无场景）下也必须存在且有说明。"""
        project, result = v1_case
        wb = _export_excel(project, result, tmp_path / "v1.xlsx")
        text = _sheet_text(wb[sheet])
        assert text.strip(), f"{sheet} 内容为空"
        assert any(
            key in text
            for key in ("尚未录入电费账单", "尚未导入或估算", "没有可用的电价计划", "尚未执行光储四场景",
                        "内置 / 已登记计划")
        ), f"{sheet} 缺少占位说明"

    def test_v24_sheets_are_appended_after_disclaimer_related_sheets(self, v1_case, tmp_path):
        """V2.4 表是**追加式**扩展：既有表的相对顺序一字未动（V1 §108 兼容承诺）。"""
        project, result = v1_case
        names = _export_excel(project, result, tmp_path / "v1.xlsx").sheetnames
        base = list(SHEET_NAMES[: len(SHEET_NAMES) - len(V24_SHEETS)])
        assert names[: len(base)] == base
        for legacy in ("政策依据", "参数来源", "账单原始数据", "消纳率分析"):
            assert base.index(legacy) < len(base)

    @pytest.mark.parametrize(
        "sheet",
        ("8760时序仿真", "负荷曲线", "光伏曲线", "能量平衡", "收益分解",
         "方案比较", "方案寻优", "数据质量"),
    )
    def test_v1_project_gets_placeholder_not_error(self, v1_case, tmp_path, sheet):
        """需要时序数据的 V2 新表，在 V1 项目中必须输出说明而不是缺表或报错。"""
        project, result = v1_case
        wb = _export_excel(project, result, tmp_path / "v1.xlsx")
        text = _sheet_text(wb[sheet])
        assert text.strip(), f"{sheet} 内容为空"
        assert any(k in text for k in ("未启用", "未执行", "不适用", "未配置")), (
            f"{sheet} 缺少占位说明"
        )

    @pytest.mark.parametrize("sheet", ("储能与调度", "分时电价", "年度汇总"))
    def test_config_and_annual_sheets_have_real_content_for_v1(self, v1_case, tmp_path, sheet):
        """配置类与年度汇总表不依赖时序数据，V1 项目下也必须有真实内容。"""
        project, result = v1_case
        wb = _export_excel(project, result, tmp_path / "v1.xlsx")
        assert len(_numbers(wb[sheet])) > 3, f"{sheet} 在 V1 项目下内容不足"

    @pytest.mark.parametrize("sheet", ("8760时序仿真", "负荷曲线", "光伏曲线", "能量平衡",
                                       "年度汇总", "收益分解", "分时电价", "储能与调度"))
    def test_v2_project_has_real_content(self, v2_case, tmp_path, sheet):
        project, result = v2_case
        wb = _export_excel(project, result, tmp_path / "v2.xlsx")
        assert len(_numbers(wb[sheet])) > 5, f"{sheet} 未写出实际数值"


class TestHourlySheetSampling:
    def test_sampled_not_full_8760(self, v2_case, tmp_path):
        """关键约束：不得把 8760 行全写进 Excel（否则工作簿过大）。"""
        project, result = v2_case
        ws = _export_excel(project, result, tmp_path / "v2.xlsx")["8760时序仿真"]
        stamps = [
            r[0].value
            for r in ws.iter_rows()
            if isinstance(r[0].value, str) and "-" in str(r[0].value) and ":" in str(r[0].value)
        ]
        # 每月 1 日 24 小时 = 288，加 1/15 与 7/15 各 24 = 336
        assert len(stamps) == 336, f"抽样行数异常：{len(stamps)}"
        assert len(stamps) < 8760

    def test_explains_where_full_data_lives(self, v2_case, tmp_path):
        project, result = v2_case
        ws = _export_excel(project, result, tmp_path / "v2.xlsx")["8760时序仿真"]
        text = _sheet_text(ws)
        assert "8760" in text
        assert ".nep" in text or "JSON" in text

    def test_hourly_values_match_result(self, v2_case, tmp_path):
        """§105：抽样行的数值必须与 CalculationResult 的列式数据逐位一致。"""
        project, result = v2_case
        rs = result.time_series_results.hourly
        ws = _export_excel(project, result, tmp_path / "v2.xlsx")["8760时序仿真"]
        checked = 0
        for row in ws.iter_rows():
            stamp = row[0].value
            if not (isinstance(stamp, str) and ":" in stamp):
                continue
            # 该行的第 2 列是负荷；用时间戳在列式结果里找同一小时比对
            ts_text = stamp.replace(" ", "T") + ":00"
            match = next(
                (i for i, t in enumerate(rs.timestamps) if t.isoformat() == ts_text), None
            )
            if match is None:
                continue
            assert float(row[1].value) == pytest.approx(rs.column("load")[match], rel=1e-12)
            checked += 1
            if checked >= 5:
                break
        assert checked >= 5, "未能匹配到抽样行"


class TestEnergyBalanceSheet:
    def test_balanced(self, v2_case, tmp_path):
        project, result = v2_case
        ws = _export_excel(project, result, tmp_path / "v2.xlsx")["能量平衡"]
        text = _sheet_text(ws)
        assert "是" in text
        assert "1.0e-06" in text or "0.000001" in text or "1e-06" in text

    def test_error_matches_result(self, v2_case, tmp_path):
        project, result = v2_case
        values = _numbers(_export_excel(project, result, tmp_path / "v2.xlsx")["能量平衡"])
        assert any(v == pytest.approx(result.energy_balance.error, abs=1e-15) for v in values)


class TestNoFormulasInV2Sheets:
    def test_no_formulas(self, v2_case, tmp_path):
        """V1 §109 / V2 §61：报表层不得写公式。"""
        project, result = v2_case
        wb = _export_excel(project, result, tmp_path / "v2.xlsx")
        offenders = [
            f"{ws.title}!{c.coordinate}"
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert offenders == []


# --------------------------------------------------------------------------- #
# 任务 B：PDF（V2 §66）
# --------------------------------------------------------------------------- #
class TestPdfSections:
    def test_v2_project_all_sections(self, v2_case):
        project, result = v2_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        text = _collect_text(story)
        for section in REPORT_SECTIONS:
            assert section in text, f"报告缺少章节：{section}"

    def test_v1_project_all_sections(self, v1_case):
        project, result = v1_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        text = _collect_text(story)
        for name in REPORT_SECTIONS:
            assert name in text, f"V1 报告缺少章节：{name}"

    def test_section_order_follows_spec(self, v2_case):
        """§66 + V2.4 §8.1：章节必须按 ``REPORT_SECTIONS`` 的顺序出现（序号由 section_title 生成）。"""
        project, result = v2_case
        text = _collect_text(
            PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        )
        positions = []
        for index in range(len(REPORT_SECTIONS)):
            title = section_title(index)
            idx = text.find(title)
            assert idx >= 0, f"未找到章节 {title}"
            positions.append(idx)
        assert positions == sorted(positions), "章节顺序与 REPORT_SECTIONS 不一致"

    def test_v24_section_titles_are_the_declared_ones(self, v2_case):
        project, result = v2_case
        text = _collect_text(
            PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        )
        for name in V24_PDF_SECTIONS:
            index = REPORT_SECTIONS.index(name)
            assert section_title(index) in text

    def test_v1_project_notes_timeseries_disabled(self, v1_case):
        project, result = v1_case
        text = _collect_text(
            PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        )
        assert "未启用时序仿真" in text


class TestPdfCharts:
    def test_three_charts_for_v2(self, v2_case):
        """§66 要求的新增图表：典型日曲线、逐月电量、SOC 曲线。"""
        project, result = v2_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        drawings = [f for f in story if isinstance(f, Drawing)]
        assert len(drawings) == 3, f"图表数量异常：{len(drawings)}"

    def test_charts_placed_in_expected_sections(self, v2_case):
        project, result = v2_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        current = ""
        placed: dict[str, int] = {}
        for item in story:
            if isinstance(item, Paragraph) and item.text in {
                section_title(index) for index in range(len(REPORT_SECTIONS))
            }:
                current = item.text
            elif isinstance(item, Drawing):
                placed[current] = placed.get(current, 0) + 1
        # V2.1 §8.1 插入账单章节、V2.2 §6.3 插入负荷估算与消纳章节、V2.4 §8.1 再插入 6 章，
        # 时序章节整体顺延；这里按声明顺序取章节标题，不硬编码中文序号。
        for name in ("负荷分析", "PV时序分析", "储能SOC分析"):
            title = section_title(REPORT_SECTIONS.index(name))
            assert placed.get(title, 0) >= 1, f"{title} 缺少图表"

    def test_no_charts_when_timeseries_disabled(self, v1_case):
        project, result = v1_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        assert [f for f in story if isinstance(f, Drawing)] == []

    def test_pdf_renders_and_stays_reasonable(self, v2_case, tmp_path):
        project, result = v2_case
        path = PdfExporter().export(project, result, tmp_path / "v2报告")
        raw = path.read_bytes()
        assert raw.startswith(b"%PDF")
        assert b"%%EOF" in raw[-2048:]
        assert raw.count(b"/Type /Page") >= 6
        # 抽样绘图，体积不应失控
        assert len(raw) < 2_000_000


# --------------------------------------------------------------------------- #
# 任务 C：结果一致性（V2 §105）
# --------------------------------------------------------------------------- #
class TestResultConsistency:
    def test_excel_metrics_match_result(self, v2_case, tmp_path):
        """Excel「财务指标」的核心数字必须等于 result 的同名字段。

        该表对 IRR / 回收期 / LCOE / DSCR 按**显示精度**写成文本
        （``_fmt(value, 6)`` 等，与 V1 行为一致），NPV、总投资等写原始数值，
        因此两类字段分别用文本匹配与数值匹配断言（V2 §105 允许显示四舍五入）。
        """
        project, result = v2_case
        ws = _export_excel(project, result, tmp_path / "v2.xlsx")["财务指标"]
        text = _sheet_text(ws)
        assert f"{result.project_irr:.6f}" in text
        assert f"{result.static_payback:.2f}" in text
        assert f"{result.lcoe:.4f}" in text
        values = _numbers(ws)
        for expected in (result.project_npv, result.total_capex, result.first_year_generation):
            assert any(v == pytest.approx(expected, rel=1e-12) for v in values), (
                f"Excel 财务指标缺少 / 不一致：{expected}"
            )

    def test_excel_revenue_matches_result_metrics(self, v2_case, tmp_path):
        project, result = v2_case
        m = result.time_series_results.metrics
        values = _numbers(_export_excel(project, result, tmp_path / "v2.xlsx")["收益分解"])
        for expected in (
            m.electricity_cost_saving,
            m.pv_self_consumption_saving,
            m.storage_arbitrage_revenue,
            m.self_consumption_rate,
            m.self_sufficiency_rate,
            m.equivalent_cycles,
            m.demand_saving,
        ):
            assert any(v == pytest.approx(expected, rel=1e-12) for v in values), (
                f"Excel 收益分解缺少 / 不一致：{expected}"
            )

    def test_both_excel_and_pdf_expose_v2_indicators(self, v2_case, tmp_path):
        """§105：Excel 与 PDF 必须同时包含 V2 新增的关键指标。"""
        project, result = v2_case
        m = result.time_series_results.metrics

        wb = _export_excel(project, result, tmp_path / "v2.xlsx")
        excel_text = "\n".join(
            _sheet_text(wb[name])
            for name in ("收益分解", "负荷曲线", "光伏曲线", "储能与调度", "年度汇总")
        )
        pdf_text = _collect_text(
            PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        )

        for label in ("自用率", "自给率", "等效循环", "需量"):
            assert label in excel_text, f"Excel 缺少指标标签：{label}"
            assert label in pdf_text, f"PDF 缺少指标标签：{label}"

        # 数值层面：PDF 用百分比格式呈现自用率与自给率
        assert f"{m.self_consumption_rate:.2%}" in pdf_text
        assert f"{m.self_sufficiency_rate:.2%}" in pdf_text
        # Excel 以原始数值写入
        values = _numbers(wb["收益分解"])
        assert any(v == pytest.approx(m.self_consumption_rate, rel=1e-12) for v in values)
        assert any(v == pytest.approx(m.self_sufficiency_rate, rel=1e-12) for v in values)

    def test_disclaimer_appears_twice_in_pdf(self, v2_case):
        project, result = v2_case
        story = PdfExporter().build_story(project, result, _styles(register_cjk_font()))
        texts = [f.text for f in story if isinstance(f, Paragraph)]
        assert sum(1 for t in texts if DISCLAIMER in t) == 2
