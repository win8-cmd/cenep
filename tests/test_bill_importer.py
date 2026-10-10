"""账单 Excel 模板 / 列映射 / 导入预览测试（V2.1 §5.3、§5.5、§9.1）。

覆盖规格书 §9.1 的账单导入矩阵：正常单月账单、缺少分项费用、总分项一致/不一致、
负值、空值、非法日期、跨月账期、重复账单、中英文列名、列顺序变化、空工作表、
错误工作表、损坏文件、重复导入，以及"示例行不得被默认当作真实数据导入"。
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cenep.calculation.errors import ValidationError
from cenep.data.bill_importer import (
    BILL_COLUMNS,
    EXAMPLE_MARKER,
    MARKER_COLUMN,
    SHEET_BILLS,
    SHEET_DICT,
    SHEET_HELP,
    SHEET_TARIFF,
    TEMPLATE_FILE_NAME,
    apply_bill_import,
    bill_template_bytes,
    build_bill_template,
    preview_bill_import,
    resolve_bill_columns,
)
from cenep.domain.bill_models import ElectricityBill, make_bill_id
from cenep.domain.enums import BillQualityStatus, DuplicateStrategy

HEADERS: list[str] = [column.header for column in BILL_COLUMNS]


def _row(marker: str = "", **fields) -> list:
    """按模板列顺序生成一行（未给出的字段留空）。"""
    values = {"": marker, **fields}
    return [values.get(column.field, "") for column in BILL_COLUMNS]


def _write_xlsx(path: Path, rows: list[list], *, sheet: str = SHEET_BILLS, sheets=None) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    for name, extra_rows in (sheets or {}).items():
        extra = workbook.create_sheet(name)
        for row in extra_rows:
            extra.append(row)
    workbook.save(path)
    return path


def _valid_pair() -> list[list]:
    """表头 + 一条完全自洽的账单。"""
    return [
        HEADERS,
        _row(
            账期起_=None,
            billing_period_start="2026-01-01",
            billing_period_end="2026-01-31",
            meter_id="M001",
            voltage_level="10kV",
            tariff_structure="单一制",
            energy_total_kwh=100_000,
            energy_sharp_kwh=10_000,
            energy_peak_kwh=30_000,
            energy_flat_kwh=40_000,
            energy_valley_kwh=20_000,
            energy_charge_yuan=65_000,
            bill_total_yuan=65_000,
        ),
    ]


class TestTemplate:
    """§5.3：模板必须含四张工作表、稳定表头、示例行与数据字典。"""

    def _template(self, tmp_path: Path) -> Path:
        return build_bill_template(tmp_path / "模板")

    def test_sheet_names_and_order(self, tmp_path):
        workbook = load_workbook(self._template(tmp_path))
        assert workbook.sheetnames == [SHEET_HELP, SHEET_BILLS, SHEET_TARIFF, SHEET_DICT]

    def test_default_file_name(self, tmp_path):
        target = build_bill_template(tmp_path)
        assert target.name == TEMPLATE_FILE_NAME

    def test_month_sheet_headers_are_stable(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_BILLS]
        actual = [cell.value for cell in worksheet[1]]
        # 模板第 1 列是"数据标记"，其余列与列定义一一对应
        assert actual == HEADERS

    def test_required_columns_present(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_BILLS]
        headers = {str(cell.value) for cell in worksheet[1]}
        for keyword in ("账期起", "账期止", "总购电量", "尖峰电量", "深谷电量", "账单总额", "账单计费需量"):
            assert any(keyword in header for header in headers), f"模板缺少列：{keyword}"

    def test_example_rows_are_marked(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_BILLS]
        markers = [worksheet.cell(row=row, column=1).value for row in (2, 3)]
        assert markers == [EXAMPLE_MARKER, EXAMPLE_MARKER]

    def test_example_rows_skipped_by_default(self, tmp_path):
        preview = preview_bill_import(self._template(tmp_path))
        assert preview.total_rows == 0
        assert preview.example_rows_skipped == 2

    def test_example_rows_are_self_consistent(self, tmp_path):
        """示例行必须自洽，否则用户照抄会一头雾水。"""
        preview = preview_bill_import(self._template(tmp_path), skip_example_rows=False)
        assert preview.total_rows == 2
        assert preview.invalid_count == 0
        assert preview.valid_count == 2
        for row in preview.rows:
            assert row.bill is not None
            assert row.bill.energy_total_kwh is not None
            assert row.bill.bill_total_yuan is not None

    def test_enum_dropdowns_present(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_BILLS]
        formulas = [dv.formula1 for dv in worksheet.data_validations.dataValidation]
        assert len(formulas) >= 3
        assert any("单一制" in (formula or "") for formula in formulas)
        assert any("kV" in (formula or "") for formula in formulas)

    def test_tariff_sheet_has_no_data_rows(self, tmp_path):
        """§0.2 红线：分时电价工作表只给表头框架，**不得**预填任何示例电价。"""
        worksheet = load_workbook(self._template(tmp_path))[SHEET_TARIFF]
        assert worksheet.max_row == 1, "分时电价工作表不得有示例数据行（禁止示例电价）"
        headers = [str(cell.value) for cell in worksheet[1]]
        assert any("尖峰单价" in header for header in headers)

    def test_month_sheet_has_no_price_columns(self, tmp_path):
        """账单事实表不得出现"单价"列（单价属于 V2.3 电价计划）。"""
        headers = [str(cell.value) for cell in load_workbook(self._template(tmp_path))[SHEET_BILLS][1]]
        assert not [header for header in headers if "单价" in header]

    def test_dictionary_sheet_documents_enums(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_DICT]
        text = "\n".join(
            str(cell.value) for row in worksheet.iter_rows() for cell in row if cell.value
        )
        assert "single_part" in text and "two_part" in text
        assert "深谷 ≠ 低谷" in text
        assert "不是负荷曲线最大值" in text

    def test_help_sheet_lists_fields(self, tmp_path):
        worksheet = load_workbook(self._template(tmp_path))[SHEET_HELP]
        text = "\n".join(
            str(cell.value) for row in worksheet.iter_rows() for cell in row if cell.value
        )
        assert "填写说明" in text
        assert "留空 = 未提供" in text
        assert "不得" in text or "请勿" in text

    def test_template_bytes_are_readable(self):
        from io import BytesIO

        workbook = load_workbook(BytesIO(bill_template_bytes()))
        assert workbook.sheetnames == [SHEET_HELP, SHEET_BILLS, SHEET_TARIFF, SHEET_DICT]

    def test_template_in_directory(self, tmp_path):
        assert build_bill_template(tmp_path).parent == tmp_path


class TestColumnMapping:
    """§5.5：支持中文/英文列名、列顺序变化与用户重命名。"""

    def test_chinese_template_headers_map_all_fields(self):
        mapping = resolve_bill_columns(HEADERS)
        assert mapping.mapping["billing_period_start"] == "账期起(YYYY-MM-DD)"
        assert mapping.mapping["energy_offpeak_kwh"] == "深谷电量(kWh)"
        assert mapping.mapping["billing_demand_kw"] == "账单计费需量(kW)"
        assert len(mapping.mapping) == len([c for c in BILL_COLUMNS if c.field])

    def test_english_headers(self):
        mapping = resolve_bill_columns(
            ["billing_period_start", "billing_period_end", "energy_total_kwh", "bill_total_yuan"]
        )
        assert mapping.mapping["billing_period_start"] == "billing_period_start"
        assert mapping.mapping["bill_total_yuan"] == "bill_total_yuan"

    def test_column_order_change_is_irrelevant(self):
        reordered = list(reversed(HEADERS))
        mapping = resolve_bill_columns(reordered)
        assert mapping.mapping["energy_total_kwh"] == "总购电量(kWh)"
        assert mapping.mapping["billing_period_start"] == "账期起(YYYY-MM-DD)"

    def test_extra_unknown_column_ignored(self):
        mapping = resolve_bill_columns([*HEADERS, "内部编号"])
        assert "内部编号" in mapping.unmapped_headers
        assert "billing_period_start" in mapping.mapping

    def test_user_renamed_columns(self):
        mapping = resolve_bill_columns(
            ["起始日期", "结束日期", "表号", "用电量合计", "应收电费"]
        )
        assert mapping.mapping["billing_period_start"] == "起始日期"
        assert mapping.mapping["meter_id"] == "表号"
        assert mapping.mapping["energy_total_kwh"] == "用电量合计"
        assert mapping.mapping["bill_total_yuan"] == "应收电费"

    def test_missing_required_column_raises_chinese(self):
        with pytest.raises(ValidationError) as exc:
            resolve_bill_columns(["总购电量(kWh)"])
        assert "缺少必需列" in exc.value.message
        assert exc.value.field.startswith("bill.import.")

    def test_empty_headers_raise(self):
        with pytest.raises(ValidationError, match="没有可用表头"):
            resolve_bill_columns(["", "   "])

    def test_explicit_mapping_wins(self):
        mapping = resolve_bill_columns(HEADERS, explicit={"energy_total_kwh": "尖峰电量(kWh)"})
        assert mapping.mapping["energy_total_kwh"] == "尖峰电量(kWh)"

    def test_explicit_mapping_unknown_header_raises(self):
        with pytest.raises(ValidationError, match="不存在"):
            resolve_bill_columns(HEADERS, explicit={"energy_total_kwh": "不存在的列"})

    def test_longest_alias_wins_for_deep_valley(self):
        """短别名（谷电量）不得抢走深谷列，深谷与低谷必须各归其位。"""
        mapping = resolve_bill_columns(
            ["低谷电量(kWh)", "深谷电量(kWh)", "账期起(YYYY-MM-DD)", "账期止(YYYY-MM-DD)"]
        )
        assert mapping.mapping["energy_valley_kwh"] == "低谷电量(kWh)"
        assert mapping.mapping["energy_offpeak_kwh"] == "深谷电量(kWh)"

    def test_short_header_prefix_matches_full_alias(self):
        """真实账单常写"账期起"，不带格式提示也应能识别。"""
        mapping = resolve_bill_columns(["账期起", "账期止", "总电量"])
        assert mapping.mapping["billing_period_start"] == "账期起"
        assert mapping.mapping["energy_total_kwh"] == "总电量"


class TestPreview:
    def test_valid_row(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        preview = preview_bill_import(path, project_id="P1")
        assert preview.total_rows == 1
        assert preview.valid_count == 1 and preview.invalid_count == 0
        bill = preview.rows[0].bill
        assert bill is not None
        assert bill.billing_month == "2026-01"
        assert bill.project_id == "P1"
        assert bill.source_type.value == "excel"
        assert bill.source_file_name == "账单.xlsx"
        assert bill.source_row_number == 2

    def test_row_numbers_match_excel_file(self, tmp_path):
        rows = [HEADERS, *_valid_pair()[1:], _row(billing_period_start="2026-02-01", billing_period_end="2026-02-28", energy_total_kwh=1000, bill_total_yuan=700, meter_id="M002")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert [row.row_number for row in preview.rows] == [2, 3]

    def test_blank_rows_do_not_shift_row_numbers(self, tmp_path):
        rows = [HEADERS, [], *_valid_pair()[1:]]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert [row.row_number for row in preview.rows] == [3]

    def test_missing_optional_fields_stay_none(self, tmp_path):
        """§2.1：留空 = 未提供（None），不得自动填 0。"""
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        bill = preview_bill_import(path).rows[0].bill
        assert bill.energy_offpeak_kwh is None
        assert bill.vat_yuan is None
        assert bill.demand_charge_yuan is None
        assert bill.billing_demand_kw is None

    def test_warning_row_when_totals_do_not_match(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", energy_total_kwh=100_000, energy_peak_kwh=10_000, bill_total_yuan=999)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.warning_count == 1
        assert preview.rows[0].status is BillQualityStatus.WARNING
        assert any("B04" in message for message in preview.rows[0].messages)

    def test_invalid_row_bad_date(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="不是日期", billing_period_end="2026-01-31")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        message = preview.rows[0].messages[0]
        assert f"工作表「{SHEET_BILLS}」第 2 行" in message
        assert "账期起" in message
        assert "不是有效日期" in message

    def test_invalid_row_missing_period(self, tmp_path):
        rows = [HEADERS, _row(energy_total_kwh=1000)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        assert "必填" in preview.rows[0].messages[0]

    def test_invalid_row_negative_energy(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", energy_total_kwh=-5)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        assert any("B01" in message for message in preview.rows[0].messages)

    def test_invalid_row_non_numeric_text(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", bill_total_yuan="待定")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        assert "不是有效数值" in preview.rows[0].messages[0]

    def test_invalid_period_order(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-02-01", billing_period_end="2026-01-31")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        assert "不能早于" in preview.rows[0].messages[0]

    def test_month_period_mismatch_is_invalid(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", billing_month="2026-02")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        assert "账单月份" in preview.rows[0].messages[0]

    def test_cross_month_period_parsed_with_notice(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-26", billing_period_end="2026-02-25", energy_total_kwh=1000, bill_total_yuan=700)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.rows[0].bill.billing_month == "2026-01"
        assert any("跨月" in message for message in preview.rows[0].messages)

    def test_number_with_units_and_thousand_separator(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", energy_total_kwh="100,000", bill_total_yuan="65,000 元")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        bill = preview_bill_import(path).rows[0].bill
        assert bill.energy_total_kwh == pytest.approx(100_000.0)
        assert bill.bill_total_yuan == pytest.approx(65_000.0)

    def test_chinese_enum_values_accepted(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", tariff_structure="两部制", billing_demand_kw=450)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        bill = preview_bill_import(path).rows[0].bill
        assert bill.tariff_structure.value == "two_part"

    def test_english_enum_values_accepted(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", tariff_structure="two_part", billing_demand_kw=450)]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        assert preview_bill_import(path).rows[0].bill.tariff_structure.value == "two_part"

    def test_unknown_enum_value_is_invalid(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", tariff_structure="三制")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 1
        message = preview.rows[0].messages[0]
        assert "不是有效取值" in message and "单一制" in message

    def test_source_label_mismatch_is_notice_not_error(self, tmp_path):
        rows = [HEADERS, _row(billing_period_start="2026-01-01", billing_period_end="2026-01-31", source_type="手动录入")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.invalid_count == 0
        bill = preview.rows[0].bill
        assert bill.source_type.value == "excel"
        assert "表中标注来源" in (bill.notes or "")

    def test_example_rows_can_be_imported_explicitly(self, tmp_path):
        rows = [HEADERS, _row(EXAMPLE_MARKER, billing_period_start="2026-01-01", billing_period_end="2026-01-31")]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        assert preview_bill_import(path).total_rows == 0
        assert preview_bill_import(path, skip_example_rows=False).total_rows == 1

    def test_missing_marker_column_warns(self, tmp_path):
        headers = [header for header in HEADERS if header != MARKER_COLUMN]
        path = _write_xlsx(tmp_path / "账单.xlsx", [headers, ["2026-01-01", "2026-01-31"]])
        preview = preview_bill_import(path)
        assert any(MARKER_COLUMN in message for message in preview.messages)

    def test_duplicate_against_existing_bill(self, tmp_path):
        existing = ElectricityBill(
            bill_id=make_bill_id(date(2026, 1, 1), date(2026, 1, 31), "M001"),
            project_id="P1",
            billing_period_start=date(2026, 1, 1),
            billing_period_end=date(2026, 1, 31),
            meter_id="M001",
        )
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        preview = preview_bill_import(path, existing_bills=[existing], project_id="P1")
        assert preview.duplicate_count == 1
        assert preview.rows[0].duplicate_of == existing.bill_id
        assert any("疑似重复" in message for message in preview.rows[0].messages)

    def test_intra_file_duplicate_detected(self, tmp_path):
        rows = [HEADERS, *_valid_pair()[1:], _valid_pair()[1]]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        preview = preview_bill_import(path)
        assert preview.duplicate_count == 1
        assert preview.rows[0].duplicate_of is None
        assert "文件第 2 行" in (preview.rows[1].duplicate_of or "")

    def test_counts_text_is_chinese_summary(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        text = preview_bill_import(path).counts_text()
        assert "识别到 1 条记录" in text and "有效 1" in text

    def test_tariff_sheet_with_data_is_reported(self, tmp_path):
        path = _write_xlsx(
            tmp_path / "账单.xlsx",
            _valid_pair(),
            sheets={SHEET_TARIFF: [["账期起", "尖峰单价"], ["2026-01-01", 1.2]]},
        )
        preview = preview_bill_import(path)
        assert any(SHEET_TARIFF in message and "V2.1 不保存" in message for message in preview.messages)

    def test_sheet_selection(self, tmp_path):
        path = _write_xlsx(
            tmp_path / "账单.xlsx",
            [["总购电量(kWh)"]],
            sheet="其他表",
            sheets={SHEET_BILLS: _valid_pair()},
        )
        preview = preview_bill_import(path, sheet=SHEET_BILLS)
        assert preview.sheet_name == SHEET_BILLS
        assert preview.total_rows == 1
        # 未指定工作表时默认选模板的"月账单"
        assert preview_bill_import(path).sheet_name == SHEET_BILLS

    def test_csv_import_supported(self, tmp_path):
        path = tmp_path / "账单.csv"
        path.write_text(
            "账期起,账期止,总购电量(kWh),账单总额(元)\n2026-01-01,2026-01-31,100000,65000\n",
            encoding="utf-8-sig",
        )
        preview = preview_bill_import(path, project_id="P1")
        assert preview.total_rows == 1
        assert preview.rows[0].bill.energy_total_kwh == pytest.approx(100_000.0)
        assert preview.rows[0].row_number == 2


class TestApplyImport:
    def _preview(self, tmp_path, *, existing=(), name="账单.xlsx"):
        path = _write_xlsx(tmp_path / name, _valid_pair())
        return preview_bill_import(path, existing_bills=existing, project_id="P1")

    def test_skip_strategy(self, tmp_path):
        existing = ElectricityBill(
            bill_id="EXIST",
            project_id="P1",
            billing_period_start=date(2026, 1, 1),
            billing_period_end=date(2026, 1, 31),
            meter_id="M001",
        )
        result = apply_bill_import(self._preview(tmp_path, existing=[existing]), strategy=DuplicateStrategy.SKIP)
        assert result.added == []
        assert result.skipped_ids == ["BILL-20260101-20260131-M001"]
        assert result.total_written == 0

    def test_replace_strategy(self, tmp_path):
        existing = ElectricityBill(
            bill_id="EXIST",
            project_id="P1",
            billing_period_start=date(2026, 1, 1),
            billing_period_end=date(2026, 1, 31),
            meter_id="M001",
        )
        result = apply_bill_import(self._preview(tmp_path, existing=[existing]), strategy=DuplicateStrategy.REPLACE)
        assert len(result.replaced) == 1
        assert result.replaced[0].bill_total_yuan == pytest.approx(65_000.0)
        assert result.added == []

    def test_keep_both_strategy_adds_suffix(self, tmp_path):
        """保留两条时，新账单必须拿到不冲突的编号（否则项目里会撞号）。"""
        colliding = ElectricityBill(
            bill_id="BILL-20260101-20260131-M001",
            project_id="P1",
            billing_period_start=date(2026, 1, 1),
            billing_period_end=date(2026, 1, 31),
            meter_id="M001",
        )
        result = apply_bill_import(
            self._preview(tmp_path, existing=[colliding]), strategy=DuplicateStrategy.KEEP_BOTH
        )
        assert len(result.kept_both) == 1
        assert result.kept_both[0].bill_id == "BILL-20260101-20260131-M001-2"

    def test_keep_both_keeps_distinct_existing_id(self, tmp_path):
        """旧账单编号不同时，保留两条只需保证编号互不相同。"""
        existing = ElectricityBill(
            bill_id="EXIST",
            project_id="P1",
            billing_period_start=date(2026, 1, 1),
            billing_period_end=date(2026, 1, 31),
            meter_id="M001",
        )
        result = apply_bill_import(
            self._preview(tmp_path, existing=[existing]), strategy=DuplicateStrategy.KEEP_BOTH
        )
        assert result.kept_both[0].bill_id != "EXIST"
        assert result.kept_both[0].bill_id != ""

    def test_invalid_rows_excluded_and_listed(self, tmp_path):
        rows = [
            HEADERS,
            *_valid_pair()[1:],
            _row(billing_period_start="坏日期", billing_period_end="2026-03-31"),
        ]
        path = _write_xlsx(tmp_path / "账单.xlsx", rows)
        result = apply_bill_import(preview_bill_import(path))
        assert result.added_count == 1
        assert result.invalid_rows == [3]

    def test_reimport_does_not_write_twice(self, tmp_path):
        """§5.5：导入不得重复写入同一条账单。"""
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        first = apply_bill_import(preview_bill_import(path, project_id="P1"))
        assert first.added_count == 1
        second = apply_bill_import(
            preview_bill_import(path, existing_bills=first.added, project_id="P1"),
            strategy=DuplicateStrategy.SKIP,
        )
        assert second.added_count == 0
        assert second.skipped_ids

    def test_summary_text(self, tmp_path):
        result = apply_bill_import(self._preview(tmp_path))
        assert "新增 1 条" in result.summary_text()


class TestFatalErrors:
    """§9.1：空工作表、错误工作表、损坏文件等必须给出中文错误，不抛裸异常。"""

    def test_missing_file(self, tmp_path):
        with pytest.raises(ValidationError, match="文件不存在"):
            preview_bill_import(tmp_path / "无.xlsx")

    def test_unsupported_suffix(self, tmp_path):
        path = tmp_path / "账单.txt"
        path.write_text("x", encoding="utf-8")
        # .txt 按 CSV 处理，表头不可识别 → 中文缺列错误
        with pytest.raises(ValidationError, match="缺少必需列|没有数据行"):
            preview_bill_import(path)

    def test_unsupported_binary_suffix(self, tmp_path):
        # V2.5 起 .pdf 是受支持的账单格式；损坏的 PDF 报"无法打开"的中文错误，
        # 而不是"不支持的文件类型"（后者留给真正的未知后缀，如 .docx）。
        path = tmp_path / "账单.pdf"
        path.write_bytes(b"%PDF-1.4")  # 只有魔数、无有效结构
        with pytest.raises(ValidationError, match="无法打开 PDF 账单"):
            preview_bill_import(path)

    def test_unsupported_unknown_suffix(self, tmp_path):
        path = tmp_path / "账单.docx"
        path.write_bytes(b"PK\x03\x04")
        with pytest.raises(ValidationError, match="不支持的文件类型"):
            preview_bill_import(path)

    def test_wrong_sheet_lists_available(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", _valid_pair())
        with pytest.raises(ValidationError) as exc:
            preview_bill_import(path, sheet="不存在的表")
        assert "不存在" in exc.value.message
        assert SHEET_BILLS in exc.value.message

    def test_header_only_sheet(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", [HEADERS])
        with pytest.raises(ValidationError, match="没有数据行"):
            preview_bill_import(path)

    def test_empty_workbook_sheet(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", [])
        with pytest.raises(ValidationError, match="没有数据行|没有可用的表头"):
            preview_bill_import(path)

    def test_help_sheet_missing_required_columns(self, tmp_path):
        template = build_bill_template(tmp_path / "模板")
        with pytest.raises(ValidationError, match="缺少必需列"):
            preview_bill_import(template, sheet=SHEET_HELP)

    def test_corrupt_xlsx_gives_chinese_error(self, tmp_path):
        path = tmp_path / "坏文件.xlsx"
        path.write_bytes("这不是一个 xlsx 文件".encode("utf-8"))
        with pytest.raises(ValidationError) as exc:
            preview_bill_import(path)
        assert "不是有效的 Excel 工作簿" in exc.value.message

    def test_corrupt_xlsx_with_zip_header(self, tmp_path):
        path = tmp_path / "坏文件2.xlsx"
        path.write_bytes(b"PK\x03\x04" + b"\x00" * 64)
        with pytest.raises(ValidationError, match="不是有效的 Excel 工作簿"):
            preview_bill_import(path)

    def test_missing_required_columns(self, tmp_path):
        path = _write_xlsx(tmp_path / "账单.xlsx", [["总购电量(kWh)"], [1000]])
        with pytest.raises(ValidationError, match="缺少必需列"):
            preview_bill_import(path)
