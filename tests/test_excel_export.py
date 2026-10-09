"""Excel 导出测试（规范 §108、§109、§154）。

核心断言：**Excel 中的数值必须与 CalculationResult 完全一致**，且导出层不得引入公式重算。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from cenep.calculation.engine import calculation_engine
from cenep.reports.excel_exporter import (
    CALIBER_TAG_ACTUAL,
    CALIBER_TAG_ESTIMATED,
    CALIBER_TAG_RECOMPUTED,
    CALIBER_TAG_SIMULATED,
    SHEET_NAMES,
    ExcelExporter,
)

#: V2.4 §8.1（阶段 7）新增的 5 张表（**追加**在既有 27 张之后）
V24_SHEETS = ("月度电费分析", "负荷数据质量", "电价版本与来源", "方案电费对比", "计算假设与警告")
#: V2.4 之前就存在的表清单 = 完整清单去掉追加的 5 张（用来断言"只是追加"）
V1_V24_BASE = list(SHEET_NAMES[: len(SHEET_NAMES) - len(V24_SHEETS)])


@pytest.fixture
def exported(golden_pv_storage, tmp_path: Path):
    result = calculation_engine.calculate(golden_pv_storage)
    exporter = ExcelExporter()
    path = exporter.export(golden_pv_storage, result, tmp_path / "报告.xlsx")
    return golden_pv_storage, result, path


class TestWorkbookStructure:
    def test_file_created_with_dev_suffix(self, exported):
        _, _, path = exported
        assert path.exists()
        assert path.suffix == ".xlsx"

    def test_suffix_is_appended_automatically(self, golden_pv_storage, tmp_path: Path):
        result = calculation_engine.calculate(golden_pv_storage)
        path = ExcelExporter().export(golden_pv_storage, result, tmp_path / "无后缀")
        assert path.name == "无后缀.xlsx"

    def test_dotted_name_not_truncated(self, golden_pv_storage, tmp_path: Path):
        """回归：文件名里的 2061.8kWp / V1.2 这类片段不得被 with_suffix 截掉。"""
        result = calculation_engine.calculate(golden_pv_storage)
        path = ExcelExporter().export(golden_pv_storage, result, tmp_path / "全屋面2061.8kWp_经济评价")
        assert path.name == "全屋面2061.8kWp_经济评价.xlsx"
        assert path.exists()

    def test_sheets_match_names_and_count(self, exported):
        """V2 §67 + V2.1 §8.1 + V2.2 §6.3 + V2.4 §8.1：V1 §108 的 13 张 + V2 新增 11 张时序表
        + V2.1 新增 2 张账单表 + V2.2 阶段 4 新增 1 张「消纳率分析」
        + V2.4 阶段 7 新增 5 张 = 32 张。

        **表数一律取自 :data:`SHEET_NAMES` 的长度**，不再硬编码数字（V2.4 §8.1）。

        这是 V2 §67「至少包含」、V2.1 §8.1、V2.2 §6.3 与 V2.4 §8.1 对 V1 §108 的
        **正当超集扩展**：V1 的 13 张全部保留且相对顺序不变，新增表在无对应数据时输出中文说明，
        因此 V1 项目、无账单项目与无负荷项目仍可正常导出（见 test_v1_project_still_exports）。
        """
        _, _, path = exported
        wb = load_workbook(path)
        assert wb.sheetnames == SHEET_NAMES
        assert len(wb.sheetnames) == len(SHEET_NAMES)

    def test_v22_self_consumption_sheet_included(self, exported):
        """V2.2 §6.3 C：消纳率分析表必须存在；未执行消纳分析时输出中文说明。"""
        _, _, path = exported
        names = set(load_workbook(path).sheetnames)
        assert "消纳率分析" in names
        ws = load_workbook(path, data_only=True)["消纳率分析"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "尚未执行负荷与消纳分析" in text

    def test_v21_bill_sheets_included(self, exported):
        """V2.1 §8.1：两张账单表必须存在于工作簿中（无账单时也照常生成）。"""
        _, _, path = exported
        names = set(load_workbook(path).sheetnames)
        for required in ("账单原始数据", "账单校验"):
            assert required in names, f"缺少 V2.1 §8.1 要求的表：{required}"

    def test_v2_sheets_included(self, exported):
        """V2 §67 要求的新增表必须存在。"""
        _, _, path = exported
        names = set(load_workbook(path).sheetnames)
        for required in (
            "储能与调度", "负荷曲线", "光伏曲线", "分时电价", "8760时序仿真",
            "能量平衡", "年度汇总", "收益分解", "方案比较", "方案寻优", "数据质量",
        ):
            assert required in names, f"缺少 V2 §67 要求的表：{required}"

    def test_v1_sheets_preserved(self, exported):
        """V1 的 13 张表一张都不能少。"""
        _, _, path = exported
        names = set(load_workbook(path).sheetnames)
        for legacy in (
            "项目概况", "基础参数", "技术参数", "电价参数", "投资参数", "运维参数",
            "融资参数", "年度现金流", "财务指标", "敏感性分析", "情景分析",
            "政策依据", "参数来源",
        ):
            assert legacy in names, f"V1 表被破坏：{legacy}"

    def test_v24_new_sheets_included(self, exported):
        """V2.4 §8.1（阶段 7）：5 张新增表必须存在；无数据时输出中文说明而非缺表。"""
        _, _, path = exported
        wb = load_workbook(path)
        names = set(wb.sheetnames)
        for required in V24_SHEETS:
            assert required in names, f"缺少 V2.4 §8.1 要求的表：{required}"
        for required in ("月度电费分析", "负荷数据质量", "方案电费对比"):
            text = "\n".join(
                str(c.value) for row in wb[required].iter_rows() for c in row if c.value is not None
            )
            assert text.strip(), f"{required} 内容为空"

    def test_v24_new_sheets_keep_v1_order(self, exported):
        """V2.4 的新表是**追加**的：既有 27 张表的相对顺序与表名一字未动。"""
        _, _, path = exported
        names = load_workbook(path).sheetnames
        assert names[: len(V1_V24_BASE)] == V1_V24_BASE
        assert names[len(V1_V24_BASE):] == list(V24_SHEETS)

    def test_v24_new_sheets_have_no_formula(self, exported):
        """§109 / V2 §61：新增表同样不得写公式。"""
        _, _, path = exported
        wb = load_workbook(path)
        offenders = [
            f"{ws.title}!{cell.coordinate}"
            for ws in wb.worksheets
            if ws.title in V24_SHEETS
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert offenders == []

    def test_v24_scenario_compare_placeholder_explains_how_to(self, exported):
        """无场景结果时，「方案电费对比」必须写明操作路径（不缺表、不臆造数值）。"""
        _, _, path = exported
        ws = load_workbook(path, data_only=True)["方案电费对比"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "尚未执行光储四场景账单对比" in text
        assert "执行四场景对比" in text

    def test_v24_caliber_legend_lists_four_data_types(self, exported):
        """§8.1：实际账单 / 软件复算 / 模型估算 / 方案模拟四类口径必须逐条说明。"""
        _, _, path = exported
        ws = load_workbook(path, data_only=True)["计算假设与警告"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        for label in CALIBER_TAG_ACTUAL, CALIBER_TAG_RECOMPUTED, CALIBER_TAG_ESTIMATED, CALIBER_TAG_SIMULATED:
            assert label in text, f"计算假设与警告缺少口径标签：{label}"

    def test_declared_sheet_names_match_actual(self, exported):
        _, result, path = exported
        wb = load_workbook(path)
        assert ExcelExporter().export_sheets(None, result) == wb.sheetnames


class TestValuesComeFromResult:
    def test_cashflow_sheet_matches_annual_results(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["年度现金流"]
        # 第 3 行是表头，第 4 行是 Year 0，第 5 行起是 Year 1..N
        assert ws.cell(row=4, column=1).value == 0
        assert ws.cell(row=4, column=13).value == pytest.approx(result.project_cashflows[0])
        assert ws.cell(row=4, column=14).value == pytest.approx(result.equity_cashflows[0])
        for offset, row in enumerate(result.annual_results):
            excel_row = 5 + offset
            assert ws.cell(row=excel_row, column=1).value == row.year
            assert ws.cell(row=excel_row, column=7).value == pytest.approx(row.total_revenue)
            assert ws.cell(row=excel_row, column=13).value == pytest.approx(row.project_cashflow)
            assert ws.cell(row=excel_row, column=14).value == pytest.approx(row.equity_cashflow)
            assert ws.cell(row=excel_row, column=15).value == pytest.approx(row.cumulative_project_cashflow)

    def test_metrics_sheet_matches_result(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["财务指标"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert f"{result.project_irr:.6f}" in text
        assert "LCOE" in text and "LCOS" in text
        assert "最低 DSCR" in text

    def test_total_capex_matches(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["投资参数"]
        values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
        assert any(v == pytest.approx(result.total_capex) for v in values)

    def test_sensitivity_rows_match(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["敏感性分析"]
        # 表头在第 3 行，数据从第 4 行开始
        data_rows = [r for r in ws.iter_rows(min_row=4) if r[0].value]
        assert len(data_rows) == len(result.sensitivity)

    def test_scenario_rows_match(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["情景分析"]
        data_rows = [r for r in ws.iter_rows(min_row=4) if r[0].value]
        assert len(data_rows) == len(result.scenarios)

    def test_parameter_sources_rows_match(self, exported):
        _, result, path = exported
        ws = load_workbook(path, data_only=True)["参数来源"]
        data_rows = [r for r in ws.iter_rows(min_row=4) if r[0].value]
        assert len(data_rows) == len(result.parameter_sources)

    def test_no_formulas_written(self, exported):
        """§109：导出层不重新计算。本实现只写数值，因此工作簿内不应出现公式。"""
        _, _, path = exported
        wb = load_workbook(path)
        formula_cells = [
            f"{ws.title}!{cell.coordinate}"
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        assert formula_cells == []


class TestPolicySheet:
    def test_without_policy(self, exported):
        _, _, path = exported
        ws = load_workbook(path, data_only=True)["政策依据"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "未关联政策" in text
        assert "请自行核对现行政策" in text

    def test_with_hubei_policy(self, golden_pv_storage, hubei_policy, tmp_path: Path):
        project = golden_pv_storage.model_copy(deep=True)
        project.policy = hubei_policy
        result = calculation_engine.calculate(project)
        path = ExcelExporter().export(project, result, tmp_path / "p.xlsx")
        ws = load_workbook(path, data_only=True)["政策依据"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "2025-10-01" in text
        assert "本测算采用政策" in text


class TestThreeProjectTypes:
    @pytest.mark.parametrize("fixture_name", ["golden_pv", "golden_storage", "golden_pv_storage"])
    def test_all_three_types_export(self, fixture_name, request, tmp_path: Path):
        project = request.getfixturevalue(fixture_name)
        result = calculation_engine.calculate(project)
        path = ExcelExporter().export(project, result, tmp_path / f"{fixture_name}.xlsx")
        wb = load_workbook(path)
        assert wb.sheetnames == SHEET_NAMES
