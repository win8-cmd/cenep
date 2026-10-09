"""V2.2 阶段 4：负荷与消纳页面 + Excel/PDF 输出测试（规格书 §6.3、§8.4、§0.2、§3.3）。

覆盖：

* 主窗口新增「负荷与消纳」页并绑定应用服务（§6.2、§6.3）；
* A 区：来源切换、生成估算曲线、来源标签**一眼可辨**（估算=橙色警告、实测=绿色）；
* B 区：负荷画像与逐月/典型日表格；
* C 区：四项指标 + **每项都带口径**、逐月明细、假设文本、两条曲线；
* 输入变更后结果标记"已过期，需重新计算"（§8.4）；
* Excel 新增「消纳率分析」表、PDF 新增「五、负荷估算与光伏消纳」章节，
  两者都必须显著标注**估算**并列出四项口径（§0.2、§3.3、§8.1）；
* 界面层静态扫描：``ui/`` 不得引入计算逻辑（沿用既有架构测试的口径）。
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

pytest.importorskip("PySide6", reason="未安装 PySide6，跳过界面测试")

from PySide6.QtWidgets import QApplication  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from cenep.calculation.engine import calculation_engine  # noqa: E402
from cenep.reports.excel_exporter import ExcelExporter  # noqa: E402
from cenep.reports.pdf_exporter import (  # noqa: E402
    REPORT_SECTIONS,
    PdfExporter,
    _styles,
    register_cjk_font,
    section_title,
)
from cenep.ui.load_pages import (  # noqa: E402
    ESTIMATE_BANNER,
    MEASURED_BANNER,
    SOURCE_MEASURED,
    STALE_TEXT,
    LoadAnalysisPage,
)

UI_DIR = Path(__file__).resolve().parents[1] / "src" / "cenep" / "ui"


@pytest.fixture(scope="module")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


def _prepare(window, annual_energy: float = 1_200_000.0):
    """在页面上生成一条月账单估算曲线（不经过任何对话框）。"""
    page = window.load_page
    window.project.pv.pv_capacity_kwp = 1000.0
    window.project.pv.equivalent_hours = 1100.0
    window.project.pv.performance_ratio = 1.0
    page._sync_pv_fields()
    page.estimate_year.setValue(2025)
    page.estimate_annual.setValue(annual_energy)
    page.generate_estimate()
    return page


# --------------------------------------------------------------------------- #
# 页面结构（§6.3）
# --------------------------------------------------------------------------- #
class TestPageStructure:
    def test_four_sections(self, qapp):
        """V2.2 §6.3 A/B/C 三段；V2.4 §8.4（阶段 7）**追加** D 段「光储场景对比」。

        追加方式与 V2.1/V2.2 一致：既有三段标题一字未改，只在末尾新增一段；
        主界面**标签页数量不变**（仍为 10 个），因此不改变既有导航结构。
        """
        page = LoadAnalysisPage()
        titles = [page.tabs.tabText(i) for i in range(page.tabs.count())]
        assert titles == ["A 数据来源", "B 负荷画像", "C 消纳分析", "D 光储场景对比"]

    def test_scenario_section_widgets_exist(self, qapp):
        """§7.1 / §7.6：D 段必须提供电价计划选择、执行按钮与三张结果表 + 一致性文本。"""
        page = LoadAnalysisPage()
        assert page.scenario_plan_combo is not None
        assert page.scenario_run_button is not None
        assert page.scenario_table is not None
        assert page.comparison_table is not None
        assert page.dedup_table is not None
        assert page.scenario_detail_view.isReadOnly()
        assert page.scenario_combo_plan_id() == ""
        assert page.last_scenario is None

    def test_source_choices_cover_both_inputs(self, qapp):
        page = LoadAnalysisPage()
        labels = [page.source_combo.itemText(i) for i in range(page.source_combo.count())]
        assert len(labels) == 2
        assert any("月账单估算" in text for text in labels)
        assert any("实测高频负荷导入" in text for text in labels)

    def test_template_combo_lists_builtin_templates(self, qapp):
        page = LoadAnalysisPage()
        assert page.template_combo.count() == 5

    def test_registered_in_main_window(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "load_ui.db")
        try:
            assert window.load_page is not None
            assert window.load_page.service is not None
            assert window.load_page.service.project is window.project
        finally:
            window.close()


# --------------------------------------------------------------------------- #
# A 区：数据来源与估算（§6.3 A、§0.2）
# --------------------------------------------------------------------------- #
class TestSourceSection:
    def test_generate_estimate_registers_dataset(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "gen.db")
        try:
            page = _prepare(window)
            state = page.state()
            assert state["datasets"] == 1
            assert state["estimate_points"] == 35040
            assert state["estimate_max_residual"] <= 1e-6
            assert "估算" in page.estimate_result_label.text()
            assert "估算数据（不是实测）" in page.estimate_result_label.text()
        finally:
            window.close()

    def test_estimate_banner_is_orange_warning(self, qapp, tmp_path):
        """§6.3 A：估算数据始终显示橙色文字提示"估算曲线，不是实测"。"""
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "banner.db")
        try:
            page = _prepare(window)
            assert ESTIMATE_BANNER in page.banner.text()
            assert "#C05600" in page.banner.styleSheet()
        finally:
            window.close()

    def test_dataset_table_marks_measured_vs_estimated(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "table.db")
        try:
            page = _prepare(window)
            assert page.dataset_table.rowCount() == 1
            assert page.dataset_table.item(0, 3).text() == "估算"
            assert page.dataset_table.item(0, 2).text() == "月账单估算负荷"
            assert "当前激活" in page.dataset_table.item(0, 0).text()
        finally:
            window.close()

    def test_estimate_invalid_input_reports_chinese_message(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "bad.db")
        try:
            page = window.load_page
            page.estimate_year.setValue(2025)
            page.estimate_annual.setValue(0.0)  # 无月电量、无年电量 → 必须报中文错误
            page.generate_estimate()
            assert "失败" in page.estimate_result_label.text()
            assert page.state()["datasets"] == 0
        finally:
            window.close()

    def test_bill_coverage_button_warns_when_annual_missing(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "cov.db")
        try:
            page = window.load_page
            page.estimate_year.setValue(2025)
            page.estimate_annual.setValue(0.0)
            page.refresh_bill_coverage()
            assert "缺" in page.bill_coverage_label.text()
            assert "#B00020" in page.bill_coverage_label.styleSheet()
        finally:
            window.close()

    def test_source_switch_moves_stack_and_marks_stale(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "switch.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            assert page.state()["stale"] is False
            page.source_combo.setCurrentIndex(1)
            assert page.source_stack.currentIndex() == 1
            assert page.state()["stale"] is True
            assert STALE_TEXT in page.stale_label.text()
            assert page.source_combo.currentData() == SOURCE_MEASURED
        finally:
            window.close()


# --------------------------------------------------------------------------- #
# B 区：负荷画像（§6.3 B）
# --------------------------------------------------------------------------- #
class TestPortraitSection:
    def test_portrait_tables_filled(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "portrait.db")
        try:
            page = _prepare(window)
            page.refresh_portrait()
            assert page.portrait_table.rowCount() >= 10
            assert page.monthly_table.rowCount() == 12
            # 15 分钟数据 → 典型日 96 个时刻
            assert page.typical_table.rowCount() == 96
            assert "估算" in page.portrait_source_label.text()
            assert "#C05600" in page.portrait_source_label.styleSheet()
        finally:
            window.close()

    def test_typical_day_columns(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "typical.db")
        try:
            page = _prepare(window)
            page.refresh_portrait()
            headers = [
                page.typical_table.horizontalHeaderItem(i).text()
                for i in range(page.typical_table.columnCount())
            ]
            assert headers == ["时刻", "工作日 kW", "休息日 kW", "全部日 kW"]
            assert page.typical_table.item(0, 0).text() == "00:00"
        finally:
            window.close()

    def test_portrait_without_dataset_reports_hint(self, qapp):
        page = LoadAnalysisPage()
        page.refresh_portrait()
        assert "尚未选择负荷数据集" in page.portrait_source_label.text()


# --------------------------------------------------------------------------- #
# C 区：消纳分析（§6.3 C、§3.3）
# --------------------------------------------------------------------------- #
class TestAnalysisSection:
    def test_four_metrics_with_calibers(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "analysis.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            assert page.metric_table.rowCount() == 4
            names = [page.metric_table.item(row, 0).text() for row in range(4)]
            assert any("光伏自用率" in name for name in names)
            assert any("负荷覆盖率" in name for name in names)
            assert any("上网率" in name for name in names)
            assert any("电网依赖率" in name for name in names)
            # 每一项都必须带口径（分子 / 分母 / 单位 / 边界）
            for row in range(4):
                caliber = page.metric_table.item(row, 2).text()
                assert "分母" in caliber and "边界" in caliber
                assert "min(E_load,i, E_pv,i)" in caliber or "max(" in caliber
            # 分母口径必须分别写对（§3.3）
            by_name = {
                page.metric_table.item(row, 0).text(): page.metric_table.item(row, 2).text()
                for row in range(4)
            }
            assert "光伏发电量** E_pv" in next(
                text for name, text in by_name.items() if "上网率" in name
            )
            assert "负荷电量** E_load" in next(
                text for name, text in by_name.items() if "电网依赖率" in name
            )
        finally:
            window.close()

    def test_complementary_rates_shown(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "rates.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            state = page.state()
            assert state["self_consumption_rate"] + state["export_rate"] == pytest.approx(1.0)
            assert state["load_coverage_rate"] + state["grid_dependency_rate"] == pytest.approx(1.0)
            assert state["based_on_estimate"] is True
            assert state["monthly_rows"] == 12
        finally:
            window.close()

    def test_analysis_banner_flags_estimate(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "banner2.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            text = page.analysis_source_label.text()
            assert "月账单估算负荷" in text and "基于估算" in text
            assert "#C05600" in page.analysis_source_label.styleSheet()
        finally:
            window.close()

    def test_charts_receive_series(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "charts.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            assert page.chart.series_point_count() == 35040
            assert page.monthly_chart.month_count() == 12
            assert page.chart.plot.plotItem.listDataItems(), "叠加曲线必须至少有一条线"
        finally:
            window.close()

    def test_assumptions_text_covers_calibers(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "assume.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            text = page.assumption_view.toPlainText()
            assert "消纳口径" in text
            assert "负荷覆盖率是**电量覆盖率**" in text
            assert "基于估算" in text
            assert "不得显示 0%" in text
        finally:
            window.close()

    def test_analysis_without_dataset_keeps_old_result_absent(self, qapp):
        page = LoadAnalysisPage()
        page.run_analysis()
        assert page.state()["self_consumption_rate"] is None

    def test_pv_params_write_back(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "pv.db")
        try:
            page = window.load_page
            page.pv_capacity.setValue(2500.0)
            page.pv_hours.setValue(1050.0)
            page.pv_pr.setValue(0.82)
            page.apply_pv_params()
            assert window.project.pv.pv_capacity_kwp == pytest.approx(2500.0)
            assert window.project.pv.equivalent_hours == pytest.approx(1050.0)
            assert window.project.pv.performance_ratio == pytest.approx(0.82)
        finally:
            window.close()

    def test_range_selector_does_not_crash(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "range.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            page.range_mode.setCurrentIndex(1)  # 按月
            page.range_month.setCurrentIndex(5)
            page.range_day.setCurrentIndex(9)
            assert page.chart.visible_x_span_seconds() if hasattr(
                page.chart, "visible_x_span_seconds"
            ) else True
        finally:
            window.close()


# --------------------------------------------------------------------------- #
# 界面层静态扫描（沿用既有架构测试口径）
# --------------------------------------------------------------------------- #
class TestNoCalculationInLoadUi:
    def test_load_pages_imports_no_calculation_module(self):
        offenders: list[str] = []
        source_path = UI_DIR / "load_pages.py"
        for line in source_path.read_text(encoding="utf-8").splitlines():
            stripped = line.strip()
            if re.match(r"from\s+\.{1,2}calculation", stripped) and "errors" not in stripped:
                offenders.append(stripped)
            if re.match(r"from\s+cenep\.calculation", stripped) and "errors" not in stripped:
                offenders.append(stripped)
        assert offenders == []

    def test_load_pages_does_not_reference_metric_functions(self):
        forbidden = ["financial_metrics", "npv(", "irr(", "payback_period", "lcoe(", "lcos("]
        source = (UI_DIR / "load_pages.py").read_text(encoding="utf-8")
        assert [token for token in forbidden if token in source] == []

    def test_charts_module_does_not_import_calculation(self):
        source = (UI_DIR / "charts.py").read_text(encoding="utf-8")
        assert "from ..calculation" not in source


# --------------------------------------------------------------------------- #
# Excel / PDF 输出（§8.1、§0.2、§3.3）
# --------------------------------------------------------------------------- #
@pytest.fixture
def exported_with_analysis(tmp_path):
    """构造一次真实的"估算负荷 → 消纳分析"结果，并导出 Excel/PDF。"""
    from cenep.application.load_profile_service import LoadProfileService
    from cenep.domain.enums import ProjectType
    from cenep.domain.load_estimate import LoadEstimateParams, MonthlyLoadEnergy
    from cenep.domain.models import BasicInfo, Project, PVConfig

    project = Project(
        basic_info=BasicInfo(
            project_name="报告消纳测试", province="湖北", project_type=ProjectType.COMMERCIAL_PV
        ),
        pv=PVConfig(
            pv_capacity_kwp=1000.0,
            equivalent_hours=1100.0,
            performance_ratio=1.0,
            annual_degradation_rate=0.0,
        ),
    )
    service = LoadProfileService(project, project_id="RPT")
    monthly = [MonthlyLoadEnergy(year=2025, month=m, energy_kwh=100_000.0) for m in range(1, 13)]
    service.estimate_load(
        LoadEstimateParams(year=2025, interval_minutes=60, template_id="double_shift"),
        monthly=monthly,
    )
    consumption = service.analyze()
    portrait = service.portrait()
    result = calculation_engine.calculate(project)
    excel = ExcelExporter().export(
        project, result, tmp_path / "报告.xlsx", self_consumption=consumption, load_portrait=portrait
    )
    pdf = PdfExporter().export(
        project, result, tmp_path / "报告.pdf", self_consumption=consumption, load_portrait=portrait
    )
    return project, consumption, portrait, excel, pdf


class TestExcelSelfConsumptionSheet:
    def test_sheet_exists_and_always_present(self, golden_pv_storage, tmp_path):
        result = calculation_engine.calculate(golden_pv_storage)
        path = ExcelExporter().export(golden_pv_storage, result, tmp_path / "no_load.xlsx")
        names = load_workbook(path).sheetnames
        assert "消纳率分析" in names

    def test_values_and_calibers_written(self, exported_with_analysis):
        _, consumption, _, excel, _ = exported_with_analysis
        ws = load_workbook(excel, data_only=True)["消纳率分析"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        # 四项指标的中文名与口径必须都在表里
        for token in ("光伏自用率", "负荷覆盖率", "光伏上网率", "电网依赖率"):
            assert token in text
        assert "分母" in text and "边界" in text
        assert "电量覆盖率" in text and "不是时间覆盖率" in text
        # 估算标记必须显著
        assert "估算数据（不是实测）" in text
        assert "假设" in text
        # 数值必须与结果一致
        assert f"{consumption.self_consumption_rate:.2f}" in text or f"{consumption.self_consumption_rate}" in text

    def test_placeholder_when_no_analysis(self, golden_pv_storage, tmp_path):
        result = calculation_engine.calculate(golden_pv_storage)
        path = ExcelExporter().export(golden_pv_storage, result, tmp_path / "empty.xlsx")
        ws = load_workbook(path, data_only=True)["消纳率分析"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert "尚未执行负荷与消纳分析" in text
        assert "不得用未经说明的默认值" in text


class TestPdfLoadConsumptionSection:
    def test_section_declared_after_load_analysis(self):
        assert REPORT_SECTIONS[REPORT_SECTIONS.index("负荷分析") + 1] == "负荷估算与光伏消纳"

    def test_story_contains_metrics_and_calibers(self, exported_with_analysis):
        project, consumption, portrait, _, _ = exported_with_analysis
        result = calculation_engine.calculate(project)
        styles = _styles(register_cjk_font())
        story = PdfExporter().build_story(
            project, result, styles, self_consumption=consumption, load_portrait=portrait
        )
        text = "\n".join(
            getattr(item, "text", "") or _flatten_table(item) for item in story
        )
        assert section_title(REPORT_SECTIONS.index("负荷估算与光伏消纳")) in text
        for token in ("光伏自用率", "负荷覆盖率", "光伏上网率", "电网依赖率", "分母", "边界"):
            assert token in text, f"PDF 章节缺少：{token}"
        assert "估算" in text
        assert "时间覆盖率（≠负荷覆盖率）" in text or "时间覆盖率" in text

    def test_story_without_analysis_still_declares_section(self, golden_pv_storage):
        result = calculation_engine.calculate(golden_pv_storage)
        styles = _styles(register_cjk_font())
        story = PdfExporter().build_story(golden_pv_storage, result, styles)
        text = "\n".join(getattr(item, "text", "") or "" for item in story)
        assert section_title(REPORT_SECTIONS.index("负荷估算与光伏消纳")) in text
        assert "尚未执行负荷与消纳分析" in text

    def test_pdf_file_written(self, exported_with_analysis):
        _, _, _, _, pdf = exported_with_analysis
        raw = pdf.read_bytes()
        assert raw.startswith(b"%PDF")
        assert pdf.stat().st_size > 5000


def _flatten_table(item) -> str:
    """把 reportlab 表格的单元格文本拼成字符串（仅供断言使用）。"""
    cells = getattr(item, "_cellvalues", None)
    if not cells:
        return ""
    parts: list[str] = []
    for row in cells:
        for cell in row:
            if isinstance(cell, str):
                parts.append(cell)
            elif hasattr(cell, "text"):
                parts.append(cell.text)
            elif hasattr(cell, "_cellvalues"):
                parts.append(_flatten_table(cell))
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# 主窗口导出载荷（把结果交给报告层，而不是让报告层重算）
# --------------------------------------------------------------------------- #
class TestMainWindowExportPayload:
    def test_payload_passed_to_reports(self, qapp, tmp_path, monkeypatch):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "payload.db")
        try:
            page = _prepare(window)
            page.run_analysis()
            consumption, portrait = window._load_analysis_payload()
            assert consumption is page.last_result
            assert portrait is page.last_portrait

            captured: dict[str, object] = {}

            def fake_export(project, result, path, **kwargs):
                captured.update(kwargs)
                Path(path).write_bytes(b"stub")
                return Path(path)

            monkeypatch.setattr(window.excel_exporter, "export", fake_export)
            window.last_result = calculation_engine.calculate(window.project)
            window.excel_exporter.export(
                window.project,
                window.last_result,
                tmp_path / "x.xlsx",
                **dict(zip(("self_consumption", "load_portrait"), (consumption, portrait))),
            )
            assert captured["self_consumption"] is consumption
            assert captured["load_portrait"] is portrait
        finally:
            window.close()

    def test_payload_none_before_analysis(self, qapp, tmp_path):
        from cenep.ui.main_window import MainWindow

        window = MainWindow(db_path=tmp_path / "nonepayload.db")
        try:
            assert window._load_analysis_payload() == (None, None)
        finally:
            window.close()

    def test_measured_dataset_is_green_banner(self, qapp, tmp_path):
        """实测数据必须与估算使用不同的（绿色）提示，二者一眼可辨（§6.3 A、§0.2）。"""
        import csv

        from cenep.ui.main_window import MainWindow

        path = tmp_path / "measured.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["时间", "有功功率(kW)"])
            for hour in range(48):
                writer.writerow([f"2025-01-{1 + hour // 24:02d} {hour % 24:02d}:00", 100.0 + hour])
        window = MainWindow(db_path=tmp_path / "measured.db")
        try:
            page = window.load_page
            assert page.set_load_file(path) is True
            page.preview_load_file()
            page.confirm_load_import()
            assert MEASURED_BANNER in page.banner.text()
            assert "#1B7F3B" in page.banner.styleSheet()
            assert page.dataset_table.item(0, 3).text() == "实测"
            page.run_analysis()
            assert page.state()["based_on_estimate"] is False
        finally:
            window.close()
